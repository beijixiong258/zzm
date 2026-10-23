"""Stage-2 real HTTP acceptance, isolated synthetic data only.

No real model requests. Public Link extraction may perform bounded HTTPS reads.
"""
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlencode
import argparse
import io
import json
import os
import secrets
import time

import httpx
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
OUT=Path(os.getenv('AICADEMIYA_TEST_ARTIFACT_DIR',str(ROOT/'artifacts/stage2')))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--base-url',required=True)
    parser.add_argument('--confirm-test-database',action='store_true')
    args=parser.parse_args()
    if not args.confirm_test_database:raise SystemExit('Use an isolated synthetic QA database.')
    events=[];checks=[];start=time.monotonic()
    nonce=secrets.token_hex(4)
    clients=[httpx.Client(base_url=args.base_url,trust_env=False,timeout=180) for _ in range(5)]
    admin,teacher,student,second,anonymous=clients
    def call(client,method,path,body=None,expected=200,**kwargs):
        response=client.request(method,path,json=body,**kwargs)
        statuses={expected} if isinstance(expected,int) else set(expected)
        events.append({'method':method,'path':path,'status':response.status_code,'expected':sorted(statuses)})
        assert response.status_code in statuses,(path,response.status_code,response.text[:500])
        if response.headers.get('content-type','').startswith('application/json'):return response.json()
        return response
    def check(name,condition=True):
        assert condition,name
        checks.append(name);print('PASS',name,flush=True)
    def versioned(path,vid):return path+'?'+urlencode({'version_id':vid})
    result={'status':'failed','synthetic_only':True,'real_model_requests':0}
    try:
        boot=call(anonymous,'GET','/api/bootstrap')
        check('AIcademiya branding and explicit demo',boot['app_name']=='AIcademiya' and boot['ai']['mode']=='demo')
        au=call(admin,'POST','/api/auth/demo',{'role':'admin'})['user']
        tu=call(teacher,'POST','/api/auth/demo',{'role':'teacher'})['user']
        call(teacher,'GET','/api/teacher/settings',expected=403)
        call(teacher,'PUT','/api/teacher/settings',{'mode':'demo','base_url':'','model':'forbidden'},expected=403)
        call(teacher,'POST','/api/teacher/settings/test',{},expected=403)
        call(admin,'GET','/api/teacher/settings')
        check('Only admin can inspect/change/test platform model settings',au['role']=='admin' and tu['role']=='teacher')
        course=call(teacher,'POST','/api/courses',{'code':'STAGE2-'+nonce,'title':'Synthetic stage-two course','description':'All data in this course is QA-only.'})
        other_course=call(admin,'POST','/api/courses',{'code':'OTHER-'+nonce,'title':'Other responsible teacher course','description':'Cross-course isolation fixture'})
        cid=course['id'];ocid=other_course['id']
        call(teacher,'GET',f'/api/teacher/courses/{ocid}/students',expected=403)
        # Real uploaded roster; passwords are synthetic and not printed/returned.
        email1=f'roster-one-{nonce}@example.invalid';email2=f'roster-two-{nonce}@example.invalid'
        temp='SyntheticRoster9';newpass='SyntheticChanged9'
        csv=f'name,email,temporary_password\nRoster One,{email1},{temp}\nRoster Two,{email2},{temp}\n'.encode()
        preview=call(teacher,'POST',f'/api/teacher/courses/{cid}/students/import/preview',files={'file':('students.csv',csv,'text/csv')})
        check('Roster preview validates and never echoes passwords',preview['valid_count']==2 and preview['error_count']==0 and temp not in json.dumps(preview))
        committed=call(teacher,'POST',f'/api/teacher/courses/{cid}/students/import/{preview["import_id"]}/commit',{})
        check('Roster commits the new students',committed['created']==2)
        roster=call(teacher,'GET',f'/api/teacher/courses/{cid}/students')
        user1=next(r['user'] for r in roster if r['user']['email']==email1)
        user2=next(r['user'] for r in roster if r['user']['email']==email2)
        login=call(student,'POST','/api/auth/login',{'email':email1,'password':temp})
        check('Imported account requires first password change',login['user']['must_change_password'])
        call(student,'GET','/api/courses',expected=403)
        call(student,'POST','/api/auth/change-password',{'current_password':temp,'new_password':newpass})
        call(student,'PATCH','/api/auth/profile',{'name':'Changed display name'})
        me=call(student,'GET','/api/auth/me')['user']
        check('Profile changes display name but not account identifier',me['name']=='Changed display name' and me['email']==email1)
        call(student,'PATCH','/api/auth/profile',{'name':'X','role':'admin'},expected=422)
        call(student,'POST','/api/auth/change-password',{'current_password':'WrongCurrent9','new_password':'AnotherNew9'},expected=400)
        call(student,'GET','/api/teacher/settings',expected=403)
        call(student,'GET','/api/admin/teachers',expected=403)
        check('Profile/current-password and student privilege gates remain enforced')
        # Enrolling an existing student must not reset their name or password.
        existing_csv=f'name,email,temporary_password\nDo Not Replace,{email1},DoNotSetThis9\n'.encode()
        preview2=call(admin,'POST',f'/api/teacher/courses/{ocid}/students/import/preview',files={'file':('existing.csv',existing_csv,'text/csv')})
        call(admin,'POST',f'/api/teacher/courses/{ocid}/students/import/{preview2["import_id"]}/commit',{})
        me=call(student,'GET','/api/auth/me')['user']
        check('Existing shared account is enrolled without identity/password takeover',me['name']=='Changed display name')
        # New editable lesson, materials and question sources.
        lesson=call(teacher,'POST',f'/api/courses/{cid}/sessions',{'title':'Versioned lesson','number':1,'instructions':'Synthetic private instruction v1'})
        sid=lesson['id']
        alpha=call(teacher,'POST',f'/api/sessions/{sid}/materials',files={'file':('alpha.txt',b'ALPHA SOURCE: reliability concerns consistency.','text/plain')},data={'visibility':'student','kind':'reading'})
        beta=call(teacher,'POST',f'/api/sessions/{sid}/materials',files={'file':('beta.txt',b'BETA SOURCE: a different topic.','text/plain')},data={'visibility':'student','kind':'reading'})
        def q(qid,prompt,**extra):
            return {'id':qid,'group_id':'group-core','ordinal':1 if qid=='q-one' else 2,'type':'open_response','prompt':prompt,
                    'options':[],'answer_slots':1,'images':[],'source_label':'Synthetic authoring source','source_material_ids':[], 'source_links':[],**extra}
        q1=q('q-one','Original question one',source_material_ids=[alpha['id']])
        q2=q('q-two','Original question two',source_links=['https://example.com/'])
        detail=call(teacher,'PATCH',f'/api/sessions/{sid}',{'status':'published','groups':[{'id':'group-core','title':'Core tasks','kind':'core'}],
             'questions':[q1,q2],'reference_answers':{'q-one':'Original reference one','q-two':'Original reference two'}})
        links=[m for m in detail['materials'] if m['kind']=='link']
        check('Question external source creates a Material Link',len(links)==1 and links[0]['url']=='https://example.com/')
        # Second question shares the Link; clearing just one must preserve it.
        q1['source_links']=['https://example.com/']
        call(teacher,'PATCH',f'/api/sessions/{sid}',{'questions':[q1,q2]})
        q2['source_links']=[]
        detail=call(teacher,'PATCH',f'/api/sessions/{sid}',{'questions':[q1,q2]})
        check('Shared Link survives clearing one question',any(m['id']==links[0]['id'] for m in detail['materials']))
        preview_file=call(teacher,'GET',alpha['preview_url'])
        check('Uploaded text has a real usable preview','ALPHA SOURCE' in preview_file.text)
        v1=call(teacher,'POST',f'/api/teacher/courses/{cid}/versions/publish',{'label':'Initial edition'})
        vid1=v1['id']
        w1=call(student,'GET',f'/api/sessions/{sid}/workspace')
        check('Initial publication creates a student-visible version',w1['version']['id']==vid1 and w1['version']['number']==1)
        call(student,'POST',versioned(f'/api/sessions/{sid}/answers/q-one',vid1),{'answer':'Answer under edition one'})
        call(student,'POST',versioned(f'/api/sessions/{sid}/answers/q-two',vid1),{'answer':'Second answer under edition one'})
        call(student,'POST',versioned(f'/api/sessions/{sid}/references/q-one',vid1),{})
        call(student,'POST',versioned(f'/api/sessions/{sid}/chat',vid1),{'message':'Synthetic follow-up for version one','question_id':'q-one','intent':'ask','web_search':False,'request_id':'stage2-ask-one'})
        call(student,'POST',versioned(f'/api/sessions/{sid}/submit',vid1),{})
        old_material_url=next(m for m in w1['session']['materials'] if m['id']==alpha['id'])['download_url']
        # Draft editing is not publication; old work must not be reinterpreted.
        changed={**q1,'prompt':'Changed question one','source_material_ids':[beta['id']]}
        call(teacher,'PATCH',f'/api/sessions/{sid}',{'questions':[changed,q2],'reference_answers':{'q-one':'Changed reference one','q-two':'Original reference two'}})
        still=call(student,'GET',versioned(f'/api/sessions/{sid}/workspace',vid1))
        check('Unpublished edits stay away from student version',still['session']['questions'][0]['prompt']=='Original question one')
        diff=call(teacher,'GET',f'/api/teacher/courses/{cid}/versions/{vid1}/diff')
        check('Version comparison shows actual content changes',bool(diff['changes']) and 'Changed question one' in json.dumps(diff))
        v2=call(teacher,'POST',f'/api/teacher/courses/{cid}/versions/publish',{'label':'Second edition'})
        vid2=v2['id']
        latest=call(student,'GET',f'/api/sessions/{sid}/workspace')
        check('New edition has its own learning run',latest['version']['id']==vid2 and not latest['attempts'] and not latest['messages'])
        old=call(student,'GET',versioned(f'/api/sessions/{sid}/workspace',vid1))
        check('Historical version keeps original question and answer',old['read_only'] and old['session']['questions'][0]['prompt']=='Original question one' and old['attempts']['q-one']['answer']=='Answer under edition one')
        call(student,'POST',versioned(f'/api/sessions/{sid}/answers/q-one',vid1),{'answer':'Must not change history'},expected=409)
        old_ref=call(student,'GET',versioned(f'/api/sessions/{sid}/references/q-one',vid1))
        check('Previously revealed reference remains original',old_ref['reference']=='Original reference one')
        reviewed=call(teacher,'POST',versioned(f'/api/teacher/sessions/{sid}/students/{user1["id"]}/review',vid1),{'status':'confirmed','note':'Reviewed original edition'})
        check('Teacher can review original submission after later publication',reviewed['submission']['status']=='confirmed')
        check('Historical material download survives later edits',call(student,'GET',old_material_url).content==b'ALPHA SOURCE: reliability concerns consistency.')
        hist=call(student,'GET','/api/history')
        check('History lists the original learning run',any(r['session']['id']==sid and r['version']['id']==vid1 for r in hist))
        # Operational availability is independent of curriculum restore.
        call(teacher,'PATCH',f'/api/sessions/{sid}',{'active':False,'starts_at':None,'ends_at':None})
        inactive=call(student,'GET',versioned(f'/api/sessions/{sid}/workspace',vid2))
        check('Deactive is readonly while history remains accessible',inactive['read_only'] and not inactive['session']['effective_active'])
        call(student,'POST',versioned(f'/api/sessions/{sid}/chat',vid2),{'message':'Should be blocked','question_id':'q-one','intent':'ask'},expected=409)
        v3=call(teacher,'POST',f'/api/teacher/courses/{cid}/versions/{vid1}/restore',{'label':'Restore initial content'})
        current=call(teacher,'GET',f'/api/sessions/{sid}')
        check('Restore creates a new edition and preserves operational closure',v3['id'] not in {vid1,vid2} and current['questions'][0]['prompt']=='Original question one' and not current['effective_active'])
        now=datetime.now(timezone.utc)
        scheduled=call(teacher,'PATCH',f'/api/sessions/{sid}',{'active':False,'starts_at':(now-timedelta(hours=1)).isoformat(),'ends_at':(now+timedelta(hours=1)).isoformat()})
        check('Inside scheduled window overrides manual inactive',scheduled['effective_active'])
        scheduled=call(teacher,'PATCH',f'/api/sessions/{sid}',{'active':True,'starts_at':(now+timedelta(hours=1)).isoformat(),'ends_at':(now+timedelta(hours=2)).isoformat()})
        check('Future schedule overrides manual active',not scheduled['effective_active'])
        call(teacher,'PATCH',f'/api/sessions/{sid}',{'active':True,'starts_at':None,'ends_at':None})
        vid3=v3['id']
        # Scoped course archive does not disable a shared platform account.
        call(teacher,'POST',f'/api/teacher/courses/{ocid}/students/archive',{'user_ids':[user1['id']]},expected=403)
        call(teacher,'POST',f'/api/teacher/courses/{cid}/students/archive',{'user_ids':[user1['id']]})
        archived=call(teacher,'GET',f'/api/teacher/courses/{cid}/students?archived=true')
        check('Course archive has recoverable30-day metadata',any(r['user']['id']==user1['id'] and r['archived_at'] and r['purge_at'] for r in archived))
        call(student,'POST',versioned(f'/api/sessions/{sid}/answers/q-one',vid3),{'answer':'Blocked after course removal'},expected={403,409})
        check('Shared student still accesses other enrolled course',any(c['id']==ocid for c in call(student,'GET','/api/courses')))
        call(teacher,'POST',f'/api/teacher/courses/{cid}/students/restore',{'user_ids':[user1['id']]})
        restored=call(student,'GET',versioned(f'/api/sessions/{sid}/workspace',vid1))
        check('Course restore keeps original historical work',restored['attempts']['q-one']['answer']=='Answer under edition one')
        call(admin,'POST','/api/admin/students/archive',{'user_ids':[user1['id']]})
        call(student,'GET','/api/auth/me',expected={401,403})
        call(student,'POST','/api/auth/login',{'email':email1,'password':newpass},expected={401,403})
        call(admin,'POST','/api/admin/students/restore',{'user_ids':[user1['id']]})
        call(student,'POST','/api/auth/login',{'email':email1,'password':newpass})
        check('Global archive revokes access; restore keeps credentials and history',call(student,'GET',versioned(f'/api/sessions/{sid}/workspace',vid1))['attempts']['q-one']['answer']=='Answer under edition one')
        # Real DOCX extraction -> reviewed editable draft -> commit, not a hard-coded sample.
        document=Document();document.add_heading('Synthetic imported worksheet',0)
        document.add_paragraph('Part A: Short answers')
        document.add_paragraph('1. Explain why consistent marking matters.')
        document.add_paragraph('2. Give a classroom example of authentic assessment.')
        stream=io.BytesIO();document.save(stream)
        imported=call(teacher,'POST',f'/api/sessions/{sid}/worksheet-import/preview',files={'file':('new-worksheet.docx',stream.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'use_ai':'false'})
        draft=imported['draft']
        check('Uploaded DOCX produces actual editable question drafts',len(draft['questions'])>=2 and 'consistent marking' in json.dumps(draft))
        body={k:draft[k] for k in ('groups','questions','reference_answers')};body['mode']='append'
        newdraft=call(teacher,'POST',f'/api/sessions/{sid}/worksheet-import/{imported["import_id"]}/commit',body)
        check('Reviewed worksheet commit modifies only draft',len(newdraft['questions'])>2 and len(call(student,'GET',f'/api/sessions/{sid}/workspace')['session']['questions'])==2)
        # Promotion is last because this fixture teacher legitimately becomes admin.
        teachers=call(admin,'GET','/api/admin/teachers')
        check('Promotion candidates include existing teacher',any(u['id']==tu['id'] for u in teachers))
        promoted=call(admin,'POST',f'/api/admin/teachers/{tu["id"]}/promote',{})['user']
        check('Only specifically selected teacher is promoted',promoted['role']=='admin' and promoted['id']==tu['id'])
        call(teacher,'GET','/api/auth/me',expected=401)
        signed_in=call(teacher,'POST','/api/auth/demo',{'role':'teacher'})['user']
        check('Promotion revokes old sessions and requires fresh authentication',signed_in['id']==tu['id'] and signed_in['role']=='admin')
        call(teacher,'GET','/api/teacher/settings')
        check('Promoted teacher can configure models after signing in again')
        result.update(status='passed',checks=checks,requests=events,request_count=len(events),seconds=round(time.monotonic()-start,2),
                      course_id=cid,session_id=sid,versions=[vid1,vid2,vid3],student_id=user1['id'])
    except Exception as exc:
        result.update(error=str(exc),checks=checks,requests=events,request_count=len(events));raise
    finally:
        for client in clients:client.close()
        OUT.mkdir(parents=True,exist_ok=True)
        (OUT/'api-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        print(json.dumps({k:v for k,v in result.items() if k not in {'checks','requests'}},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
