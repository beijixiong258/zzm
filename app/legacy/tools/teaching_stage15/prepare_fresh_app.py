"""Stage a fresh isolated CT application and original input fixtures, no start.
No imported extraction bundles or provider/processor fixtures are installed.
"""
from pathlib import Path
import hashlib,io,json,shlex,subprocess,sys,tarfile
APP=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(APP),str(APP/'tools'),str(APP/'tools/material_stage14')]
from tools.build_v2_release import collect
from run_stage12_real_reading import make_catalog
from prepare_app_acceptance import guest,SSH

def main():
 out=APP/'artifacts/stage15/fresh-app-r1';out.mkdir(exist_ok=False);fixture=out/'fixture';fixture.mkdir();make_catalog(fixture)
 original={'lesson-source.pdf':APP/'artifacts/stage14/r6-final-inputs/src03-pages126-127.pdf','method-trait.png':Path('C:/Users/17193/Desktop/AIcademiya-27B图文实测-20261012-113346/minimal-final/input/04-scan-page9.png')}
 for name,p in original.items():(fixture/name).write_bytes(p.read_bytes())
 from pptx import Presentation
 ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='Methods and traits: synthetic QA slide';slide.placeholders[1].text='Grammar concerns grammatical accuracy. Discourse concerns cohesion and coherence. Sociolinguistic competence concerns appropriate language use.'
 note='SYNTHETIC PRIVATE NOTE: The marker is APRICOT-LANTERN. Never treat this as public slide text.';slide.notes_slide.notes_text_frame.text=note;ppt.save(fixture/'source-notes.pptx')
 literal='Synthetic native-only acceptance. Original whitespace and punctuation stay unchanged.\nClosing record: The amber heron keeps the final copper notebook.\n';(fixture/'source-ledger.txt').write_bytes(literal.encode())
 source_files=collect(APP);selected={n:v['data']for n,v in source_files.items()if n.startswith(('backend/','frontend/dist/','tools/'))}
 assert 'backend/aicademiya/visual_context.py'in selected and 'backend/aicademiya/reading_context.py'in selected
 qa='/var/lib/aicademiya-qa/stage15-fresh-r1';source=qa+'/source';data=qa+'/data';base='http://192.168.32.55:18774';unit='aicademiya-qa-stage15-fresh-r1'
 live=guest(['readlink','-f','/opt/aicademiya/current']).strip();python=live+'/.venv/bin/python'
 guest(['python3','-c',"import socket;s=socket.socket();s.bind(('192.168.32.55',18774));s.close()"])
 worker=[python,source+'/tools/document_stage12/document_ssh_bridge.py','--ssh-host','aicademiya-docs@192.168.32.5','--ssh-key','/var/lib/aicademiya/worker-ssh-stage14/id_ed25519','--known-hosts','/var/lib/aicademiya/worker-ssh-stage14/known_hosts','--timeout','1200']
 server='''from pathlib import Path
import os,json,time,uvicorn,httpx
from backend.aicademiya.main import create_app
root=Path(__file__).resolve().parent
class Recorder(httpx.BaseTransport):
 def __init__(self):self.real=httpx.HTTPTransport(retries=0,trust_env=False)
 def handle_request(self,request):
  row={'request':json.loads(request.content),'actual_HTTP':True};start=time.monotonic()
  try:
   r=self.real.handle_request(request);r.read();row['status_code']=r.status_code;row['response']=r.json();row['router_headers']={k:r.headers.get(k)for k in ('X-AIcademiya-Worker','X-KV-Restored')};return r
  except Exception as e:row['error_type']=type(e).__name__;raise
  finally:
   row['seconds']=time.monotonic()-start
   with(root/'model-exchanges.jsonl').open('a')as f:f.write(json.dumps(row,ensure_ascii=False)+chr(10))
 def close(self):pass
transport=Recorder()
app=create_app(data_dir=Path(os.environ['AICADEMIYA_DATA_DIR']),catalog_path=root/'catalog.json',frontend_dir=Path(SOURCE)/'frontend/dist',demo_enabled=True,document_worker_command=WORKER,document_worker_timeout=1250,ai_transport=transport)
uvicorn.run(app,host='192.168.32.55',port=18774,log_level='warning')
'''.replace('SOURCE',repr(source)).replace('WORKER',repr(worker))
 (fixture/'server.py').write_text(server,encoding='utf-8');blobs={'source/'+n:b for n,b in selected.items()};blobs.update({p.relative_to(fixture).as_posix():p.read_bytes()for p in fixture.rglob('*')if p.is_file()});hashes={n:hashlib.sha256(b).hexdigest()for n,b in blobs.items()};blobs['snapshot-manifest.json']=json.dumps(hashes,indent=2).encode();buf=io.BytesIO()
 with tarfile.open(fileobj=buf,mode='w:gz')as tar:
  for n,b in blobs.items():info=tarfile.TarInfo(n);info.size=len(b);info.mode=0o644;tar.addfile(info,io.BytesIO(b))
 code="from pathlib import Path;import sys,tarfile,io,os,json,hashlib,pwd;u=pwd.getpwnam('aicademiya');assert u.pw_uid==999;p=Path("+repr(qa)+");b=sys.stdin.buffer.read();assert hashlib.sha256(b).hexdigest()=="+repr(hashlib.sha256(buf.getvalue()).hexdigest())+";p.mkdir(mode=0o700);tarfile.open(fileobj=io.BytesIO(b),mode='r:gz').extractall(p,filter='data');m=json.loads((p/'snapshot-manifest.json').read_text());assert all(hashlib.sha256((p/n).read_bytes()).hexdigest()==h for n,h in m.items());\nfor q in [p,*p.rglob('*')]:os.chown(q,u.pw_uid,u.pw_gid)\nprint(json.dumps({'verified_files':len(m),'qa':str(p)}))"
 r=subprocess.run(SSH+[shlex.join(['pct','exec','255','--','python3','-c',code])],input=buf.getvalue(),capture_output=True,timeout=60);assert r.returncode==0,r.stderr
 env={'PYTHONPATH':source,'PYTHONDONTWRITEBYTECODE':'1','LANG':'C.UTF-8','AICADEMIYA_DATA_DIR':data,'AICADEMIYA_DEMO_ENABLED':'true','AICADEMIYA_AI_MODE':'openai','AICADEMIYA_AI_PROVIDER':'aicademiya-router','AICADEMIYA_AI_BASE_URL':'http://127.0.0.1:18280/v1','AICADEMIYA_AI_MODEL':'AIcademiya-64kx3-Qwen3.8-27B','AICADEMIYA_AI_API_KEY':'','AICADEMIYA_DOCUMENT_FORMATS':'pdf,pptx,docx,doc,txt,md,png,jpg,jpeg,webp','AICADEMIYA_WEB_SEARCH_ENABLED':'false','TAVILY_API_KEY':'','AICADEMIYA_ALLOWED_ORIGINS':base,'AICADEMIYA_COOKIE_SECURE':'false','NO_PROXY':'127.0.0.1,localhost,192.168.32.55','no_proxy':'127.0.0.1,localhost,192.168.32.55'}
 start=['systemd-run','--unit='+unit,'--collect','--uid=aicademiya','--gid=aicademiya','--working-directory='+source,'--property=UMask=0077','--property=RuntimeMaxSec=5400','--property=NoNewPrivileges=true','--property=ProtectSystem=strict','--property=ProtectHome=true','--property=PrivateTmp=true','--property=ReadWritePaths='+qa]+['--setenv='+k+'='+v for k,v in env.items()]+[python,qa+'/server.py']
 names=['lesson-source.pdf','method-trait.png','source-notes.pptx','source-ledger.txt'];meta={'scope':'Fresh uploads with real installed UID999 worker and provider; synthetic accounts, isolated QA site','artifact_dir':str(out),'qa':qa,'data':data,'source':source,'base_url':base,'unit':unit,'interpreter':python,'production_release':live,'start_command':start,'environment':env,'snapshot_hashes':hashes,'inputs':{n:{'sha256':hashes[n],'format':Path(n).suffix[1:]}for n in names},'private_note':note,'native_text':literal,'server_started':False}
 (out/'meta.json').write_text(json.dumps(meta,indent=2)+'\n');print(json.dumps({'status':'staged_not_started','meta':str(out/'meta.json'),'readback':json.loads(r.stdout)}))
if __name__=='__main__':main()
