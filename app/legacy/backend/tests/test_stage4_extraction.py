"""Deterministic extraction contract tests; no model, DB or shared server.

Real installed CPU OCR is exercised separately by artifacts/stage4/ocr/
run_cpu_evidence.py. These tests replace recognition only, not document parsing.
"""
from dataclasses import FrozenInstanceError
import io
import json
from pathlib import Path
import zipfile

import pytest
import pymupdf
from PIL import Image

from backend.aicademiya import material_extraction as me
from backend.aicademiya import materials
from backend.aicademiya import worksheet_import

REPORT_KEYS = {"schema_version", "method", "units_total", "units_processed", "ocr_units",
               "images_detected", "chars_extracted", "excerpt_chars", "truncated",
               "visual_review_required", "failed_units", "notes"}


def raster(color="white", size=(260, 100)):
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def pdf(path, *, native="", image=True, pages=1, drawings=False):
    with pymupdf.open() as document:
        for index in range(pages):
            page = document.new_page(width=300, height=400)
            if native:
                page.insert_text((20, 35), native + f" page {index + 1}")
            if image:
                page.insert_image(pymupdf.Rect(20, 70, 280, 170), stream=raster())
            if drawings:
                page.draw_line((20, 190), (220, 260))
        document.save(path)
    return path


def assert_report(result):
    assert set(result) == {"status", "excerpt", "warning", "extraction_report"}
    report = result["extraction_report"]
    assert set(report) == REPORT_KEYS
    assert report["schema_version"] == 1
    assert 0 <= report["units_processed"] <= report["units_total"]
    assert report["excerpt_chars"] == len(result["excerpt"]) <= 50000
    assert report["chars_extracted"] >= report["excerpt_chars"]
    assert report["truncated"] == (report["chars_extracted"] > report["excerpt_chars"])
    assert isinstance(report["failed_units"], list) and isinstance(report["notes"], list)
    json.dumps(result)
    return report


@pytest.fixture
def fake_ocr(monkeypatch):
    calls = []
    def recognize(png):
        assert png.startswith(b"\x89PNG")
        calls.append(png)
        return ["Synthetic raster body"]
    monkeypatch.setattr(me, "_ocr_lines", recognize)
    return calls


def test_report_is_frozen():
    report = me.ExtractionReport(1, "native", 1, 1, 0, 0, 3, 3, False, False, (), ())
    with pytest.raises(FrozenInstanceError):
        report.method = "changed"
    assert set(report.to_dict()) == REPORT_KEYS


def test_material_ocr_reuses_constructor_not_mutable_worksheet_instance(monkeypatch):
    from functools import lru_cache
    from types import SimpleNamespace
    created = []
    class Engine:
        def __call__(self, png):
            return SimpleNamespace(txts=["Independent synthetic recognition", " "])
    @lru_cache(maxsize=1)
    def constructor():
        engine = Engine()
        created.append(engine)
        return engine
    monkeypatch.setattr(worksheet_import, "_local_ocr_engine", constructor)
    worksheet_engine = constructor()
    me._material_ocr_engine.cache_clear()
    try:
        material_engine = me._material_ocr_engine()
        assert material_engine is not worksheet_engine
        assert me._material_ocr_engine() is material_engine and len(created) == 2
        assert me._ocr_lines(b"synthetic input") == ["Independent synthetic recognition"]
    finally:
        me._material_ocr_engine.cache_clear()


def test_text_preserves_legacy_tuple_and_utf8_bom(tmp_path):
    path = tmp_path / "native.txt"
    path.write_bytes(b"\xef\xbb\xbf  Synthetic native text.\n")
    result = materials.extract_document(path)
    assert materials.extract(path) == (result["status"], result["excerpt"], result["warning"])
    assert result["status"] == "ready" and result["warning"] is None
    assert result["excerpt"] == "  Synthetic native text.\n"
    report = assert_report(result)
    assert report["units_total"] == report["units_processed"] == 1
    assert report["method"] == "native" and not report["visual_review_required"]


@pytest.mark.parametrize("suffix", [".png", ".jpg", ".jpeg", ".webp"])
def test_standalone_images_decode_to_cpu_ocr_png(tmp_path, suffix, fake_ocr):
    path = tmp_path / ("synthetic" + suffix)
    Image.new("RGB", (260, 100), "white").save(path)
    before = path.read_bytes()
    result = materials.extract_document(path)
    report = assert_report(result)
    assert result["status"] == "needs_review" and "Synthetic raster body" in result["excerpt"]
    assert report["method"] == "cpu_ocr" and report["visual_review_required"]
    assert report["units_total"] == report["units_processed"] == report["images_detected"] == report["ocr_units"] == 1
    assert before == path.read_bytes() and suffix in materials.SUPPORTED
    assert suffix not in worksheet_import.SUPPORTED


def test_pdf_native_header_never_suppresses_image_body_ocr(tmp_path, fake_ocr):
    path = pdf(tmp_path / "mixed.pdf", native="Long native header with substantial teaching text")
    before = path.read_bytes()
    result = materials.extract_document(path)
    report = assert_report(result)
    assert "Long native header" in result["excerpt"] and "Synthetic raster body" in result["excerpt"]
    assert report["method"] == "native+cpu_ocr"
    assert report["units_total"] == report["units_processed"] == 2
    assert report["images_detected"] == report["ocr_units"] == len(fake_ocr) == 1
    assert before == path.read_bytes()


def test_native_only_pdf_no_ocr_and_vector_review(tmp_path, fake_ocr):
    native = materials.extract_document(pdf(tmp_path / "native.pdf", native="Native table label", image=False))
    assert native["status"] == "ready" and not fake_ocr
    assert_report(native)
    vector = materials.extract_document(pdf(tmp_path / "vector.pdf", native="Axes label", image=False, drawings=True))
    assert vector["status"] == "needs_review" and not fake_ocr
    assert vector["extraction_report"]["visual_review_required"]


def test_rotated_pdf_image_region_is_rendered_in_visible_page_coordinates(tmp_path, fake_ocr):
    path = tmp_path / "rotated.pdf"
    with pymupdf.open() as document:
        page = document.new_page(width=300, height=400)
        page.insert_text((20, 35), "Native rotated header")
        page.insert_image(pymupdf.Rect(20, 70, 280, 170), stream=raster("black"))
        page.set_rotation(90)
        document.save(path)
    result = materials.extract_document(path)
    report = assert_report(result)
    assert report["units_total"] == report["units_processed"] == 2
    rendered = Image.open(io.BytesIO(fake_ocr[0]))
    assert rendered.height > rendered.width
    assert rendered.convert("RGB").getpixel((rendered.width // 2, rendered.height // 2)) == (0, 0, 0)


def test_empty_pdf_render_ocr_and_empty_recognition_are_truthful(tmp_path, monkeypatch):
    monkeypatch.setattr(me, "_ocr_lines", lambda png: [])
    result = materials.extract_document(pdf(tmp_path / "blank.pdf", image=False))
    report = assert_report(result)
    assert result["status"] == "needs_review" and result["excerpt"] == ""
    assert report["units_total"] == report["units_processed"] == 2
    assert report["images_detected"] == 0 and report["ocr_units"] == 1
    assert "OCR" in result["warning"] and any("found no text" in note for note in report["notes"])


def test_duplicate_native_raster_text_omitted_without_losing_unique_lines(tmp_path, monkeypatch):
    monkeypatch.setattr(me, "_ocr_lines", lambda png: ["SYNTHETIC NATIVE HEADER, PAGE 1", "Raster-only body text", "Raster-only body text"])
    result = materials.extract_document(pdf(tmp_path / "overlap.pdf", native="Synthetic native header"))
    assert result["excerpt"].lower().count("synthetic native header") == 1
    assert result["excerpt"].count("Raster-only body text") == 1
    assert_report(result)


def test_repeated_embedded_images_reuse_recognition(tmp_path, fake_ocr):
    from docx import Document
    from docx.shared import Inches
    document = Document()
    document.add_paragraph("Native prose")
    document.add_picture(io.BytesIO(raster()), width=Inches(2))
    document.add_picture(io.BytesIO(raster()), width=Inches(2))
    path = tmp_path / "repeat.docx"
    document.save(path)
    result = materials.extract_document(path)
    report = assert_report(result)
    assert report["images_detected"] == 2 and report["ocr_units"] == len(fake_ocr) == 1
    assert report["units_total"] == report["units_processed"] == 3
    assert result["excerpt"].count("Synthetic raster body") == 1
    assert any("reused" in note for note in report["notes"])


def test_docx_native_nested_tables_headers_and_embedded_images(tmp_path, fake_ocr):
    from docx import Document
    from docx.shared import Inches
    document = Document()
    document.add_paragraph("Native document body")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Native cell alpha"
    table.cell(0, 1).text = "Native cell beta"
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested table value"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(raster()), width=Inches(1))
    document.sections[0].header.paragraphs[0].text = "Synthetic header text"
    document.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(raster("yellow")), width=Inches(1))
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer text"
    path = tmp_path / "embedded.docx"
    document.save(path)
    result = materials.extract_document(path)
    for text in ("Native document body", "Native cell alpha", "Native cell beta", "Nested table value", "Synthetic header text", "Synthetic footer text", "Synthetic raster body"):
        assert text in result["excerpt"]
    report = assert_report(result)
    assert report["units_total"] == report["units_processed"] == 5
    assert report["images_detected"] == report["ocr_units"] == 2
    assert report["method"] == "native+cpu_ocr"


def test_docx_external_image_is_skipped_without_fetching(tmp_path, fake_ocr):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Inches
    document = Document()
    document.add_paragraph("Native body with linked image")
    document.add_picture(io.BytesIO(raster()), width=Inches(1))
    blip = document.element.body.xpath(".//a:blip")[0]
    del blip.attrib[qn("r:embed")]
    blip.set(qn("r:link"), document.part.relate_to("https://invalid.example/private.png", RT.IMAGE, is_external=True))
    path = tmp_path / "external.docx"
    document.save(path)
    result = materials.extract_document(path)
    report = assert_report(result)
    assert not fake_ocr and result["status"] == "needs_review"
    assert report["units_total"] == 2 and report["units_processed"] == 1
    assert report["images_detected"] == 1
    assert any("external resources are never fetched" in note for note in report["notes"])
    assert "invalid.example" not in json.dumps(result)


def presentation(path, *, grouped=False, image=True, table=True):
    from pptx import Presentation
    from pptx.util import Inches
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    shapes = slide.shapes
    if grouped:
        shapes = shapes.add_group_shape().shapes.add_group_shape().shapes
        shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Nested grouped slide text"
    if image:
        shapes.add_picture(io.BytesIO(raster()), Inches(1), Inches(2), width=Inches(2))
    if table:
        frame = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1))
        if grouped:
            # python-pptx has no GroupShapes.add_table; move the native table
            # element into the group to exercise an actual grouped table.
            shapes._spTree.append(frame._element)
        cells = frame.table
        cells.cell(0, 0).text = "Only native table alpha"
        cells.cell(0, 1).text = "Only native table beta"
    document.save(path)
    return path


def test_image_only_pptx_is_not_discarded(tmp_path, fake_ocr):
    result = materials.extract_document(presentation(tmp_path / "image.pptx", table=False))
    report = assert_report(result)
    assert result["status"] == "needs_review" and "Synthetic raster body" in result["excerpt"]
    assert report["units_total"] == report["units_processed"] == 2
    assert report["method"] == "cpu_ocr"


def test_table_only_pptx_is_native_text(tmp_path, fake_ocr):
    result = materials.extract_document(presentation(tmp_path / "table.pptx", image=False))
    assert result["status"] == "ready" and "Only native table alpha | Only native table beta" in result["excerpt"]
    assert not fake_ocr and assert_report(result)["method"] == "native"


def test_nested_grouped_slide_text_tables_and_images(tmp_path, fake_ocr):
    path = presentation(tmp_path / "grouped.pptx", grouped=True)
    result = materials.extract_document(path)
    assert all(text in result["excerpt"] for text in ("Nested grouped slide text", "Only native table alpha", "Only native table beta", "Synthetic raster body"))
    report = assert_report(result)
    assert report["units_total"] == report["units_processed"] == 2
    assert report["images_detected"] == report["ocr_units"] == 1
    assert report["visual_review_required"] and any("recursively" in note for note in report["notes"])


def test_slide_background_image_is_ocr_read_not_only_picture_shapes(tmp_path, fake_ocr):
    from pptx import Presentation
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    _, rid = slide.part.get_or_add_image_part(io.BytesIO(raster()))
    background = OxmlElement("p:bg")
    props = OxmlElement("p:bgPr")
    fill = OxmlElement("a:blipFill")
    blip = OxmlElement("a:blip")
    blip.set(qn("r:embed"), rid)
    fill.append(blip)
    stretch = OxmlElement("a:stretch")
    stretch.append(OxmlElement("a:fillRect"))
    fill.append(stretch)
    props.append(fill)
    background.append(props)
    slide._element.cSld.insert(0, background)
    path = tmp_path / "background.pptx"
    document.save(path)
    result = materials.extract_document(path)
    report = assert_report(result)
    assert "Synthetic raster body" in result["excerpt"]
    assert report["images_detected"] == report["ocr_units"] == 1
    assert report["units_total"] == report["units_processed"] == 2


def test_punctuation_ocr_is_not_silently_discarded():
    assert me._new_ocr_text(["++", "--", "++"], "Native words") == "++\n--"


def test_ocr_unavailable_preserves_native_and_reports_all_failures(tmp_path, monkeypatch):
    calls = []
    def unavailable(png):
        calls.append(1)
        raise ImportError("PRIVATE_PATH_OR_TEXT")
    monkeypatch.setattr(me, "_ocr_lines", unavailable)
    result = materials.extract_document(pdf(tmp_path / "missing.pdf", native="Native survives", pages=3))
    report = assert_report(result)
    assert result["status"] == "needs_review" and "Native survives" in result["excerpt"]
    assert report["units_total"] == 6 and report["units_processed"] == 3
    assert report["ocr_units"] == 0 and len(report["failed_units"]) == 3 and len(calls) == 1
    assert "PRIVATE_PATH_OR_TEXT" not in json.dumps(result)


def test_ocr_failed_image_is_not_a_successful_processed_unit(tmp_path, monkeypatch):
    def fail(png):
        raise RuntimeError("private filesystem and document text")
    monkeypatch.setattr(me, "_ocr_lines", fail)
    path = tmp_path / "failed.png"
    path.write_bytes(raster())
    result = materials.extract_document(path)
    report = assert_report(result)
    assert result["status"] == "needs_review" and result["excerpt"] == ""
    assert report["units_total"] == 1 and report["units_processed"] == report["ocr_units"] == 0
    assert report["method"] == "cpu_ocr_failed"
    assert len(report["failed_units"]) == 1
    assert "private filesystem" not in json.dumps(result)


def test_limits_report_skipped_pages_and_ocr_images(tmp_path, monkeypatch, fake_ocr):
    monkeypatch.setattr(me, "MAX_NATIVE_UNITS", 2)
    monkeypatch.setattr(me, "MAX_OCR_UNITS", 1)
    result = materials.extract_document(pdf(tmp_path / "limited.pdf", native="Retained native text", pages=4))
    report = assert_report(result)
    assert report["units_total"] == 6 and report["units_processed"] == 3
    assert report["images_detected"] == 2 and report["ocr_units"] == 1
    assert any("PDF pages 3-4" in note for note in report["notes"])
    assert any("PDF page 2 image 1" in note and "OCR attempt limit" in note for note in report["notes"])
    assert "Retained native text page 2" in result["excerpt"]
    assert "page 3" not in result["excerpt"]
    assert not report["failed_units"]


def test_large_image_skip_reports_pixels_not_empty_success(tmp_path, monkeypatch, fake_ocr):
    monkeypatch.setattr(me, "MAX_IMAGE_PIXELS", 100)
    path = tmp_path / "oversized.png"
    path.write_bytes(raster())
    result = materials.extract_document(path)
    report = assert_report(result)
    assert result["status"] == "needs_review" and not fake_ocr
    assert report["method"] == "not_extracted"
    assert report["units_processed"] == 0 and any("pixel decode limit" in note for note in report["notes"])


def test_downscaled_raster_limit_is_reported(tmp_path, monkeypatch, fake_ocr):
    monkeypatch.setattr(me, "MAX_RENDER_PIXELS", 400)
    path = tmp_path / "scaled.png"
    path.write_bytes(raster())
    report = assert_report(materials.extract_document(path))
    assert Image.open(io.BytesIO(fake_ocr[0])).width * Image.open(io.BytesIO(fake_ocr[0])).height <= 400
    assert any("downscaled" in note for note in report["notes"])


def test_truncation_keeps_processing_later_native_units(tmp_path, monkeypatch, fake_ocr):
    monkeypatch.setattr(me, "MAX_EXCERPT", 30)
    result = materials.extract_document(pdf(tmp_path / "long.pdf", native="Synthetic native text exceeding excerpt length", pages=3))
    report = assert_report(result)
    assert len(result["excerpt"]) == 30 and report["truncated"]
    assert report["units_total"] == report["units_processed"] == 6
    assert report["chars_extracted"] > 100 and report["images_detected"] == 3


def test_text_cap_and_markdown_unfetched_images(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("X" * 51000, encoding="utf-8")
    result = materials.extract_document(path)
    report = assert_report(result)
    assert len(result["excerpt"]) == 50000 and report["chars_extracted"] == 51000
    assert result["status"] == "needs_review"
    path = tmp_path / "external.md"
    path.write_text("![private figure](https://invalid.example/image)\nSynthetic text", encoding="utf-8")
    result = materials.extract_document(path)
    report = assert_report(result)
    assert result["status"] == "needs_review" and report["visual_review_required"]
    assert "not fetched" in result["warning"]


def test_animation_frames_have_truthful_coverage(tmp_path, fake_ocr):
    path = tmp_path / "animated.webp"
    Image.new("RGB", (260, 100), "white").save(path, save_all=True, append_images=[Image.new("RGB", (260, 100), "yellow")], duration=100, loop=0)
    result = materials.extract_document(path)
    report = assert_report(result)
    assert report["images_detected"] == report["units_total"] == report["units_processed"] == report["ocr_units"] == 2
    assert any("animation timing" in note for note in report["notes"])


@pytest.mark.parametrize("suffix", [".pdf", ".docx", ".pptx", ".png", ".txt", ".exe"])
def test_bad_input_never_exposes_parser_exception_or_private_path(tmp_path, suffix):
    path = tmp_path / ("PRIVATE_SOURCE_NAME" + suffix)
    path.write_bytes(b"\xff\xfe corrupt")
    result = materials.extract_document(path)
    report = assert_report(result)
    assert result["status"] == "error" and result["warning"]
    assert "PRIVATE_SOURCE_NAME" not in json.dumps(result)


def test_expanded_office_guard_and_input_limit(tmp_path, monkeypatch):
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "X" * 1000)
    monkeypatch.setattr(me, "MAX_EXPANDED_BYTES", 100)
    result = materials.extract_document(path)
    assert result["status"] == "error" and any("expanded Office" in note for note in result["extraction_report"]["notes"])
    path = tmp_path / "large.txt"
    path.write_bytes(b"X" * 101)
    monkeypatch.setattr(me, "MAX_INPUT_BYTES", 100)
    result = materials.extract_document(path)
    assert result["status"] == "error" and result["excerpt"] == ""


def test_password_pdf_never_prompts_or_uploads(tmp_path):
    path = tmp_path / "encrypted.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((50, 50), "Private synthetic encrypted text")
        document.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="synthetic-owner-only", user_pw="synthetic-reader-only")
    result = materials.extract_document(path)
    assert result["status"] == "error" and "encrypted PDF" in result["warning"]
    assert not result["excerpt"]
