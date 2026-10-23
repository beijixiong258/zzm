"""Bounded native/CPU-OCR material previews; never fetch linked resources.

Report schema v1 is fixed by STAGE4_CONTRACT.md. Units are native pages/slides/
DOCX stories/text files plus detected image occurrences (or image frames).
A blank/vector-only PDF page can add a page-render OCR unit. Successful blank
reads count as processed; OCR failures do not. Skips are explicit in notes.
"""
from dataclasses import dataclass
from functools import lru_cache
import hashlib
import io
from pathlib import Path
import re
from threading import Lock
import zipfile
from xml.etree import ElementTree as ET

MAX_EXCERPT = 50000
MAX_NATIVE_UNITS = 200
MAX_OCR_UNITS = 64
MAX_IMAGE_PIXELS = 16000000
MAX_RENDER_PIXELS = 4000000
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_INPUT_BYTES = 25 * 1024 * 1024
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".pptx", ".ppt", ".xls", ".xlsx"} | IMAGE_SUFFIXES

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
V = "{urn:schemas-microsoft-com:vml}"
_OCR_LOCK = Lock()


@dataclass(frozen=True)
class ExtractionReport:
    schema_version: int
    method: str
    units_total: int
    units_processed: int
    ocr_units: int
    images_detected: int
    chars_extracted: int
    excerpt_chars: int
    truncated: bool
    visual_review_required: bool
    failed_units: tuple[str, ...]
    notes: tuple[str, ...]

    def to_dict(self):
        # Explicit fields preserve the frozen public schema and JSON list types.
        return {"schema_version": self.schema_version, "method": self.method,
                "units_total": self.units_total, "units_processed": self.units_processed,
                "ocr_units": self.ocr_units, "images_detected": self.images_detected,
                "chars_extracted": self.chars_extracted, "excerpt_chars": self.excerpt_chars,
                "truncated": self.truncated, "visual_review_required": self.visual_review_required,
                "failed_units": list(self.failed_units), "notes": list(self.notes)}


@lru_cache(maxsize=1)
def _material_ocr_engine():
    # Reuse the unchanged installed-model constructor/preflight, not its
    # cached worksheet instance: RapidOCR mutates its detector preprocess_op
    # per input. A material-only instance avoids racing worksheet extraction.
    from .worksheet_import import _local_ocr_engine
    return _local_ocr_engine.__wrapped__()


def _ocr_lines(png):
    # Serial recognition bounds CPU work and protects RapidOCR's mutable
    # per-image state across simultaneous material upload/reprocess requests.
    with _OCR_LOCK:
        result = _material_ocr_engine()(png)
    return [str(text) for text in (result.txts or []) if str(text).strip()]


def _normalized(text):
    return re.sub(r"[^\w]", "", text.casefold())


def _new_ocr_text(lines, native):
    """Remove exact normalized native/OCR overlap, not approximate meaning."""
    native_lines = {_normalized(line) for line in native.splitlines() if line.strip()}
    native_flat = _normalized(native)
    result = []
    seen = set()
    for line in lines:
        line = str(line).strip()
        token = _normalized(line)
        if not line:
            continue
        # Math/punctuation-only recognitions are still source text, not empty.
        if not token:
            token = re.sub(r"\s", "", line.casefold())
            if token not in seen and not any(line == native_line.strip() for native_line in native.splitlines()):
                result.append(line)
            seen.add(token)
            continue
        if token in seen:
            continue
        seen.add(token)
        if token in native_lines or (len(token) >= 12 and token in native_flat):
            continue
        result.append(line)
    return "\n".join(result)


class _Skipped(ValueError):
    """Only controlled, teacher-safe reasons may enter report notes."""


class _Extraction:
    def __init__(self):
        self.total = self.processed = self.ocr_count = self.ocr_attempts = self.images = 0
        self.chars = self.text_chars = 0
        self.excerpt = ""
        self.native_text = False
        self.visual = False
        self.failed = []
        self.notes = ["Units include native pages/slides/DOCX stories/text files and detected image occurrences or frames; page-render OCR is an additional unit. Images on unvisited units are not counted.",
                      "Processed counts completed reads, including blank reads and cached OCR reuse; ocr_units counts completed CPU OCR calls, including calls finding no text.",
                      "chars_extracted counts deduplicated text plus provenance labels/separators before the 50,000-character excerpt cap. The preview is not a promise that every extracted character enters every model turn."]
        self.review_reasons = []
        self.ocr_cache = {}
        self.ocr_unavailable = False

    def note(self, text):
        if text not in self.notes:
            self.notes.append(text)

    def review(self, text, *, visual=False):
        self.note(text)
        if text not in self.review_reasons:
            self.review_reasons.append(text)
        self.visual |= visual

    def add(self, text, label, *, native=False):
        if not text.strip():
            return
        text = text.strip() if label else text
        self.text_chars += len(text)
        self.native_text |= native
        fragment = ("\n\n" if self.chars else "") + (f"[{label}]\n" if label else "") + text
        self.chars += len(fragment)
        remaining = max(0, MAX_EXCERPT - len(self.excerpt))
        self.excerpt += fragment[:remaining]

    def fail(self, label, reason):
        self.failed.append(f"{label}: {reason}")

    def skip(self, label, reason):
        self.review(f"Skipped: {label} ({reason}).")

    def native_limit(self, total, label):
        if total > MAX_NATIVE_UNITS:
            self.skip(f"{label} {MAX_NATIVE_UNITS + 1}-{total}", f"native-unit limit {MAX_NATIVE_UNITS}; images in these units were not inspected")

    def ocr(self, label, make_png, native=""):
        self.review("CPU OCR recognizes raster text, not diagram/chart relationships, color coding, handwriting accuracy or complete visual meaning; compare with the original.", visual=True)
        if self.ocr_unavailable:
            self.fail(label, "local CPU OCR unavailable")
            return ""
        if self.ocr_attempts >= MAX_OCR_UNITS:
            self.skip(label, f"OCR attempt limit {MAX_OCR_UNITS}")
            return ""
        try:
            png = make_png()
            key = hashlib.sha256(png).digest()
            if key in self.ocr_cache:
                lines = self.ocr_cache[key]
                self.note("Identical raster inputs reused this document's OCR result; each occurrence remains a processed unit, without another OCR call.")
            else:
                self.ocr_attempts += 1
                lines = _ocr_lines(png)
                self.ocr_count += 1
                self.ocr_cache[key] = lines
            self.processed += 1
            text = _new_ocr_text(lines, native)
            if not lines:
                self.note(f"{label}: CPU OCR found no text; this is not evidence that the image has no instructional meaning.")
            elif not text:
                self.note(f"{label}: normalized OCR text already present in native/earlier text was omitted to avoid duplication.")
            return text
        except _Skipped as exc:
            self.skip(label, str(exc))
        except ImportError:
            self.ocr_unavailable = True
            self.fail(label, "local CPU OCR unavailable; installed models/dictionary required, no download attempted")
        except Exception:
            self.fail(label, "CPU OCR or raster decoding failed")
        return ""

    def result(self, *, error=False, unsupported=False):
        truncated = self.chars > len(self.excerpt)
        if truncated:
            self.review("Preview limited to 50,000 characters; later units were still processed within the documented limits. The complete original is retained.")
        if self.failed:
            self.review(f"{len(self.failed)} extraction unit(s) failed; inspect failed_units and the retained original.")
        if not self.text_chars and not error:
            self.review("No extractable text was found; OCR is not proof of an empty document. Teacher inspection is required.")
        status = "error" if error else "needs_review" if self.review_reasons else "ready"
        if unsupported:
            method = "unsupported"
        elif error:
            method = "error"
        elif self.ocr_count:
            method = "native+cpu_ocr" if self.native_text else "cpu_ocr"
        elif self.native_text or self.processed:
            method = "native"
        else:
            method = "cpu_ocr_unavailable" if self.ocr_unavailable else "cpu_ocr_failed" if self.ocr_attempts else "not_extracted"
        report = ExtractionReport(1, method, self.total, self.processed, self.ocr_count,
                                  self.images, self.chars, len(self.excerpt), truncated,
                                  self.visual, tuple(self.failed), tuple(self.notes))
        # Full per-unit detail is in the structured report; keep the UI warning
        # bounded without hiding that partial reads/skips occurred.
        warnings = list(self.review_reasons)
        skipped = sum(note.startswith("Skipped:") for note in self.notes)
        if skipped:
            warnings = [w for w in warnings if not w.startswith("Skipped:")]
            warnings.append(f"Extraction skipped {skipped} unit/range(s); see coverage notes for exact labels and limits.")
        return {"status": status, "excerpt": self.excerpt,
                "warning": " ".join(warnings) if warnings else None,
                "extraction_report": report.to_dict()}


def _image_png(blob, extraction):
    from PIL import Image
    with Image.open(io.BytesIO(blob)) as image:
        if getattr(image, "n_frames", 1) > 1:
            extraction.review("Embedded animated images use the first frame only; later frames and temporal meaning are not interpreted.", visual=True)
        return _frame_png(image, extraction)


def _frame_png(image, extraction):
    from PIL import Image, ImageOps
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise _Skipped(f"image exceeds {MAX_IMAGE_PIXELS:,}-pixel decode limit")
    image = ImageOps.exif_transpose(image).convert("RGBA")
    if image.width * image.height > MAX_RENDER_PIXELS:
        scale = (MAX_RENDER_PIXELS / (image.width * image.height)) ** 0.5
        image = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))))
        extraction.note(f"Large OCR rasters were downscaled to at most {MAX_RENDER_PIXELS:,} pixels; small text may be lost.")
    background = Image.new("RGBA", image.size, "white")
    background.alpha_composite(image)
    output = io.BytesIO()
    background.convert("RGB").save(output, format="PNG")
    return output.getvalue()


def _office_guard(path):
    from .extraction_limits import office_guard
    try:
        office_guard(path, max_expanded=MAX_EXPANDED_BYTES)
    except ValueError as exc:
        message = str(exc)
        if message.startswith("Expanded Office"):
            message = "expanded Office" + message[len("Expanded Office"):]
        raise _Skipped(message) from None


def _render_pdf(page, rect, extraction):
    import pymupdf
    rect = pymupdf.Rect(rect) & page.rect
    if rect.is_empty or rect.is_infinite:
        raise _Skipped("raster region is outside the visible PDF page")
    scale = min(2.0, (MAX_RENDER_PIXELS / max(1, rect.get_area())) ** 0.5)
    if scale < 2:
        extraction.note(f"PDF OCR rasters were downscaled to at most {MAX_RENDER_PIXELS:,} pixels; small text may be lost.")
    return page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), clip=rect, colorspace=pymupdf.csRGB, alpha=False).tobytes("png")


def _pdf(path, extraction):
    import pymupdf
    with pymupdf.open(path) as document:
        if document.needs_pass:
            raise _Skipped("encrypted PDF requires an unlocked/re-exported copy")
        extraction.total = len(document)
        extraction.native_limit(len(document), "PDF pages")
        extraction.note("PDF native text is read in approximate page order; table alignment/merged cells and visual layout are not reconstructed.")
        for index in range(min(len(document), MAX_NATIVE_UNITS)):
            label = f"PDF page {index + 1}"
            try:
                page = document[index]
                text = page.get_text(sort=True)
                extraction.add(text, label, native=True)
                extraction.processed += 1
            except Exception:
                extraction.fail(label, "native page parsing failed; images were not inspected")
                continue
            try:
                # Display-list occurrences include inline images (xref=0), not
                # unused resource-dictionary images. Any native header density
                # is irrelevant: every displayed raster region is considered.
                images = page.get_image_info()
                extraction.images += len(images)
                extraction.total += len(images)
                for number, image in enumerate(images, 1):
                    image_label = f"{label} image {number} (CPU OCR)"
                    extra = extraction.ocr(image_label, lambda bbox=image["bbox"]: _render_pdf(page, pymupdf.Rect(bbox) * page.rotation_matrix, extraction), text)
                    extraction.add(extra, image_label)
                    text += "\n" + extra
                drawings = bool(page.get_drawings())
                annotations = bool(list(page.annots() or []))
                if drawings or annotations:
                    extraction.review("PDF vector drawings/annotations remain in the original; text labels alone do not capture their visual meaning.", visual=True)
                if not text.strip() and not images:
                    extraction.total += 1
                    extra = extraction.ocr(label + " page render (CPU OCR)", lambda: _render_pdf(page, page.rect, extraction))
                    extraction.add(extra, label + " page render (CPU OCR)")
            except Exception:
                extraction.fail(label + " visual inventory", "images/vector content could not be inspected")


def _word_text(node):
    """Body-order native paragraphs/tables, including nested textbox text."""
    if node.tag == W + "p":
        return ["".join(item.text or "" if item.tag == W + "t" else "\t" if item.tag == W + "tab" else "\n" if item.tag in {W + "br", W + "cr"} else "" for item in node.iter()).strip()]
    if node.tag == W + "tbl":
        rows = []
        for row in node:
            if row.tag == W + "tr":
                cells = ["\n".join(_word_text(cell)).strip() for cell in row if cell.tag == W + "tc"]
                rows.append(" | ".join(cells))
        return rows
    return [text for child in node for text in _word_text(child)]


def _office_images(root, part, label, extraction, native=""):
    nodes = [node for node in root.iter() if node.tag in {A + "blip", V + "imagedata"}]
    extraction.images += len(nodes)
    extraction.total += len(nodes)
    if nodes:
        extraction.note("Office OCR reads embedded source image pixels, not a rendered Office page/slide; cropping, stacking, image effects and hidden/overlaid regions are not resolved, so text outside a visible crop may appear.")
    for number, node in enumerate(nodes, 1):
        image_label = f"{label} image {number} (CPU OCR)"
        rid = node.get(R + "embed") or node.get(R + "id")
        if not rid or rid not in part.rels:
            extraction.skip(image_label, "linked/missing image relationship; external resources are never fetched")
            extraction.visual = True
            continue
        relation = part.rels[rid]
        if relation.is_external:
            extraction.skip(image_label, "external image; external resources are never fetched")
            extraction.visual = True
            continue
        extra = extraction.ocr(image_label, lambda rel=relation: _image_png(rel.target_part.blob, extraction), native)
        extraction.add(extra, image_label)
        native += "\n" + extra
    return native


def _docx(path, extraction):
    from docx import Document
    _office_guard(path)
    document = Document(path)
    stories = [("DOCX body", document.part)]
    known = {str(document.part.partname)}
    counters = {}
    for relation in document.part.rels.values():
        kind = relation.reltype.rsplit("/", 1)[-1]
        if kind not in {"header", "footer", "footnotes", "endnotes"} or relation.is_external:
            continue
        part = relation.target_part
        if str(part.partname) not in known:
            known.add(str(part.partname))
            counters[kind] = counters.get(kind, 0) + 1
            stories.append((f"DOCX {kind} {counters[kind]}", part))
    extraction.total = len(stories)
    extraction.native_limit(len(stories), "DOCX stories")
    extraction.note("DOCX units are body/header/footer/footnote/endnote stories, not rendered pages. Native tables are read by row; pagination, list numbering and complex textbox layout are not reconstructed.")
    for label, part in stories[:MAX_NATIVE_UNITS]:
        try:
            root = ET.fromstring(part.blob)
            text = "\n".join(_word_text(root))
            extraction.add(text, label, native=True)
            extraction.processed += 1
        except Exception:
            extraction.fail(label, "native story parsing failed; images were not inspected")
            continue
        _office_images(root, part, label, extraction, text)
        if any(node.tag in {W + "object", W + "txbxContent"} or (node.tag == A + "graphicData" and any(word in node.get("uri", "") for word in ("chart", "diagram"))) for node in root.iter()):
            extraction.review("DOCX textboxes/charts/SmartArt/embedded objects may have incomplete layout or visual meaning; inspect the original.", visual=True)


def _slide_text(shapes, extraction):
    from pptx.enum.shapes import MSO_SHAPE_TYPE as TYPE
    parts = []
    for shape in sorted(shapes, key=lambda item: (item.top, item.left)):
        if shape.shape_type == TYPE.GROUP:
            parts.extend(_slide_text(shape.shapes, extraction))
            extraction.review("Grouped slide shapes were traversed recursively; text order is approximate, not a reconstruction of diagram relationships.", visual=True)
        if shape.has_text_frame and shape.text.strip():
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            extraction.note("Native slide tables were read by row; merged-cell alignment and visual formatting require the original.")
        if shape.shape_type in {TYPE.CHART, TYPE.AUTO_SHAPE, TYPE.FREEFORM, TYPE.LINE, TYPE.MEDIA, TYPE.EMBEDDED_OLE_OBJECT, TYPE.LINKED_OLE_OBJECT}:
            extraction.review("Slide charts/vector shapes/connectors/media/embedded objects are not semantically interpreted; native labels and OCR do not capture their full meaning.", visual=True)
    return parts


def _pptx(path, extraction):
    from pptx import Presentation
    _office_guard(path)
    document = Presentation(path)
    extraction.total = len(document.slides)
    extraction.native_limit(len(document.slides), "Slides")
    extraction.note("Slide units include hidden slides. Speaker notes, animations, master/layout text and theme-generated graphics are not extracted; original presentation remains authoritative.")
    for index in range(min(len(document.slides), MAX_NATIVE_UNITS)):
        label = f"Slide {index + 1}"
        try:
            slide = document.slides[index]
            text = "\n".join(_slide_text(slide.shapes, extraction))
            extraction.add(text, label, native=True)
            extraction.processed += 1
        except Exception:
            extraction.fail(label, "native slide parsing failed; images were not inspected")
            continue
        # XML references include grouped pictures, table/shape fills, and slide
        # backgrounds, not only top-level PICTURE shapes.
        for suffix, part in (("", slide.part), (" layout", slide.slide_layout.part), (" master", slide.slide_layout.slide_master.part)):
            try:
                root = ET.fromstring(part.blob)
                text = _office_images(root, part, label + suffix, extraction, text)
                if suffix and any(node.tag == A + "blip" for node in root.iter()):
                    extraction.note("Layout/master image resources were OCR-read as template images; visibility/placement on individual slides may differ.")
                if any(node.tag == A + "graphicData" and "diagram" in node.get("uri", "") for node in root.iter()):
                    extraction.review("Slide SmartArt relationships/diagram content are not fully extracted or interpreted.", visual=True)
            except Exception:
                extraction.fail(label + suffix + " visual inventory", "embedded images could not be inspected")


def _xlsx(path, extraction):
    from tools.material_stage16.workbook_inventory import read_workbook, ledger_rows
    native = read_workbook(path)
    extraction.total = native['sheet_count']
    for note in native['warnings']: extraction.note(note)
    extraction.note('Workbook units are original sheets, not PDF pages. Preview text is a native cell ledger with sheet/cell locators; formulas and stored caches are separate, never evaluated. Comments remain outside ordinary cell reading.')
    extraction.native_limit(native['sheet_count'], 'Workbook sheets')
    preview_native = dict(native, sheets=native['sheets'][:MAX_NATIVE_UNITS])
    for row in ledger_rows(preview_native): extraction.add(row, None, native=True)
    extraction.processed = min(native['sheet_count'], MAX_NATIVE_UNITS)
    extraction.images = native['media_count']
    extraction.review('Excel layout, charts, cell formatting and original media are not reconstructed by this bounded text preview; complete processing retains all stored cells and source-media previews.', visual=True)


def _legacy_office(path, extraction):
    # Reuse the controlled worker converter only where it is locally available;
    # never start SSH/model services or silently treat binary bytes as OOXML.
    import platform
    import shutil
    import tempfile
    if platform.system() != 'Linux' or not (shutil.which('libreoffice') or shutil.which('soffice')):
        raise _Skipped('legacy Office preview requires the configured complete-document worker or a local bounded LibreOffice converter; none was started')
    from tools.material_stage16.legacy_office import binary_inventory
    from tools.material_stage13.parse_material import convert
    fmt = path.suffix.lower().lstrip('.')
    binary_inventory(path, fmt)
    with tempfile.TemporaryDirectory(prefix='aicademiya-office-preview-') as directory:
        temp = Path(directory); source = temp/('source.'+fmt)
        shutil.copyfile(path, source); source.chmod(0o400)
        derived, _ = convert(source, temp, 'pptx' if fmt == 'ppt' else 'xlsx', 120,
                             {**__import__('os').environ, 'CUDA_VISIBLE_DEVICES':'', 'HF_HUB_OFFLINE':'1', 'SAL_USE_VCLPLUGIN':'svp'})
        extraction.note('Legacy '+fmt.upper()+' native preview text is LibreOffice-converted OOXML, not original binary text/caches. Original bytes are retained; no macros or link refresh are allowed.')
        if fmt == 'ppt': _pptx(derived, extraction)
        else: _xlsx(derived, extraction)


def _standalone_image(path, extraction):
    from PIL import Image
    with Image.open(path) as image:
        frames = getattr(image, "n_frames", 1)
        extraction.total = extraction.images = frames
        if frames > 1:
            extraction.review("Standalone image frames are OCR-read separately; animation timing and temporal meaning are not interpreted.", visual=True)
        for index in range(frames):
            label = f"Image frame {index + 1} (CPU OCR)"
            if index >= MAX_NATIVE_UNITS:
                extraction.skip(f"Image frames {index + 1}-{frames}", f"frame inspection limit {MAX_NATIVE_UNITS}")
                break
            image.seek(index)
            extraction.add(extraction.ocr(label, lambda: _frame_png(image, extraction)), label)


def extract_document(path):
    path = Path(path)
    extraction = _Extraction()
    if path.suffix.lower() not in SUPPORTED:
        extraction.review("Unsupported format; use TXT, MD, PDF, DOCX, PPTX, PPT, XLS, XLSX, PNG, JPG, JPEG or WebP; legacy DOC uses complete processing.")
        return extraction.result(error=True, unsupported=True)
    try:
        if path.stat().st_size > MAX_INPUT_BYTES:
            raise _Skipped("document exceeds the 25 MiB input limit")
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            extraction.total = 1
            text = path.read_bytes().decode("utf-8-sig")
            extraction.add(text, None, native=True)
            extraction.processed = 1
            if suffix == ".md" and re.search(r"!\[[^\]]*\]\([^)]+\)", text):
                extraction.review("Markdown image references were retained as text only; local/external images were not fetched or interpreted.", visual=True)
        elif suffix == ".pdf":
            _pdf(path, extraction)
        elif suffix == ".docx":
            _docx(path, extraction)
        elif suffix == ".pptx":
            _pptx(path, extraction)
        elif suffix == ".xlsx":
            _xlsx(path, extraction)
        elif suffix in {".ppt", ".xls"}:
            _legacy_office(path, extraction)
        else:
            _standalone_image(path, extraction)
    except UnicodeDecodeError:
        extraction.fail("Text", "UTF-8 decoding failed")
        extraction.review("Text could not be decoded as UTF-8; re-export it as UTF-8.")
        return extraction.result(error=True)
    except _Skipped as exc:
        extraction.skip("Document", str(exc))
        extraction.review("Document was not extracted: " + str(exc) + ".")
        return extraction.result(error=True)
    except Exception:
        # Never expose parser exception strings, private paths or source text.
        extraction.fail("Document", "format/file/parser error")
        extraction.review("Document extraction failed. The original is retained for teacher inspection.")
        return extraction.result(error=not extraction.processed)
    return extraction.result()
