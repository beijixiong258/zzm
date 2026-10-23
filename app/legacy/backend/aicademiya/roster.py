"""Course-scoped rosters, hash-only atomic imports and requested retention rules."""
import csv
import io
import json
import zipfile
from datetime import datetime, timedelta, timezone
from fastapi import File, Form, HTTPException, UploadFile, Depends
from . import models as m
from .db import new_id, now, password_hash, public_user
from .services import course_access

MAX_ROWS = 1000
MAX_BYTES = 2 * 1024 * 1024


def expired(value, at=None):
    return bool(value and datetime.fromisoformat(value.replace("Z", "+00:00")) <= (at or datetime.now(timezone.utc)))


def purge_learning(conn, uid, cid=None):
    for table in ("chat_requests", "messages", "revisions", "attempts", "workspaces"):
        if cid:
            conn.execute(f"DELETE FROM {table} WHERE user_id=? AND session_id IN (SELECT id FROM sessions WHERE course_id=?)", (uid, cid))
        else:
            conn.execute(f"DELETE FROM {table} WHERE user_id=?", (uid,))


def maintenance(conn, *, at=None, dry_run=False, limit=100):
    at = at or datetime.now(timezone.utc)
    stamp = at.astimezone(timezone.utc).isoformat()
    # Trash pauses destructive retention, not the original independent deadlines.
    accounts = conn.execute("""SELECT id,purge_at FROM users u
        WHERE role='student' AND archived_at IS NOT NULL AND purge_at<=?
        AND NOT EXISTS (SELECT 1 FROM enrollments e JOIN course_trash t ON t.course_id=e.course_id WHERE e.user_id=u.id)
        AND NOT EXISTS (SELECT 1 FROM courses c WHERE c.owner_id=u.id)
        AND NOT EXISTS (SELECT 1 FROM course_staff cs WHERE cs.user_id=u.id)
        AND NOT EXISTS (SELECT 1 FROM teacher_previews tp WHERE tp.user_id=u.id)
        AND NOT EXISTS (SELECT 1 FROM course_trash t WHERE t.trashed_by=u.id)
        AND NOT EXISTS (SELECT 1 FROM import_staging i WHERE i.owner_id=u.id AND i.expires_at>?)
        AND NOT EXISTS (SELECT 1 FROM account_imports i WHERE i.owner_id=u.id AND i.expires_at>?)
        ORDER BY purge_at,id LIMIT ?""", (stamp, stamp, stamp, limit)).fetchall()
    enrollments = conn.execute("""SELECT course_id,user_id,purge_at FROM enrollments e
        WHERE archived_at IS NOT NULL AND purge_at<=?
        AND NOT EXISTS (SELECT 1 FROM course_trash t WHERE t.course_id=e.course_id)
        ORDER BY purge_at,user_id,course_id LIMIT ?""", (stamp, limit)).fetchall()
    preview = {"accounts": [r["id"] for r in accounts], "enrollments": [{"course_id": r["course_id"], "user_id": r["user_id"]} for r in enrollments], "dry_run": dry_run}
    if dry_run:
        return preview
    for row in enrollments:
        purge_learning(conn, row["user_id"], row["course_id"])
        conn.execute("DELETE FROM enrollments WHERE course_id=? AND user_id=?", (row["course_id"], row["user_id"]))
    for row in accounts:
        uid = row["id"]
        # A former Admin/Teacher can now be a Student. Expired owned import
        # receipts must not make deletion fail; durable staff references above
        # retain the disabled identity instead of discarding teaching history.
        conn.execute('DELETE FROM import_staging WHERE owner_id=? AND expires_at<=?', (uid, stamp))
        conn.execute('DELETE FROM account_imports WHERE owner_id=? AND expires_at<=?', (uid, stamp))
        purge_learning(conn, uid)
        conn.execute("DELETE FROM enrollments WHERE user_id=?", (uid,))
        conn.execute("DELETE FROM auth_sessions WHERE user_id=?", (uid,))
        conn.execute("DELETE FROM users WHERE id=? AND role='student'", (uid,))
    for row in conn.execute("SELECT id FROM import_staging WHERE expires_at<=? ORDER BY expires_at,id LIMIT ?", (stamp, limit)).fetchall():
        conn.execute("DELETE FROM import_staging WHERE id=?", (row["id"],))
    for row in conn.execute('SELECT id FROM account_imports WHERE expires_at<=? ORDER BY expires_at,id LIMIT ?', (stamp, limit)).fetchall():
        conn.execute('DELETE FROM account_imports WHERE id=?', (row['id'],))
    return preview


def roster(conn, archived=False, cid=None):
    if cid:
        rows = conn.execute("SELECT u.*,e.archived_at AS enrollment_archived,e.purge_at AS enrollment_purge FROM users u JOIN enrollments e ON e.user_id=u.id WHERE e.course_id=? AND u.role='student' ORDER BY u.name,u.id", (cid,)).fetchall()
    else:
        rows = conn.execute("SELECT * FROM users WHERE role='student' ORDER BY name,id").fetchall()
    result = []
    for r in rows:
        stamp, deadline = (r["enrollment_archived"], r["enrollment_purge"]) if cid else (r["archived_at"], r["purge_at"])
        if bool(stamp) != archived:
            continue
        query = "SELECT 1 FROM workspaces w WHERE user_id=? AND (last_active IS NOT NULL OR status<>'not_started')"
        args = [r["id"]]
        if cid:
            query += " AND session_id IN (SELECT id FROM sessions WHERE course_id=?)"
            args.append(cid)
        result.append({"user": public_user(r, conn), "archived_at": stamp, "purge_at": deadline, "has_work": bool(conn.execute(query, args).fetchone())})
    return result


def batch(conn, ids, restore=False, cid=None):
    rows = []
    for uid in ids:
        user = conn.execute("SELECT * FROM users WHERE id=? AND role='student'", (uid,)).fetchone()
        relation = conn.execute("SELECT * FROM enrollments WHERE user_id=? AND course_id=?", (uid, cid)).fetchone() if cid else user
        if not user or not relation:
            raise HTTPException(404, "Every target must be a student in the authorized scope")
        if expired(relation["purge_at"]):
            raise HTTPException(410, "Archived record has expired and can no longer be restored")
        rows.append((uid, relation))
    stamp = now()
    deadline = (datetime.fromisoformat(stamp) + timedelta(days=30)).isoformat()
    affected = []
    for uid, row in rows:
        if restore == bool(row["archived_at"]):
            values = (None, None) if restore else (stamp, deadline)
            if cid:
                conn.execute("UPDATE enrollments SET archived_at=?,purge_at=? WHERE course_id=? AND user_id=?", (*values, cid, uid))
            else:
                conn.execute("UPDATE users SET archived_at=?,purge_at=? WHERE id=?", (*values, uid))
                if not restore:
                    conn.execute("DELETE FROM auth_sessions WHERE user_id=?", (uid,))
            affected.append(uid)
    return {"affected_ids": affected}


def parse_rows(data, suffix, *, global_accounts=False):
    max_columns = 4 if global_accounts else 3
    if suffix == ".csv":
        try:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("gb18030")
            rows = list(csv.reader(io.StringIO(text)))
        except (UnicodeError, csv.Error):
            raise HTTPException(422, "Provide a valid UTF-8 or GB18030 CSV") from None
    elif suffix == ".xlsx":
        try:
            from .extraction_limits import office_guard
            try:
                office_guard(io.BytesIO(data), max_expanded=20 * MAX_BYTES)
            except ValueError as exc:
                if "limit" in str(exc):
                    raise HTTPException(413, "XLSX expanded size exceeds the import limit") from None
                raise
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist()) > 20 * MAX_BYTES or len(z.infolist()) > 500:
                    raise HTTPException(413, "XLSX expanded size exceeds the import limit")
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=not global_accounts)
            sheet = book.active
            try:
                if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > max_columns:
                    raise HTTPException(413, 'Import is limited to 1000 rows and the supported columns')
                # Do not trust a producer-supplied worksheet dimension for bounds.
                sheet.reset_dimensions()
                rows = []
                for row_index, cells in enumerate(sheet.iter_rows(max_col=max_columns + 1)):
                    if row_index > MAX_ROWS:
                        raise HTTPException(413, 'Import is limited to 1000 rows; no rows were imported')
                    if global_accounts and any(c.data_type == 'f' for c in cells):
                        raise HTTPException(422, 'Import cells must contain literal values, not formulas')
                    values = [c.value for c in cells]
                    while values and values[-1] is None:
                        values.pop()
                    rows.append(values)
                while rows and not any(v is not None and str(v).strip() for v in rows[-1]):
                    rows.pop()
            finally:
                book.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, "Provide a valid XLSX workbook") from None
    else:
        raise HTTPException(415, "Supported roster formats: CSV and XLSX")
    if not rows:
        raise HTTPException(422, "The roster is empty")
    aliases = {"姓名": "name", "邮箱": "email", "电子邮箱": "email", "电子邮件": "email",
               "临时密码": "temporary_password", "角色": "role"}
    headers = [aliases.get(str(x or "").strip().lower(), str(x or "").strip().lower()) for x in rows[0]]
    permitted = {'name', 'email', 'temporary_password'} | ({'role'} if global_accounts else set())
    if len(set(headers)) != len(headers) or not {'name', 'email'} <= set(headers) or not set(headers) <= permitted:
        raise HTTPException(422, 'Headers must be name,email and optional ' + ('role,' if global_accounts else '') + 'temporary_password')
    if len(rows) > MAX_ROWS + 1:
        raise HTTPException(413, "Roster is limited to 1000 rows; no rows were imported")
    result = []
    for number, values in enumerate(rows[1:], 2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if len(values) > len(headers):
            raise HTTPException(422, "A roster row contains excess columns")
        result.append((number, {k: (str(v) if v is not None else '') if k == 'temporary_password' else str(v or '').strip()
                                for k, v in zip(headers, values)}))
    if not result:
        raise HTTPException(422, "No student rows were found")
    return result


def classify(conn, cid, email):
    user = conn.execute("SELECT * FROM users WHERE email=?", (email,)).fetchone()
    if not user:
        return "create", None, None
    if user["role"] != "student":
        return "error", user["id"], None
    if user["archived_at"]:
        return "error", user["id"], "global_archive"
    e = conn.execute("SELECT * FROM enrollments WHERE course_id=? AND user_id=?", (cid, user["id"])).fetchone()
    return ("restore" if e and e["archived_at"] else "skip" if e else "enroll"), user["id"], e["archived_at"] if e else None


def ensure_staging_capacity(conn, uid, kind):
    maximum = 10 if kind == "worksheet" else 30
    count = conn.execute("SELECT count(*) FROM import_staging WHERE owner_id=? AND kind=? AND expires_at>?", (uid, kind, now())).fetchone()[0]
    if count >= maximum:
        raise HTTPException(429, "Too many uncommitted import previews; commit a preview or wait for expiry")


def preview_import(conn, cid, uid, rows, default_password=None):
    ensure_staging_capacity(conn, uid, "roster")
    staged, emails = [], set()
    for number, raw in rows:
        name, email = raw.get("name", ""), raw.get("email", "").lower()
        action, existing, enrollment = classify(conn, cid, email)
        errors, warnings = [], []
        if not name or len(name) > 100 or "@" not in email or " " in email or len(email) > 254:
            errors.append("Provide a name (1–100 characters) and valid email")
        if email in emails:
            errors.append("Duplicate email in this roster")
        emails.add(email)
        if action == "error":
            errors.append("Email belongs to staff or a globally archived account; no account was changed")
        hashed = None
        if action == "create":
            password = raw.get("temporary_password") or default_password or ""
            try:
                m.validate_new_password(password)
                hashed = password_hash(password)
            except ValueError:
                errors.append("New accounts require 8–512 characters including a letter and digit")
        else:
            warnings.append("Existing account name and password will not be changed")
        if action == "restore":
            relation = conn.execute("SELECT purge_at FROM enrollments WHERE course_id=? AND user_id=?", (cid, existing)).fetchone()
            if expired(relation[0]):
                errors.append("Archived enrollment has expired")
        staged.append({"row_number": number, "name": name, "email": email, "action": "error" if errors else action,
                       "errors": errors, "warnings": warnings, "password_hash": hashed,
                       "expected_user_id": existing, "expected_enrollment": enrollment})
    iid = new_id()
    deadline = (datetime.now(timezone.utc) + timedelta(minutes=30)).isoformat()
    conn.execute("INSERT INTO import_staging VALUES(?,?,?,?,?,?,?)", (iid, uid, cid, None, "roster", deadline, json.dumps(staged)))
    visible = [{k: r[k] for k in ("row_number", "name", "email", "action", "errors", "warnings")} for r in staged]
    errors = sum(bool(r["errors"]) for r in staged)
    return {"import_id": iid, "rows": visible, "valid_count": len(staged)-errors, "error_count": errors}


def staging(conn, iid, uid, cid, kind, sid=None):
    r = conn.execute("SELECT * FROM import_staging WHERE id=? AND owner_id=? AND course_id=? AND kind=?", (iid, uid, cid, kind)).fetchone()
    if not r or (sid and r["session_id"] != sid):
        raise HTTPException(404, "Import preview not found in this scope")
    if expired(r["expires_at"]):
        raise HTTPException(410, "Import preview expired; upload again")
    return json.loads(r["payload_json"])


def commit_import(conn, cid, uid, iid):
    rows = staging(conn, iid, uid, cid, "roster")
    if any(r["errors"] for r in rows):
        raise HTTPException(422, "Import contains errors; fix and upload again")
    for r in rows:
        if classify(conn, cid, r["email"]) != (r["action"], r["expected_user_id"], r["expected_enrollment"]):
            raise HTTPException(409, "Roster changed since preview; upload again")
        if r["action"] == "restore":
            e = conn.execute("SELECT purge_at FROM enrollments WHERE course_id=? AND user_id=?", (cid, r["expected_user_id"])).fetchone()
            if expired(e[0]):
                raise HTTPException(410, "Archived enrollment has expired")
    counts = {"created": 0, "enrolled": 0, "restored": 0, "skipped": 0}
    for r in rows:
        action, target = r["action"], r["expected_user_id"]
        if action == "create":
            target = new_id()
            conn.execute("INSERT INTO users(id,name,email,role,password_hash,must_change_password) VALUES(?,?,?,'student',?,1)", (target, r["name"], r["email"], r["password_hash"]))
            conn.execute("INSERT INTO enrollments(course_id,user_id) VALUES(?,?)", (cid, target))
            counts["created"] += 1
        elif action == "enroll":
            conn.execute("INSERT INTO enrollments(course_id,user_id) VALUES(?,?)", (cid, target))
            counts["enrolled"] += 1
        elif action == "restore":
            conn.execute("UPDATE enrollments SET archived_at=NULL,purge_at=NULL WHERE course_id=? AND user_id=?", (cid, target))
            counts["restored"] += 1
        else:
            counts["skipped"] += 1
    conn.execute("DELETE FROM import_staging WHERE id=?", (iid,))
    return counts


def register_roster_routes(app, db, teacher_user, admin_user):
    @app.get("/api/teacher/courses/{course_id}/students")
    def course_roster(course_id: str, archived: bool = False, user=Depends(teacher_user)):
        with db.connect() as conn:
            course_access(conn, course_id, user)
            return roster(conn, archived, course_id)

    @app.get("/api/admin/students")
    def global_roster(archived: bool = False, user=Depends(admin_user)):
        with db.connect() as conn:
            from .course_staff import live_actor
            live_actor(conn, user, admin=True)
            return roster(conn, archived)

    @app.post("/api/teacher/courses/{course_id}/students/archive")
    def course_archive(course_id: str, payload: m.StudentBatch, user=Depends(teacher_user)):
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            course_access(conn, course_id, user)
            return batch(conn, payload.user_ids, cid=course_id)

    @app.post("/api/teacher/courses/{course_id}/students/restore")
    def course_restore(course_id: str, payload: m.StudentBatch, user=Depends(teacher_user)):
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            course_access(conn, course_id, user)
            return batch(conn, payload.user_ids, True, course_id)

    @app.post("/api/admin/students/archive")
    def global_archive(payload: m.StudentBatch, user=Depends(admin_user)):
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            from .course_staff import live_actor
            live_actor(conn, user, admin=True)
            return batch(conn, payload.user_ids)

    @app.post("/api/admin/students/restore")
    def global_restore(payload: m.StudentBatch, user=Depends(admin_user)):
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            from .course_staff import live_actor
            live_actor(conn, user, admin=True)
            return batch(conn, payload.user_ids, True)

    @app.post("/api/teacher/courses/{course_id}/students/import/preview")
    def import_preview(course_id: str, file: UploadFile = File(...), default_temporary_password: str | None = Form(None), user=Depends(teacher_user)):
        from pathlib import Path
        with db.connect() as conn:
            course_access(conn, course_id, user)
        data = file.file.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            raise HTTPException(413, "Roster exceeds the 2 MiB limit")
        rows = parse_rows(data, Path(file.filename or "").suffix.lower())
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            course_access(conn, course_id, user)
            return preview_import(conn, course_id, user["id"], rows, default_temporary_password)

    @app.post("/api/teacher/courses/{course_id}/students/import/{import_id}/commit")
    def import_commit(course_id: str, import_id: str, payload: m.Empty, user=Depends(teacher_user)):
        with db.connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            course_access(conn, course_id, user)
            return commit_import(conn, course_id, user["id"], import_id)
