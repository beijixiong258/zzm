"""Actual local instruction extraction through authenticated multipart API.

Synthetic native DOCX/TXT/PDF and one genuine installed CPU OCR PDF; no LLM.
"""
import io
import json
from pathlib import Path
import pytest
import pymupdf
from backend.aicademiya import material_extraction as extraction
from .conftest import expect
from .test_stage3_persistence import records
from .test_stage4_extraction import REPORT_KEYS

URL = "/api/teacher/sessions/s4/instruction-import/preview"


def post_file(teacher, path, filename=None):
    return teacher.post(URL, files={"file": (filename or path.name, path.read_bytes())})


def assert_preview(result):
    assert set(result) == {"filename","text","status","warning","extraction_report","max_instruction_chars"}
    assert result["max_instruction_chars"] == 30000
    assert set(result["extraction_report"]) == REPORT_KEYS
    assert result["extraction_report"]["excerpt_chars"] == len(result["text"]) <= 50000
    assert isinstance(result["extraction_report"]["notes"], list)


@pytest.mark.parametrize("extension", ["txt", "docx", "pdf"])
def test_actual_native_extraction_owner_private_temp_no_attachment_save_or_publication(teacher, student, admin, app, tmp_path, monkeypatch, extension):
    path = tmp_path / ("native-instructions." + extension)
    token = "SYNTHETIC INSTRUCTION NATIVE TEXT"
    if extension == "txt":
        path.write_text(token + "\nKeep learner responses intact.", encoding="utf-8-sig")
    elif extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph(token)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0,0).text = "Progressive hints"
        table.cell(0,1).text = "No formal grades"
        doc.save(path)
    else:
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_text((40,50), token)
            page.insert_text((40,75), "Keep teacher references private.")
            doc.save(path)
    original = path.read_bytes()
    with app.state.database.connect() as conn:
        before = records(conn)
    root_before = set(app.state.database.path.parent.iterdir())
    def no_model(*a, **kw):
        raise AssertionError("Instruction preview never calls a model")
    monkeypatch.setattr(app.state.tutor, "complete", no_model)
    monkeypatch.setattr(app.state.tutor, "feedback", no_model)
    result = expect(post_file(teacher, path, "C:\\private\\" + path.name))
    assert_preview(result)
    assert token in result["text"] and result["status"] == "ready"
    assert result["extraction_report"]["method"] == "native"
    assert result["filename"] == path.name
    assert str(app.state.database.path.parent) not in json.dumps(result)
    assert not any(p.name.startswith("instruction-preview-") for p in app.state.database.path.parent.iterdir())
    assert set(app.state.database.path.parent.iterdir()) == root_before
    assert path.read_bytes() == original
    expect(student.post(URL, files={"file": (path.name, original)}),403)
    expect(admin.post(URL, files={"file": (path.name, original)}),403)
    with app.state.database.connect() as conn:
        assert records(conn) == before
    (tmp_path / ("native-" + extension + "-verification.json")).write_text(json.dumps({"format":extension,"status":result["status"],"report":result["extraction_report"],"original_unchanged":True,"all_tables_unchanged":True,"private_temp_removed":True},indent=2), encoding="utf-8")


@pytest.mark.parametrize("length", [40000,65000])
def test_preview_field_and_extraction_limits_visible_never_silent_field_truncation(teacher, app, tmp_path, length):
    path = tmp_path / "long.txt"
    path.write_text("x" * length, encoding="utf-8")
    result = expect(post_file(teacher, path))
    assert_preview(result)
    assert len(result["text"]) == min(length,50000)
    assert result["status"] == "needs_review" and "30,000" in result["warning"]
    assert result["extraction_report"]["truncated"] == (length > 50000)
    assert result["extraction_report"]["chars_extracted"] == length
    # Import returns editable text only. Applying/saving still requires an explicit valid draft write.
    before = expect(teacher.get("/api/sessions/s4"))["instructions"]
    expect(teacher.patch("/api/sessions/s4", json={"instructions":result["text"]}),422)
    assert expect(teacher.get("/api/sessions/s4"))["instructions"] == before
    edited = result["text"][:80] + "\nTeacher's explicit edited instruction."
    expect(teacher.patch("/api/sessions/s4", json={"instructions":edited}))
    assert expect(teacher.get("/api/sessions/s4"))["instructions"] == edited


@pytest.mark.parametrize("filename,data,status", [("broken.docx",b"not a zip","error"),("broken.pdf",b"not a pdf","error"),("bad.txt",b"\xff\xfe\xff","error"),("empty.txt",b"","needs_review")])
def test_parser_error_empty_and_encoding_are_visible_and_preserve_saved_content(teacher, app, filename, data, status):
    with app.state.database.connect() as conn:
        before = records(conn)
    result = expect(teacher.post(URL,files={"file":(filename,data)}))
    assert_preview(result)
    assert result["status"] == status and result["text"] == "" and result["warning"]
    assert not list(app.state.database.path.parent.glob("instruction-preview-*"))
    with app.state.database.connect() as conn:
        assert records(conn) == before


def test_unsupported_and_25_mib_limit_before_extraction(teacher, app, monkeypatch):
    from backend.aicademiya import teacher_workbench_routes as routes
    def no_parser(*a, **kw):
        raise AssertionError("Rejected input must not enter extraction")
    monkeypatch.setattr(routes,"extract_document",no_parser)
    expect(teacher.post(URL,files={"file":("unsupported.exe",b"not instructions")}),422)
    expect(teacher.post(URL,files={"file":("large.txt",b"x" * (25*1024*1024+1))}),413)
    assert not list(app.state.database.path.parent.glob("instruction-preview-*"))


def test_actual_native_partial_pdf_coverage_notes_and_unit_limit(teacher, tmp_path, monkeypatch):
    monkeypatch.setattr(extraction,"MAX_NATIVE_UNITS",2)
    path = tmp_path / "partial.pdf"
    with pymupdf.open() as doc:
        for i in range(3):
            doc.new_page().insert_text((40,50),f"Native instruction page {i+1}")
        doc.save(path)
    result = expect(post_file(teacher,path))
    assert result["status"] == "needs_review"
    assert result["extraction_report"]["units_total"] == 3 and result["extraction_report"]["units_processed"] == 2
    assert "page 3" not in result["text"] and any("native-unit limit 2" in x for x in result["extraction_report"]["notes"])
    assert "skipped" in result["warning"].lower()


def test_actual_installed_cpu_ocr_instruction_pdf_truthful_review_not_llm(teacher, tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB",(1500,400),"white")
    candidates = [Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf")]
    font_path = next((p for p in candidates if p.exists()), None)
    font = ImageFont.truetype(str(font_path),48) if font_path else ImageFont.load_default(size=48)
    draw = ImageDraw.Draw(image)
    draw.text((40,70),"Instruction preview CPU OCR",fill="black",font=font)
    draw.text((40,150),"Keep student answers unchanged.",fill="black",font=font)
    raster = io.BytesIO()
    image.save(raster,format="PNG")
    path = tmp_path / "scanned-instructions.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page(width=750,height=300)
        page.insert_image(pymupdf.Rect(0,0,750,200),stream=raster.getvalue())
        doc.save(path)
    original = path.read_bytes()
    result = expect(post_file(teacher,path))
    assert_preview(result)
    assert result["status"] == "needs_review" and result["extraction_report"]["ocr_units"] == 1
    assert result["extraction_report"]["method"] == "cpu_ocr"
    assert result["extraction_report"]["visual_review_required"]
    assert "instruction preview" in result["text"].casefold() and "student answers" in result["text"].casefold()
    assert path.read_bytes() == original
    (tmp_path / "instruction-cpu-ocr-verification.json").write_text(json.dumps({"status":result["status"],"text":result["text"],"report":result["extraction_report"],"real_installed_cpu_ocr":True,"llm_calls":0,"original_unchanged":True},indent=2),encoding="utf-8")
