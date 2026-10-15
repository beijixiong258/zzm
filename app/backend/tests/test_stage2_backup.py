"""Synthetic stage2 data relocation; no model, live account or source deletion."""
import importlib.util
from contextlib import closing
import io
import json
from pathlib import Path
import sqlite3
import subprocess
import sys

import pytest
from fastapi.testclient import TestClient
from backend.aicademiya.main import create_app
from .conftest import expect, publish

APP = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("stage2_backup_tool", APP / "tools" / "backup_data.py")
backup_tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(backup_tool)


def worksheet_docx():
    from docx import Document
    import pymupdf
    document = Document()
    document.add_heading("Synthetic backup worksheet", level=0)
    document.add_paragraph("1. Explain the figure using evidence.")
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 12, 12))
    image.clear_with(120)
    document.add_picture(io.BytesIO(image.tobytes("png")))
    document.add_paragraph("Answer: Synthetic private reference.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def commit_payload(preview):
    return {**{k: preview["draft"][k] for k in ("groups", "questions", "reference_answers")}, "mode": "replace"}


def make_legacy_manifest(backup_dir, omitted=()):
    """Use V1's original live-only fingerprints/inventory, no new metadata."""
    path = backup_dir / "manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest["version"] = 1
    manifest.pop("path_bindings", None)
    manifest.pop("directories", None)
    for name in omitted:
        manifest["files"].pop(name)
    manifest["tables"] = backup_tool.inspect_database(backup_dir / backup_tool.DB_NAME, manifest["material_bindings"])
    refresh_database_manifest(backup_dir, manifest)


def refresh_database_manifest(backup_dir, manifest):
    db = backup_dir / backup_tool.DB_NAME
    manifest["files"][backup_tool.DB_NAME] = {"bytes": db.stat().st_size, "sha256": backup_tool.digest(db)}
    (backup_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")


@pytest.fixture
def path_data(tmp_path):
    """A historical-only external original and figure directory, plus staging."""
    root, external = tmp_path / "data", tmp_path / "permitted-materials"
    uploads = root / "uploads"
    uploads.mkdir(parents=True)
    external.mkdir()
    for directory in (uploads / "current-assets", uploads / "empty-assets", external / "historic-assets"):
        directory.mkdir()
    (uploads / "current.txt").write_text("Synthetic current original", encoding="utf-8")
    (uploads / "pending.txt").write_text("Synthetic pending original", encoding="utf-8")
    (external / "anchor.txt").write_text("Synthetic external anchor", encoding="utf-8")
    (external / "historic.txt").write_text("Synthetic historical-only original", encoding="utf-8")
    (uploads / "current-assets" / "figure.png").write_bytes(b"Synthetic current figure bytes")
    (external / "historic-assets" / "figure.png").write_bytes(b"Synthetic historic figure bytes")
    metadata = {"id": "same-material", "assets_dir": str(uploads / "current-assets"), "assets": [{"key": "figure.png"}], "excerpt": "Synthetic metadata"}
    historic = {**metadata, "assets_dir": str(external / "historic-assets")}
    saved = {"course": {"id": "course"}, "sessions": [{"id": "session", "questions": [{"id": "q", "prompt": "Synthetic historic prompt"}], "_materials": [{"id": "same-material", "file_path": str(external / "historic.txt"), "content_json": json.dumps(historic)}]}]}
    staged = {"file_path": str(uploads / "pending.txt"), "material": {"id": "pending", "assets_dir": str(uploads / "empty-assets"), "assets": []}, "draft": {"note": str(root / "keep-literal")}}
    with closing(sqlite3.connect(root / backup_tool.DB_NAME)) as conn, conn:
        conn.executescript("""
        CREATE TABLE materials(id TEXT PRIMARY KEY, file_path TEXT, content_json TEXT);
        CREATE TABLE course_versions(id TEXT PRIMARY KEY, snapshot_json TEXT);
        CREATE TABLE import_staging(id TEXT PRIMARY KEY,kind TEXT,payload_json TEXT);
        CREATE TABLE meta(key TEXT PRIMARY KEY,value TEXT);
        """)
        conn.execute("INSERT INTO materials VALUES('same-material',?,?)", (str(uploads / "current.txt"), json.dumps(metadata)))
        conn.execute("INSERT INTO materials VALUES('anchor',?,'{}')", (str(external / "anchor.txt"),))
        conn.execute("INSERT INTO course_versions VALUES('historic',?)", (json.dumps(saved),))
        conn.execute("INSERT INTO import_staging VALUES('pending','worksheet',?)", (json.dumps(staged),))
        conn.execute("INSERT INTO import_staging VALUES('roster','roster',?)", (json.dumps({"file_path": str(root / "not-a-filesystem-reference")}),))
        conn.execute("INSERT INTO meta VALUES('kv_scope_secret','synthetic-path-fixture-secret')")
    return root, external


def stored_paths(conn):
    paths = []
    for row in conn.execute("SELECT file_path,content_json FROM materials"):
        if row[0]:
            paths.append(row[0])
        raw = json.loads(row[1])
        if raw.get("assets_dir"):
            paths.append(raw["assets_dir"])
    for row in conn.execute("SELECT snapshot_json FROM course_versions"):
        for session in json.loads(row[0])["sessions"]:
            for material in session.get("_materials", []):
                if material.get("file_path"):
                    paths.append(material["file_path"])
                raw = json.loads(material["content_json"])
                if raw.get("assets_dir"):
                    paths.append(raw["assets_dir"])
    for row in conn.execute("SELECT payload_json FROM import_staging WHERE kind='worksheet'"):
        raw = json.loads(row[0])
        paths.extend([raw["file_path"], raw["material"]["assets_dir"]])
    return paths


@pytest.mark.parametrize("format_version", [1, 2])
def test_stage2_backup_relocates_history_assets_and_pending_import(app, teacher, student, admin, tmp_path, catalog, monkeypatch, format_version):
    source = app.state.database.path.parent
    original_docx = worksheet_docx()
    preview = expect(teacher.post("/api/sessions/s4/worksheet-import/preview", files={"file": ("figure.docx", original_docx)}))
    assert preview["method"] == "deterministic" and preview["assets"]
    mid = preview["material"]["id"]
    expect(teacher.post(f"/api/sessions/s4/worksheet-import/{preview['import_id']}/commit", json=commit_payload(preview)))
    expect(teacher.patch(f"/api/sessions/s4/materials/{mid}", json={"visibility": "student"}))
    historic = publish(teacher)["id"]
    work = expect(student.get("/api/sessions/s4/workspace?version_id=" + historic))
    qid = work["session"]["questions"][0]["id"]
    for answer in ("Synthetic initial response", "Synthetic revised response"):
        expect(student.post(f"/api/sessions/s4/answers/{qid}?version_id={historic}", json={"answer": answer}))
    expect(student.post(f"/api/sessions/s4/references/{qid}?version_id={historic}"))
    image_url = work["session"]["questions"][0]["images"][0]["url"]
    asset_bytes = student.get(image_url).content
    assert asset_bytes.startswith(b"\x89PNG")
    # A real publication and content restore produce another immutable edition.
    expect(teacher.patch("/api/sessions/s4", json={"title": "Synthetic changed draft"}))
    changed = publish(teacher)["id"]
    restored_edition = expect(teacher.post(f"/api/teacher/courses/synthetic-course/versions/{historic}/restore", json={}))
    assert restored_edition["id"] not in (historic, changed)
    pending = expect(teacher.post("/api/sessions/s4/worksheet-import/preview", files={"file": ("pending.docx", original_docx)}))
    # Exercise paths that differ for the same material ID in different editions.
    with app.state.database.connect() as conn:
        staged = json.loads(conn.execute("SELECT payload_json FROM import_staging WHERE id=?", (pending["import_id"],)).fetchone()[0])
        raw = json.loads(conn.execute("SELECT content_json FROM materials WHERE id=?", (mid,)).fetchone()[0])
        raw["assets_dir"] = staged["material"]["assets_dir"]
        conn.execute("UPDATE materials SET file_path=?,content_json=? WHERE id=?", (staged["file_path"], json.dumps(raw), mid))
        # Preserve non-worksheet staging and secrets byte-for-byte, not as paths.
        conn.execute("INSERT INTO import_staging VALUES('roster-fixture','demo-teacher','synthetic-course',NULL,'roster','2099-01-01T00:00:00+00:00',?)", (json.dumps({"rows": [], "note": str(source / "not-a-path-field")}),))
        conn.execute("INSERT OR REPLACE INTO settings VALUES('ai_api_key','synthetic-backup-key')")
        conn.execute("INSERT INTO messages(id,session_id,user_id,version_id,role,content,question_id,created_at,mode) VALUES('backup-message','s4','demo-student-1',?,'user','Synthetic saved message',?,'2026-01-01T00:00:00+00:00','fixture')", (historic, qid))
    publish(teacher)
    with app.state.database.connect() as conn:
        original_paths = stored_paths(conn)
        untouched = {table: conn.execute(f'SELECT * FROM "{table}" ORDER BY rowid').fetchall() for table in ("users", "auth_sessions", "settings", "meta", "workspaces", "attempts", "revisions", "messages", "chat_requests", "enrollments")}
        untouched = {k: [tuple(r) for r in rows] for k, rows in untouched.items()}
    backup_dir, restore_dir = tmp_path / "stage2-backup", tmp_path / "stage2-restored"
    backup_tool.backup(source, backup_dir, catalog.parent / "materials")
    if format_version == 1:
        make_legacy_manifest(backup_dir)
    result = backup_tool.restore(backup_dir, restore_dir)
    assert result["records_match"] and result["files_match"]
    with closing(sqlite3.connect(restore_dir / backup_tool.DB_NAME)) as conn, conn:
        paths = stored_paths(conn)
        outside = sum(not Path(p).is_relative_to(restore_dir) for p in paths)
        assert outside == 0, f"{outside} historic/material/staging paths still point outside restore"
        assert all(Path(p).exists() for p in paths)
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            conn.execute("UPDATE course_versions SET label='forbidden'")
        conn.rollback()
        with pytest.raises(sqlite3.IntegrityError, match="immutable"):
            conn.execute("DELETE FROM course_versions")
        conn.rollback()
        for table, rows in untouched.items():
            assert conn.execute(f'SELECT * FROM "{table}" ORDER BY rowid').fetchall() == rows
    with app.state.database.connect() as conn:
        assert stored_paths(conn) == original_paths
    # Make both synthetic source roots unavailable without deleting their bytes.
    hidden_source, hidden_seed = source.with_name("source-unavailable"), catalog.parent.with_name("seed-unavailable")
    source.rename(hidden_source)
    catalog.parent.rename(hidden_seed)
    try:
        monkeypatch.setenv("AICADEMIYA_AI_MODE", "demo")
        restored_app = create_app(data_dir=restore_dir, catalog_path=catalog, demo_enabled=True, frontend_dir=tmp_path / "no-dist")
        # Reuse saved synthetic cookies: this verifies account/session preservation.
        with TestClient(restored_app) as reader:
            reader.cookies.update(student.cookies)
            history = expect(reader.get("/api/history"))
            assert any(r["version"]["id"] == historic for r in history)
            historic_work = expect(reader.get("/api/sessions/s4/workspace?version_id=" + historic))
            assert historic_work["attempts"][qid]["answer"] == "Synthetic revised response"
            assert historic_work["attempts"][qid]["reference_revealed"]
            assert reader.get(image_url).content == asset_bytes
            download = reader.get(f"/api/materials/{mid}/download?version_id={historic}")
            assert download.status_code == 200 and download.content == original_docx
            reader.cookies.clear()
            reader.cookies.update(teacher.cookies)
            pending_asset = reader.get(pending["assets"][0]["url"])
            assert pending_asset.status_code == 200 and pending_asset.content == asset_bytes
            expect(reader.post(f"/api/sessions/s4/worksheet-import/{pending['import_id']}/commit", json=commit_payload(pending)))
            pending_download = reader.get(f"/api/materials/{pending['material']['id']}/download")
            assert pending_download.status_code == 200 and pending_download.content == original_docx
            republished = expect(reader.post(f"/api/teacher/courses/synthetic-course/versions/{historic}/restore", json={}))
            assert republished["id"] != historic
            assert reader.get(image_url).content == asset_bytes
    finally:
        hidden_source.rename(source)
        hidden_seed.rename(catalog.parent)
    manifest_text = (backup_dir / "manifest.json").read_text(encoding="utf-8")
    assert "synthetic-backup-key" not in manifest_text and "Synthetic saved message" not in manifest_text


@pytest.mark.parametrize("legacy,portable", [(False, False), (True, False), (True, True)])
def test_external_history_empty_dirs_and_stdlib_cli(path_data, tmp_path, legacy, portable):
    root, external = path_data
    backup_dir, restored = tmp_path / "backup", tmp_path / "restored"
    # -I -S disables application/site imports; the tool must work with stdlib.
    process = subprocess.run([sys.executable, "-I", "-S", str(APP / "tools" / "backup_data.py"), "backup", "--data-dir", str(root), "--materials-root", str(external), "--destination", str(backup_dir)], capture_output=True, text=True)
    assert process.returncode == 0, process.stderr
    if portable:
        # Labelled foreign-host V1 fixture: paths use POSIX syntax on Windows.
        with closing(sqlite3.connect(backup_dir / backup_tool.DB_NAME)) as conn, conn:
            for table in ("materials", "course_versions", "import_staging"):
                cursor = conn.execute(f'SELECT * FROM "{table}"')
                columns = [c[0] for c in cursor.description]
                def foreign_path(key, value, kind, assets):
                    path = Path(value)
                    if path.is_relative_to(root):
                        return "/unavailable/old-data/" + path.relative_to(root).as_posix()
                    return "/unavailable/old-materials/" + path.relative_to(external).as_posix()
                for row in cursor.fetchall():
                    mapped = backup_tool.map_record_paths(table, columns, row, foreign_path)
                    changed = [(c, v) for c, v, old in zip(columns, mapped, row) if v != old]
                    if changed:
                        fields = ",".join(f'"{c}"=?' for c, _ in changed)
                        conn.execute(f'UPDATE "{table}" SET {fields} WHERE id=?', [v for _, v in changed] + [row[columns.index("id")]])
    if legacy:
        make_legacy_manifest(backup_dir)
    hidden_root, hidden_external = tmp_path / "unavailable-data", tmp_path / "unavailable-materials"
    root.rename(hidden_root)
    external.rename(hidden_external)
    try:
        process = subprocess.run([sys.executable, "-I", "-S", str(APP / "tools" / "backup_data.py"), "restore", "--source", str(backup_dir), "--destination", str(restored)], capture_output=True, text=True)
        assert process.returncode == 0, process.stderr
        assert json.loads(process.stdout)["records_match"] and json.loads(process.stdout)["files_match"]
        with closing(sqlite3.connect(restored / backup_tool.DB_NAME)) as conn, conn:
            assert all(Path(p).is_relative_to(restored) and Path(p).exists() for p in stored_paths(conn))
            staged = json.loads(conn.execute("SELECT payload_json FROM import_staging WHERE id='pending'").fetchone()[0])
            assert staged["draft"]["note"] == str(root / "keep-literal")
            roster = json.loads(conn.execute("SELECT payload_json FROM import_staging WHERE id='roster'").fetchone()[0])
            assert roster["file_path"] == str(root / "not-a-filesystem-reference")
        assert (restored / "uploads" / "empty-assets").is_dir()
        assert (restored / "materials" / "historic.txt").read_bytes() == (hidden_external / "historic.txt").read_bytes()
        assert (restored / "materials" / "historic-assets" / "figure.png").read_bytes() == (hidden_external / "historic-assets" / "figure.png").read_bytes()
        assert "synthetic-path-fixture-secret" not in process.stdout + process.stderr
    finally:
        hidden_root.rename(root)
        hidden_external.rename(external)


@pytest.mark.parametrize("missing", ["materials/historic.txt", "materials/historic-assets/figure.png"])
def test_v1_missing_historical_bytes_refuses_false_match(path_data, tmp_path, missing):
    root, external = path_data
    backup_dir = tmp_path / "backup"
    backup_tool.backup(root, backup_dir, external)
    # The old tool would omit historical-only external originals/assets. Leave
    # the synthetic bytes alone, but exclude them from the V1 saved inventory.
    make_legacy_manifest(backup_dir, omitted=[missing])
    with pytest.raises(ValueError, match="Legacy backup|Unbacked worksheet asset"):
        backup_tool.restore(backup_dir, tmp_path / "restored")
    assert not (tmp_path / "restored").exists()


def corrupt_path(conn, field, wrong):
    if field == "current-file":
        conn.execute("UPDATE materials SET file_path=? WHERE id='same-material'", (wrong,))
    elif field == "current-assets":
        raw = json.loads(conn.execute("SELECT content_json FROM materials WHERE id='same-material'").fetchone()[0])
        raw["assets_dir"] = wrong
        conn.execute("UPDATE materials SET content_json=? WHERE id='same-material'", (json.dumps(raw),))
    elif field.startswith("historic"):
        raw = json.loads(conn.execute("SELECT snapshot_json FROM course_versions").fetchone()[0])
        material = raw["sessions"][0]["_materials"][0]
        if field == "historic-file":
            material["file_path"] = wrong
        else:
            info = json.loads(material["content_json"])
            info["assets_dir"] = wrong
            material["content_json"] = json.dumps(info)
        conn.execute("UPDATE course_versions SET snapshot_json=?", (json.dumps(raw),))
    else:
        raw = json.loads(conn.execute("SELECT payload_json FROM import_staging WHERE id='pending'").fetchone()[0])
        if field == "staged-file":
            raw["file_path"] = wrong
        else:
            raw["material"]["assets_dir"] = wrong
        conn.execute("UPDATE import_staging SET payload_json=? WHERE id='pending'", (json.dumps(raw),))


@pytest.mark.parametrize("field", ["current-file", "current-assets", "historic-file", "historic-assets", "staged-file", "staged-assets"])
@pytest.mark.parametrize("phase", ["backup", "restored"])
def test_verification_checks_each_actual_path_not_just_binding_ids(path_data, tmp_path, field, phase):
    root, external = path_data
    backup_dir, restored = tmp_path / "backup", tmp_path / "restored"
    backup_tool.backup(root, backup_dir, external)
    manifest = backup_tool.validate_backup(backup_dir)
    if phase == "restored":
        backup_tool.restore(backup_dir, restored)
    db = (backup_dir if phase == "backup" else restored) / backup_tool.DB_NAME
    with closing(sqlite3.connect(db)) as conn, conn:
        corrupt_path(conn, field, str(tmp_path / "unverified-original-path"))
    if phase == "backup":
        # Rehash the container to exercise record/path checks, not byte checks.
        refresh_database_manifest(backup_dir, manifest)
        with pytest.raises(ValueError, match="Path record verification"):
            backup_tool.restore(backup_dir, restored)
        assert not restored.exists()
    else:
        with pytest.raises(ValueError, match="Path record verification"):
            backup_tool.inspect_database(db, {}, manifest["path_bindings"], restored)


@pytest.mark.parametrize("field", ["historic-prompt", "staged-draft", "material-excerpt", "meta-secret"])
def test_normalization_keeps_non_path_records_checked(path_data, tmp_path, field):
    root, external = path_data
    backup_dir = tmp_path / "backup"
    backup_tool.backup(root, backup_dir, external)
    manifest = backup_tool.validate_backup(backup_dir)
    with closing(sqlite3.connect(backup_dir / backup_tool.DB_NAME)) as conn, conn:
        if field == "historic-prompt":
            raw = json.loads(conn.execute("SELECT snapshot_json FROM course_versions").fetchone()[0])
            raw["sessions"][0]["questions"][0]["prompt"] = "Changed non-path content"
            conn.execute("UPDATE course_versions SET snapshot_json=?", (json.dumps(raw),))
        elif field == "staged-draft":
            raw = json.loads(conn.execute("SELECT payload_json FROM import_staging WHERE id='pending'").fetchone()[0])
            raw["draft"]["note"] = "Changed staged draft"
            conn.execute("UPDATE import_staging SET payload_json=? WHERE id='pending'", (json.dumps(raw),))
        elif field == "material-excerpt":
            raw = json.loads(conn.execute("SELECT content_json FROM materials WHERE id='same-material'").fetchone()[0])
            raw["excerpt"] = "Changed material excerpt"
            conn.execute("UPDATE materials SET content_json=? WHERE id='same-material'", (json.dumps(raw),))
        else:
            conn.execute("UPDATE meta SET value='Changed synthetic secret'")
    refresh_database_manifest(backup_dir, manifest)
    with pytest.raises(ValueError, match="Backup record verification"):
        backup_tool.restore(backup_dir, tmp_path / "restored")
    assert not (tmp_path / "restored").exists()


@pytest.mark.parametrize("missing", ["historic-original", "historic-figure", "staged-original"])
def test_missing_referenced_bytes_fail_without_manifest(path_data, tmp_path, missing):
    root, external = path_data
    path = {"historic-original": external / "historic.txt", "historic-figure": external / "historic-assets" / "figure.png", "staged-original": root / "uploads" / "pending.txt"}[missing]
    hidden = path.with_name(path.name + ".unavailable")
    path.rename(hidden)
    try:
        with pytest.raises(ValueError, match="regular file|Unbacked worksheet asset"):
            backup_tool.backup(root, tmp_path / "backup", external)
        assert not (tmp_path / "backup" / "manifest.json").exists()
    finally:
        hidden.rename(path)


@pytest.mark.parametrize("staging_only", [False, True])
def test_v1_pre_stage2_and_pending_only_without_live_upload_bindings(tmp_path, staging_only):
    root, backup_dir, restored = tmp_path / "data", tmp_path / "backup", tmp_path / "restored"
    uploads = root / "uploads"
    uploads.mkdir(parents=True)
    (uploads / "original.txt").write_text("Synthetic legacy original", encoding="utf-8")
    (uploads / "empty-assets").mkdir()
    with closing(sqlite3.connect(root / backup_tool.DB_NAME)) as conn, conn:
        conn.execute("CREATE TABLE materials(id TEXT PRIMARY KEY,file_path TEXT)")
        if staging_only:
            conn.execute("CREATE TABLE import_staging(id TEXT PRIMARY KEY,kind TEXT,payload_json TEXT)")
            payload = {"file_path": str(uploads / "original.txt"), "material": {"id": "staged", "assets_dir": str(uploads / "empty-assets"), "assets": []}}
            conn.execute("INSERT INTO import_staging VALUES('staged','worksheet',?)", (json.dumps(payload),))
        else:
            conn.execute("INSERT INTO materials VALUES('legacy',?)", (str(uploads / "original.txt"),))
    backup_tool.backup(root, backup_dir)
    make_legacy_manifest(backup_dir)
    hidden = tmp_path / "unavailable"
    root.rename(hidden)
    try:
        result = backup_tool.restore(backup_dir, restored)
        assert result["records_match"] and result["files_match"]
        with closing(sqlite3.connect(restored / backup_tool.DB_NAME)) as conn:
            if staging_only:
                raw = json.loads(conn.execute("SELECT payload_json FROM import_staging").fetchone()[0])
                assert Path(raw["material"]["assets_dir"]).is_dir()
                path = raw["file_path"]
            else:
                path = conn.execute("SELECT file_path FROM materials").fetchone()[0]
            assert Path(path).is_relative_to(restored) and Path(path).read_text() == "Synthetic legacy original"
    finally:
        hidden.rename(root)


@pytest.mark.parametrize("table", ["materials", "course_versions", "import_staging"])
def test_missing_path_bindings_cannot_hide_unchecked_rows(path_data, tmp_path, table):
    root, external = path_data
    backup_dir = tmp_path / "backup"
    backup_tool.backup(root, backup_dir, external)
    path = backup_dir / "manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest["path_bindings"] = {k: v for k, v in manifest["path_bindings"].items() if json.loads(k)[0] != table}
    path.write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="Unbacked path record"):
        backup_tool.restore(backup_dir, tmp_path / "restored")
    assert not (tmp_path / "restored").exists()
