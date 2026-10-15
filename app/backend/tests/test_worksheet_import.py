"""Real generated documents and explicitly labelled model fixtures.

No account data/provider calls. Optional provided-document acceptance is enabled
only with STAGE2_REAL_WORKSHEET and writes its private assets into tmp_path.
"""
import hashlib
import io
import json
import os
from pathlib import Path
import re

import pytest
import pymupdf

from backend.aicademiya import worksheet_import as wi


TEXT = """Synthetic astronomy worksheet
Task 1. Fill in the blanks
1. A planet orbits a ____ and has ____.
Answer: star; gravity
Word bank: star, gravity
Task 2. Multiple choice
2. Which object gives off light?
a) Star
b) Moon
Answer: a) Star
Task 3. Explain
3. Explain why we can see the Moon.
Use an example in your explanation.
Answer: Reflected sunlight.
"""


def figure_bytes():
    document = pymupdf.open()
    page = document.new_page(width=140, height=100)
    page.draw_rect(pymupdf.Rect(10, 10, 130, 90), color=(0, 0.5, 0))
    page.draw_line((20, 80), (120, 20), color=(0, 0, 1), width=3)
    png = page.get_pixmap().tobytes("png")
    document.close()
    return png


def generate_document(path, figure=True):
    """Generate genuinely parseable, plainly synthetic fixtures."""
    suffix = path.suffix
    if suffix in {".txt", ".md"}:
        path.write_text(TEXT if suffix == ".txt" else "# " + TEXT.replace("Task ", "## Task "), encoding="utf-8")
    elif suffix == ".docx":
        from docx import Document
        from docx.shared import Inches
        document = Document()
        for line in TEXT.splitlines():
            document.add_paragraph(line)
        if figure:
            document.add_picture(io.BytesIO(figure_bytes()), width=Inches(1.5))
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "4. Describe the diagram."
        document.save(path)
    elif suffix == ".pdf":
        document = pymupdf.open()
        page = document.new_page()
        y = 40
        for line in TEXT.splitlines():
            page.insert_text((40, y), line, fontsize=11)
            y += 22
        if figure:
            page.insert_image(pymupdf.Rect(40, y, 180, y + 100), stream=figure_bytes())
        document.save(path)
        document.close()
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(5)).text_frame
        frame.text = TEXT
        if figure:
            slide.shapes.add_picture(io.BytesIO(figure_bytes()), Inches(0.5), Inches(6), width=Inches(1.5))
        document.save(path)
    return path


@pytest.mark.parametrize("suffix", [".txt", ".md", ".docx", ".pdf", ".pptx"])
def test_actual_documents_preserve_questions_options_references_and_source(tmp_path, suffix):
    path = generate_document(tmp_path / ("synthetic_worksheet" + suffix))
    before = path.read_bytes()
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["method"] == "deterministic", result["warnings"]
    assert result["needs_review"] is True
    assert len(result["groups"]) == 3
    assert len(result["questions"]) == (4 if suffix == ".docx" else 3)
    first, second, third = result["questions"][:3]
    assert first["type"] == "fill_in_blank" and first["answer_slots"] == 2
    assert "Word bank: star, gravity" in first["hint"]
    assert second["type"] == "single_choice" and second["options"] == ["a) Star", "b) Moon"]
    assert third["type"] == "open_response" and "Use an example" in third["prompt"]
    assert result["reference_answers"] == {first["id"]: "star; gravity", second["id"]: "a) Star", third["id"]: "Reflected sunlight."}
    assert "Answer: a) Star" in result["source_text"]
    assert "Answer:" not in "\n".join(q["prompt"] for q in result["questions"])
    for question in result["questions"]:
        assert question["source_material_ids"] == question["source_links"] == []
        assert path.name in question["source_label"]
        assert re.fullmatch(r"[A-Za-z0-9_.-]+", question["id"])
    assert before == path.read_bytes()
    if suffix in {".docx", ".pdf", ".pptx"}:
        assert len(result["assets"]) == 1
        asset = result["assets"][0]
        assert re.fullmatch(r"figure_[0-9a-f]{24}\.png", asset["key"])
        assert (tmp_path / "assets" / asset["key"]).read_bytes().startswith(b"\x89PNG")
        assert all("url" not in image for q in result["questions"] for image in q["images"])
        assert result["questions"][2]["images"]


def test_word_automatic_numbering_and_bullet_options(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("Synthetic automatic numbering")
    document.add_paragraph("Task 1. Select an answer")
    document.add_paragraph("Which number is even?", style="List Number")
    document.add_paragraph("a) Two", style="List Bullet")
    document.add_paragraph("b) Three", style="List Bullet")
    document.add_paragraph("Explain the choice.", style="List Number")
    path = tmp_path / "synthetic_numbering.docx"
    document.save(path)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 2
    assert result["questions"][0]["options"] == ["a) Two", "b) Three"]
    assert "1. Which" in result["source_text"] and "2. Explain" in result["source_text"]


def test_matching_figures_and_explicit_labels(tmp_path):
    from docx import Document
    from docx.shared import Inches
    document = Document()
    document.add_paragraph("Synthetic matching worksheet")
    document.add_paragraph("Task 1. Matching")
    document.add_paragraph("1. Rising line")
    document.add_paragraph("2. Falling line")
    document.add_picture(io.BytesIO(figure_bytes()), width=Inches(1))
    other = pymupdf.open(stream=figure_bytes(), filetype="png")
    other_pdf = other.convert_to_pdf()
    other.close()
    other = pymupdf.open(stream=other_pdf, filetype="pdf")
    other[0].insert_text((20, 40), "B")
    document.add_picture(io.BytesIO(other[0].get_pixmap().tobytes("png")), width=Inches(1))
    other.close()
    document.add_paragraph("A          B")
    path = tmp_path / "synthetic_matching.docx"
    document.save(path)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 2 and len(result["assets"]) == 2
    for q in result["questions"]:
        assert q["type"] == "image_matching" and q["options"] == ["A", "B"]
        assert [image["label"] for image in q["images"]] == ["A", "B"]
        assert "A          B" not in q["prompt"]


@pytest.mark.parametrize("text", ["", "Synthetic background reading.\nThis is not a question.", "Notes: what does the author think?\nThis is background, not a numbered task."])
def test_no_fake_question_success(tmp_path, text):
    path = tmp_path / "synthetic_background.txt"
    path.write_text(text, encoding="utf-8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["questions"] == [] and result["needs_review"]
    assert any("No unambiguous questions" in warning for warning in result["warnings"])
    assert text.splitlines()[0] in result["source_text"] if text else not result["source_text"]


def test_answer_key_ambiguity_and_reference_privacy(tmp_path):
    path = tmp_path / "synthetic_key.txt"
    path.write_text("Task 1. Explain\n1. Explain water.\nTask 2. Explain\n1. Explain light.\nAnswer Key\n1. PRIVATE_AMBIGUOUS_KEY", encoding="utf-8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 2 and result["reference_answers"] == {}
    assert "PRIVATE_AMBIGUOUS_KEY" in result["source_text"]
    assert all("PRIVATE" not in q["prompt"] for q in result["questions"])
    assert any("mapped uniquely" in warning for warning in result["warnings"])


def test_numbered_answer_key_mapping(tmp_path):
    path = tmp_path / "synthetic_numbered_key.txt"
    path.write_text("1. Explain light.\n2. Explain water.\nAnswer Key\n1. Energy.\n2. Molecules.", encoding="utf-8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert list(result["reference_answers"].values()) == ["Energy.", "Molecules."]


@pytest.mark.parametrize("suffix", [".exe", ".pdf", ".docx"])
def test_unsupported_or_corrupt(tmp_path, suffix):
    path = tmp_path / ("synthetic_invalid" + suffix)
    path.write_bytes(b"NOT_A_VALID_DOCUMENT")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["method"] in {"unsupported", "error"}
    assert not result["questions"] and result["warnings"] and result["needs_review"]


def test_utf16_and_invalid_encoding(tmp_path):
    path = tmp_path / "synthetic_utf16.txt"
    path.write_text("1. Explain gravity.", encoding="utf-16")
    assert len(wi.preview_worksheet(path, assets_dir=tmp_path / "assets")["questions"]) == 1
    path.write_bytes(b"\xffNOT_UTF8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["method"] == "error" and result["warnings"]


def scan_fixture(path):
    source = pymupdf.open()
    page = source.new_page()
    for y, text in ((60, "Synthetic scanned worksheet"), (110, "Task 1. Explain"), (165, "1. Explain why plants need light."), (220, "2. Describe a healthy leaf.")):
        page.insert_text((40, y), text, fontsize=18)
    image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
    source.close()
    scanned = pymupdf.open()
    page = scanned.new_page()
    page.insert_image(page.rect, stream=image)
    scanned.save(path)
    scanned.close()
    return path


def test_actual_cpu_ocr_scan_no_network_no_public_page_asset(tmp_path, monkeypatch):
    pytest.importorskip("rapidocr")
    pytest.importorskip("onnxruntime")
    # This is real OCR inference, not a mocked recognition result. Network access
    # is patched to fail to prove that installed model files suffice.
    from rapidocr.utils.download_file import DownloadFile
    def no_download(*args, **kwargs):
        raise AssertionError("OCR must not download any model/dictionary")
    monkeypatch.setattr(DownloadFile, "run", no_download)
    wi._local_ocr_engine.cache_clear()
    path = scan_fixture(tmp_path / "synthetic_scan.pdf")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["method"] == "deterministic+cpu_ocr", result["warnings"]
    assert len(result["questions"]) == 2, result["source_text"]
    assert "plants need light" in result["questions"][0]["prompt"].lower()
    assert "healthy leaf" in result["questions"][1]["prompt"].lower()
    assert result["assets"] == [] and result["needs_review"]
    assert "CPU OCR" in result["questions"][0]["source_label"]
    engine = wi._local_ocr_engine()
    for stage in (engine.text_det, engine.text_cls, engine.text_rec):
        assert stage.session.session.get_providers() == ["CPUExecutionProvider"]


def test_ocr_unavailable_is_truthful_fixture(tmp_path, monkeypatch):
    # Explicit failure fixture, not a claim about the installed engine.
    def unavailable(png):
        raise ImportError("Labelled missing-local-engine fixture")
    monkeypatch.setattr(wi, "_ocr_lines", unavailable)
    result = wi.preview_worksheet(scan_fixture(tmp_path / "synthetic_scan.pdf"), assets_dir=tmp_path / "assets")
    assert not result["questions"] and not result["assets"]
    assert any("unavailable" in warning for warning in result["warnings"])


def test_pdf_full_page_image_and_vector_warning(tmp_path):
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, stream=figure_bytes(), keep_proportion=False)
    page.insert_text((40, 50), "1. Explain the figure and identify the evidence that supports your conclusion.")
    page.draw_line((20, 20), (100, 100))
    path = tmp_path / "synthetic_large_figure.pdf"
    document.save(path)
    document.close()
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 1 and result["assets"] == []
    assert any("Page-sized" in warning for warning in result["warnings"])
    assert any("Vector drawings" in warning for warning in result["warnings"])


def ai_draft():
    return {"title": "Labelled callable fixture", "groups": [{"id": "g", "title": "Synthetic core", "kind": "core"}], "questions": [{"id": "q", "group_id": "g", "type": "open_response", "prompt": "Explain a synthetic concept."}], "reference_answers": {"q": "Explicit synthetic reference"}, "warnings": ["Verify the inferred grouping."]}


def test_injected_callable_normalization_no_model_service(tmp_path):
    path = generate_document(tmp_path / "synthetic_ai.txt")
    calls = []
    def labelled_fixture(messages):
        calls.append(messages)
        return "```json\n" + json.dumps(ai_draft()) + "\n```"
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets", llm=labelled_fixture)
    assert len(calls) == 1 and calls[0][0]["role"] == "system"
    assert "source_text" in calls[0][1]["content"]
    assert result["method"] == "llm" and len(result["questions"]) == 1
    assert result["questions"][0]["ordinal"] == 1
    assert result["questions"][0]["source_material_ids"] == []
    assert "synthetic_ai.txt" in result["questions"][0]["source_label"]
    assert result["reference_answers"] == {"q": "Explicit synthetic reference"}
    assert any("possible omissions" in warning for warning in result["warnings"])
    assert any("fewer options" in warning for warning in result["warnings"])
    assert any("Verify the inferred grouping" in warning for warning in result["warnings"])
    assert "Optional structured extractor response" in result["source_text"]
    assert result["needs_review"]


@pytest.mark.parametrize("mutate", [
    lambda d: d.update(questions=[]),
    lambda d: d["questions"].append(d["questions"][0].copy()),
    lambda d: d["questions"][0].update(group_id="unknown"),
    lambda d: d["questions"][0].update(type="invented_type"),
    lambda d: d["questions"][0].update(prompt=" "),
    lambda d: d["questions"][0].update(options=[{"text": "unsupported"}]),
    lambda d: d["questions"][0].update(images=[{"asset_key": "../private.png"}]),
    lambda d: d["questions"][0].update(images=[{"url": "https://example.org/figure.png"}]),
    lambda d: d.update(reference_answers={"unknown": "Not mapped"}),
    lambda d: d["questions"][0].update(answer_slots=True),
])
def test_invalid_callable_fixture_keeps_deterministic_draft(tmp_path, mutate):
    path = generate_document(tmp_path / "synthetic_invalid_ai.txt")
    draft = ai_draft()
    mutate(draft)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets", llm=lambda messages: json.dumps(draft))
    assert result["method"] == "deterministic" and len(result["questions"]) == 3
    assert any("invalid draft" in warning for warning in result["warnings"])


def test_ai_fixture_preserves_known_assets_and_unknown_content(tmp_path):
    path = generate_document(tmp_path / "synthetic_figures.docx")
    def fixture(messages):
        source = json.loads(messages[1]["content"])
        draft = ai_draft()
        draft["extra_teaching_context"] = "SYNTHETIC_CONTEXT_KEEP"
        draft["questions"][0].update(images=[{"asset_key": source["assets"][0]["key"], "label": "A", "alt": "Line diagram"}], source_links=["https://example.org"])
        return json.dumps(draft)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets", llm=fixture)
    assert result["method"] == "llm" and result["questions"][0]["images"][0]["label"] == "A"
    assert result["questions"][0]["source_links"] == []
    assert "SYNTHETIC_CONTEXT_KEEP" in result["source_text"]
    assert any("Unsupported structured fields" in warning for warning in result["warnings"])


def test_ai_exception_is_not_fake_model_success(tmp_path):
    path = generate_document(tmp_path / "synthetic_failure.txt")
    def failure(messages):
        raise RuntimeError("Private fixture details must not appear")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets", llm=failure)
    assert result["method"] == "deterministic" and len(result["questions"]) == 3
    assert "Private fixture details" not in str(result)


def test_bounded_ai_source_not_silently_truncated(tmp_path, monkeypatch):
    path = generate_document(tmp_path / "synthetic_bounded.txt")
    monkeypatch.setattr(wi, "MAX_AI_TEXT", 20)
    def forbidden(messages):
        raise AssertionError("Over-limit payload must not be sent")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets", llm=forbidden)
    assert len(result["questions"]) == 3
    assert any("payload limit" in warning for warning in result["warnings"])
    assert "Reflected sunlight" in result["source_text"]


def test_size_limits_are_visible(tmp_path, monkeypatch):
    path = generate_document(tmp_path / "synthetic_limits.txt")
    monkeypatch.setattr(wi, "MAX_BYTES", 10)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert result["method"] == "error" and not result["questions"]
    assert any("size limit" in warning for warning in result["warnings"])


def test_assets_directory_preserves_existing_files_and_safe_dedup(tmp_path):
    path = generate_document(tmp_path / "synthetic_preservation.docx")
    assets_dir = tmp_path / "assets"
    assets_dir.mkdir()
    keep = assets_dir / "caller_owned.txt"
    keep.write_text("Keep this caller asset", encoding="utf-8")
    first = wi.preview_worksheet(path, assets_dir=assets_dir)
    second = wi.preview_worksheet(path, assets_dir=assets_dir)
    assert keep.read_text(encoding="utf-8") == "Keep this caller asset"
    assert first["assets"] == second["assets"] and first["questions"] == second["questions"]


def test_multiline_reference_does_not_become_student_prompt(tmp_path):
    path = tmp_path / "synthetic_multiline_key.txt"
    path.write_text("Task 1. Explain\n1. Explain gravity.\nAnswer: PRIVATE_FIRST_LINE\nPRIVATE_SECOND_LINE\nExplain the private reasoning here.\n2. Explain light.", encoding="utf-8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 2
    answer = result["reference_answers"][result["questions"][0]["id"]]
    assert answer == "PRIVATE_FIRST_LINE\nPRIVATE_SECOND_LINE\nExplain the private reasoning here."
    assert "PRIVATE" not in " ".join(q["prompt"] for q in result["questions"])


@pytest.mark.parametrize("suffix", [".md", ".docx"])
def test_generic_section_headings_are_editable_groups(tmp_path, suffix):
    path = tmp_path / ("synthetic_sections" + suffix)
    if suffix == ".md":
        path.write_text("# Synthetic title\n## Reading discussion\n1. Explain the author.\n## Reflection\n2. Describe your thinking.", encoding="utf-8")
    else:
        from docx import Document
        document = Document()
        document.add_paragraph("Synthetic title", style="Title")
        document.add_paragraph("Reading discussion", style="Heading 1")
        document.add_paragraph("1. Explain the author.")
        document.add_paragraph("Reflection", style="Heading 1")
        document.add_paragraph("2. Describe your thinking.")
        document.save(path)
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert [g["title"] for g in result["groups"]] == ["Reading discussion", "Reflection"]
    assert [q["type"] for q in result["questions"]] == ["reading_response", "open_response"]


def test_existing_asset_collision_is_preserved_and_not_referenced(tmp_path):
    path = generate_document(tmp_path / "synthetic_asset_conflict.docx")
    assets_dir = tmp_path / "assets"
    first = wi.preview_worksheet(path, assets_dir=assets_dir)
    asset = assets_dir / first["assets"][0]["key"]
    asset.write_bytes(b"CALLER_FILE_MUST_NOT_BE_OVERWRITTEN")
    second = wi.preview_worksheet(path, assets_dir=assets_dir)
    assert asset.read_bytes() == b"CALLER_FILE_MUST_NOT_BE_OVERWRITTEN"
    assert second["assets"] == []
    assert not any(q["images"] for q in second["questions"])
    assert any("conflicts" in warning for warning in second["warnings"])


def test_actual_scan_with_sparse_overlay_uses_cpu_ocr(tmp_path):
    pytest.importorskip("rapidocr")
    path = scan_fixture(tmp_path / "synthetic_watermark_scan.pdf")
    document = pymupdf.open(path)
    document[0].insert_text((40, 800), "PAGE 1")
    overlaid = tmp_path / "synthetic_watermark_overlay.pdf"
    document.save(overlaid)
    document.close()
    result = wi.preview_worksheet(overlaid, assets_dir=tmp_path / "assets")
    assert result["method"] == "deterministic+cpu_ocr"
    assert len(result["questions"]) == 2, result["source_text"]
    assert "PAGE 1" in result["source_text"] and result["assets"] == []


@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
def test_multiline_numbered_prompt_not_split_into_fake_questions(tmp_path, suffix):
    path = tmp_path / ("synthetic_multiline_prompt" + suffix)
    text = "1. Explain gravity.\nGive an example of its effects.\nDescribe the forces in your example."
    if suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_paragraph(text)
        document.save(path)
    else:
        document = pymupdf.open()
        document.new_page().insert_textbox(pymupdf.Rect(40, 40, 500, 300), text)
        document.save(path)
        document.close()
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 1, result["questions"]
    assert result["questions"][0]["prompt"] == text[3:]


def test_markdown_figure_paths_retained_not_followed(tmp_path):
    path = tmp_path / "synthetic_image_reference.md"
    path.write_text("1. Explain this figure.\n![Synthetic](http://127.0.0.1/private.png)", encoding="utf-8")
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "assets")
    assert len(result["questions"]) == 1 and result["assets"] == []
    assert "http://127.0.0.1/private.png" in result["source_text"]
    assert "http://127.0.0.1" not in result["questions"][0]["prompt"]
    assert "manual asset assignment" in result["questions"][0]["prompt"]
    assert any("not fetched" in warning for warning in result["warnings"])


@pytest.mark.skipif(not os.environ.get("STAGE2_REAL_WORKSHEET"), reason="Provided worksheet acceptance is explicitly opt-in/read-only")
def test_provided_student_worksheet_read_only(tmp_path):
    path = Path(os.environ["STAGE2_REAL_WORKSHEET"])
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = wi.preview_worksheet(path, assets_dir=tmp_path / "private_real_assets")
    assert result["method"] == "deterministic"
    assert len(result["groups"]) == 4 and len(result["questions"]) == 22
    from collections import Counter
    assert Counter(q["type"] for q in result["questions"]) == {"fill_in_blank": 6, "single_choice": 5, "image_matching": 3, "open_response": 8}
    assert sum(q["answer_slots"] for q in result["questions"] if q["type"] == "fill_in_blank") == 8
    assert all(len(q["options"]) == 4 for q in result["questions"] if q["type"] == "single_choice")
    assert len(result["assets"]) == 3
    assert all(len(q["images"]) == 3 and q["options"] == ["A", "B", "C"] for q in result["questions"] if q["type"] == "image_matching")
    assert result["reference_answers"] == {} and result["needs_review"]
    assert before == hashlib.sha256(path.read_bytes()).hexdigest()
