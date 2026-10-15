"""Synthetic Office bundles verify app transport/auth/revision, not rendering/OCR."""
import io,json,shutil,zipfile
from pathlib import Path
import pymupdf,pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from docx import Document
from backend.aicademiya.main import create_app
from backend.aicademiya.document_bundle import source_identity,sha256_file,validate_bundle,BundleError
from backend.aicademiya.reading_context import execute_plan
from backend.aicademiya.guided_reading import SourceFeedback,source_citations
from .test_stage12_document_api import synthetic_bundle,pdf_bytes,wait_job,endpoint
from .conftest import expect,publish

def office_bytes(kind='pptx'):
 stream=io.BytesIO()
 if kind=='pptx':
  p=Presentation()
  for n in range(2):p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(0,0,1000000,1000000).text=f'SYNTHETIC slide {n+1}'
  p.save(stream)
 elif kind=='docx':
  d=Document();d.add_paragraph('SYNTHETIC Word text');d.save(stream)
 else:return bytes.fromhex('d0cf11e0a1b11ae1')+b'\0'*504
 return stream.getvalue()

def office_bundle(source,root,wrong_format=False):
 rendered=root.parent/(root.name+'-SYNTHETIC.pdf');rendered.write_bytes(pdf_bytes(2));value=synthetic_bundle(rendered,root)
 shutil.copy2(rendered,root/'preview.pdf');value.update(source_sha256=sha256_file(source),source_format='pdf'if wrong_format else source.suffix[1:],preview_pdf='preview.pdf',coverage={'synthetic':True,'source_units':2})
 (root/'manifest.json').write_text(json.dumps(value));return value

@pytest.fixture
def env(tmp_path,catalog,monkeypatch):
 monkeypatch.setenv('AICADEMIYA_DOCUMENT_FORMATS','pdf,pptx,docx,doc');monkeypatch.setenv('AICADEMIYA_AI_MODE','demo')
 app=create_app(data_dir=tmp_path/'data',catalog_path=catalog,demo_enabled=True,frontend_dir=tmp_path/'dist',document_processor=lambda s,r,c:office_bundle(s,r))
 with TestClient(app)as teacher:
  expect(teacher.post('/api/auth/demo',json={'role':'teacher'}))
  yield app,teacher

def upload_office(teacher,kind='pptx',content=None,processing='full_document'):
 return teacher.post('/api/sessions/s4/materials',data={'processing':processing},files={'file':('synthetic.'+kind,office_bytes(kind)if content is None else content,'application/octet-stream')})

@pytest.mark.parametrize('kind',['pptx','docx','doc'])
def test_source_identity_and_original_format(tmp_path,kind):
 p=tmp_path/('source.'+kind);p.write_bytes(office_bytes(kind));digest,count=source_identity(p)
 assert digest==sha256_file(p);assert count==(2 if kind=='pptx'else None)
 if kind!='doc':
  bad=tmp_path/('wrong.'+('docx'if kind=='pptx'else'pptx'));bad.write_bytes(p.read_bytes())
  with pytest.raises(BundleError):source_identity(bad)

@pytest.mark.parametrize('kind',['pptx','docx','doc'])
def test_office_upload_complete_derived_pdf_and_original_preserved(env,kind):
 app,teacher=env;content=office_bytes(kind);m=expect(upload_office(teacher,kind,content));mid=m['id'];job=wait_job(teacher,mid);assert job['status']=='completed'
 assert job['page_count']==job['processed_pages']==2
 summary=expect(teacher.get(endpoint(mid)));info=summary['document'];assert info['source_format']==kind and info['has_rendered_pdf']
 rid=info['revision_id'];url=endpoint(mid)+'/pdf';r=teacher.get(url,params={'revision_id':rid});assert r.status_code==200 and r.headers['content-type']=='application/pdf'and r.headers['content-disposition']=='inline'
 assert teacher.get(m['download_url']).content==content
 expect(teacher.get(endpoint(mid)+'/assets/preview.pdf',params={'revision_id':rid}),404)
 expect(teacher.get(url,params={'revision_id':'wrong'}),409)
 with TestClient(app)as anonymous:expect(anonymous.get(url,params={'revision_id':rid}),401)
 with TestClient(app)as admin:
  expect(admin.post('/api/auth/demo',json={'role':'admin'}));expect(admin.get(url,params={'revision_id':rid}),403)
 version=publish(teacher);vid=version['id']
 with TestClient(app)as student:
  expect(student.post('/api/auth/demo',json={'role':'student'}));r=student.get(url,params={'revision_id':rid,'version_id':vid});assert r.status_code==200
  result=expect(teacher.post(endpoint(mid)+'/reviews',json={'request_id':'new-review','expected_revision':rid,'changes':[{'block_id':'p1-b0','quality':'approved','reason':'Synthetic review','text':'Synthetic corrected text'}]}));new=result['document']['revision_id']
  expect(student.get(url,params={'revision_id':new,'version_id':vid}),409)
  assert student.get(url,params={'revision_id':rid,'version_id':vid}).content==r.content

@pytest.mark.parametrize('bad',['missing','png','count'])
def test_rendered_pdf_must_be_real_complete_pdf(tmp_path,bad):
 source=tmp_path/'source.pptx';source.write_bytes(office_bytes());root=tmp_path/'bundle';office_bundle(source,root)
 if bad=='missing':(root/'preview.pdf').rename(root/'retained.pdf')
 elif bad=='png':shutil.copy2(root/'pages/p1.png',root/'preview.pdf')
 else:(root/'preview.pdf').write_bytes(pdf_bytes(1))
 with pytest.raises(BundleError):validate_bundle(root,sha256_file(source),2)

def test_configured_formats_are_not_silently_enabled(env):
 app,teacher=env;app.state.document_queue.formats=('pdf',)
 expect(upload_office(teacher),415)
 assert expect(teacher.get('/api/document-processing'))['formats']==['pdf']

def test_large_office_full_accepted_but_preview_limit_retained(env):
 app,teacher=env;data=io.BytesIO(office_bytes())
 with zipfile.ZipFile(data,'a',compression=zipfile.ZIP_STORED)as z:z.writestr('ppt/media/synthetic-padding.bin',b'\0'*(26*1024*1024))
 content=data.getvalue();assert len(content)>25*1024*1024
 expect(upload_office(teacher,content=content,processing='preview'),413)
 m=expect(upload_office(teacher,content=content));assert wait_job(teacher,m['id'])['status']=='completed'

def test_wrong_worker_format_cannot_commit(env):
 app,teacher=env;app.state.document_queue.processor=lambda s,r,c:office_bundle(s,r,True)
 m=expect(upload_office(teacher));assert wait_job(teacher,m['id'])['status']=='failed'
 assert expect(teacher.get(endpoint(m['id'])))['document']is None

def test_office_citation_uses_exact_revision_bound_pdf(tmp_path):
 source=tmp_path/'s.pptx';source.write_bytes(office_bytes());manifest=office_bundle(source,tmp_path/'bundle')
 d={'material_id':'mid','revision_id':'rid','title':'Source slides','kind':'reading','visibility':'student','manifest':manifest}
 evidence=execute_plan([d],{'requests':[{'material_id':'mid','pages':[1],'query':'','block_ids':[]}]})
 q=evidence['blocks'][0];answer=SourceFeedback(text='Synthetic answer',document_citations=[{'ref':q['ref'],'quote':q['text']}])
 c=source_citations(answer,evidence,{'_version_id':'edition'})[0]
 assert c['url']=='/api/materials/mid/document/pdf?version_id=edition&revision_id=rid#page=1'
 assert '_rendered_pdf'not in c
