"""Core authoring integration with pure helpers; no live link/model requests."""
import io
import json
from pathlib import Path
import httpx
import pytest
from backend.aicademiya import link_materials
from .conftest import expect, publish, question, demo_client

VERSIONS = "/api/teacher/courses/synthetic-course/versions"


@pytest.fixture
def public_link_fixture(monkeypatch):
    calls = []
    def extraction(url, **kwargs):
        calls.append(url)
        return {"title": "Labelled protocol link fixture", "url": url, "excerpt": "LABELLED_PUBLIC_LINK_EVIDENCE reliability and assessment",
                "status": "needs_review", "warning": "Labelled protocol fixture: not a real website retrieval"}
    monkeypatch.setattr(link_materials, "extract_link", extraction)
    return calls


def editable_question(teacher):
    q = question("q1"); q.pop("teacher_internal")
    return q


def test_link_sources_shared_dependencies_historical_retention_and_manual_links(teacher, student, public_link_fixture):
    q1, q2 = editable_question(teacher), editable_question(teacher)
    q2.update(id="q2", ordinal=2)
    q1["source_links"] = ["10.1234/Exact-DOI"]
    q2["source_links"] = ["https://doi.org/10.1234/Exact-DOI"]
    draft = expect(teacher.patch("/api/sessions/s4", json={"questions": [q1, q2]}))
    link = next(m for m in draft["materials"] if m["kind"] == "link")
    mid = link["id"]
    assert link["url"] == "https://doi.org/10.1234/Exact-DOI"
    assert public_link_fixture == [link["url"]]
    assert all(q["source_material_ids"] == [] and q["source_links"] == [link["url"]] for q in draft["questions"])
    assert draft["questions"][0]["source_label"] == "Synthetic worksheet q1"
    expect(teacher.delete("/api/sessions/s4/materials/"+mid), 409)
    q1["source_links"] = []
    updated = expect(teacher.patch("/api/sessions/s4", json={"questions": [q1, q2]}))
    assert any(m["id"] == mid for m in updated["materials"])
    version = publish(teacher)["id"]
    expect(student.get("/api/sessions/s4/workspace?version_id="+version))
    q2["source_links"] = []
    updated = expect(teacher.patch("/api/sessions/s4", json={"questions": [q1, q2]}))
    assert any(m["id"] == mid for m in updated["materials"])  # Immutable snapshot still refers to it.
    manual = expect(teacher.post("/api/sessions/s4/materials/link", json={"url": "https://example.org/manual", "visibility": "teacher", "title": "Independent manual Link"}))
    expect(teacher.patch("/api/sessions/s4", json={"questions": [q1, q2]}))
    assert any(m["id"] == manual["id"] for m in expect(teacher.get("/api/sessions/s4"))["materials"])
    expect(teacher.delete("/api/sessions/s4/materials/"+mid))
    assert "LABELLED_PUBLIC_LINK_EVIDENCE" in teacher.get(link["preview_url"]).text
    assert all(m["id"] != mid for m in expect(teacher.get("/api/sessions/s4"))["materials"])
    historical_preview = student.get(f"/api/materials/{mid}/preview?version_id="+version)
    assert historical_preview.status_code == 200 and "LABELLED_PUBLIC_LINK_EVIDENCE" in historical_preview.text
    expect(student.get(f"/api/materials/{manual['id']}/preview?version_id="+version), 404)


def test_unpublished_managed_link_is_detached_but_retained_when_unreferenced(teacher, public_link_fixture):
    q = editable_question(teacher); q["source_links"] = ["https://example.org/new-source"]
    draft = expect(teacher.patch("/api/sessions/s4", json={"questions": [q]}))
    mid = next(m["id"] for m in draft["materials"] if m["kind"] == "link")
    q["source_links"] = []
    updated = expect(teacher.patch("/api/sessions/s4", json={"questions": [q]}))
    assert all(m["id"] != mid for m in updated["materials"])
    retained = next(m for m in expect(teacher.get("/api/courses/synthetic-course/materials"))["materials"] if m["id"] == mid)
    assert retained["session_ids"] == []


@pytest.mark.parametrize("value", ["10.123/TooShort", "10.1234/has a space", "doi:invalid", "http://127.0.0.1/private", "http://localhost/private", "https://name:secret@example.org/", "javascript:alert(1)"])
def test_bad_identifiers_and_private_links_rejected_before_any_fetch(teacher, public_link_fixture, value):
    q = editable_question(teacher); q["source_links"] = [value]
    expect(teacher.patch("/api/sessions/s4", json={"questions": [q]}), 422)
    expect(teacher.post("/api/sessions/s4/materials/link", json={"url": value, "visibility": "student"}), 422)
    assert not public_link_fixture


def test_source_material_session_isolation_and_group_editing(teacher, student):
    groups = [{"id": "core-new", "title": "Edited core", "kind": "core"}, {"id": "custom-reading-uuid", "title": "Arbitrary branch", "kind": "reading"}]
    q = editable_question(teacher); q.update(group_id="custom-reading-uuid", source_material_ids=["s2-public"])
    expect(teacher.patch("/api/sessions/s4", json={"groups": groups, "questions": [q]}), 422)
    q["source_material_ids"] = []
    draft = expect(teacher.patch("/api/sessions/s4", json={"groups": groups, "questions": [q]}))
    assert draft["groups"] == groups and draft["questions"][0]["source_label"]
    version = publish(teacher)["id"]
    expect(student.post("/api/sessions/s4/reading?version_id="+version, json={"group_id": "custom-reading-uuid"}))
    assert expect(student.post("/api/sessions/s4/answers/q1?version_id="+version, json={"answer": "Arbitrary stable group works"}))["progress"]["answered"] == 1
    expect(teacher.patch("/api/sessions/s4", json={"groups": groups+groups}), 422)
    expect(teacher.patch("/api/sessions/s4", json={"questions": [{**q, "prompt": "   "}]}), 422)
    expect(teacher.patch("/api/sessions/s4", json={"questions": [{**q, "group_id": "missing"}]}), 422)
    expect(teacher.patch("/api/sessions/s4", json={"questions": [{**q, "images": [{"url": "https://example.org/tracker.png"}]}]}), 422)


def test_material_metadata_preview_diff_and_historical_authorization(teacher, student, public_link_fixture):
    first = expect(teacher.post("/api/sessions/s4/materials", files={"file": ("preview.txt", b"<script>safe text only</script> synthetic evidence")}))
    version = publish(teacher)["id"]
    expect(student.get("/api/sessions/s4/workspace?version_id="+version))
    preview = student.get(f"/api/materials/{first['id']}/preview?version_id="+version)
    assert preview.status_code == 200 and preview.headers["content-type"].startswith("text/plain")
    assert preview.headers["x-content-type-options"] == "nosniff"
    changed = expect(teacher.patch("/api/sessions/s4/materials/"+first["id"], json={"title": "Edited material", "visibility": "teacher"}))
    assert changed["title"] == "Edited material" and changed["visibility"] == "teacher"
    diffs = expect(teacher.get(VERSIONS+"/"+version+"/diff"))["changes"]
    assert any(d["path"].endswith(".metadata.title") and d["before"] == "preview.txt" and d["after"] == "Edited material" for d in diffs)
    new = publish(teacher)["id"]
    expect(student.get(f"/api/materials/{first['id']}/preview?version_id="+new), 403)
    assert student.get(f"/api/materials/{first['id']}/download?version_id="+version).content.startswith(b"<script>")
    link = expect(teacher.post("/api/sessions/s4/materials/link", json={"url": "https://example.org/old", "visibility": "student"}))
    link = expect(teacher.patch("/api/sessions/s4/materials/"+link["id"], json={"url": "10.9999/Exact", "title": "Changed DOI"}))
    assert link["url"] == "https://doi.org/10.9999/Exact" and link["warning"]
    expect(teacher.patch("/api/sessions/s4/materials/"+first["id"], json={"url": "https://example.org/not-link"}), 422)


def worksheet_docx():
    from docx import Document
    import pymupdf
    document = Document()
    document.add_heading("Synthetic reviewed worksheet", level=0)
    document.add_paragraph("1. Explain the figure using evidence.")
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 12, 12)); image.clear_with(120)
    document.add_picture(io.BytesIO(image.tobytes("png")))
    document.add_paragraph("Answer: Synthetic teacher-only reference for the figure.")
    document.add_paragraph("2. Explain formative reliability.")
    output = io.BytesIO(); document.save(output)
    return output.getvalue()


def test_real_worksheet_preview_review_commit_assets_private_original_and_versions(app, teacher, student):
    initial = expect(student.get("/api/sessions/s4/workspace"))["version"]["id"]
    response = teacher.post("/api/sessions/s4/worksheet-import/preview", files={"file": ("synthetic-worksheet.docx", worksheet_docx())})
    preview = expect(response)
    assert preview["method"] == "deterministic" and preview["needs_review"] and len(preview["draft"]["questions"]) == 2 and preview["assets"]
    asset = preview["assets"][0]
    mid = preview["material"]["id"]
    assert teacher.get(asset["url"]).content.startswith(b"\x89PNG")
    expect(student.get(asset["url"]), 404)
    assert preview["material"]["download_url"] is None and preview["material"]["preview_url"] is None
    assert "teacher-only reference" not in student.get("/api/sessions/s4/workspace").text
    payload = {k: preview["draft"][k] for k in ("groups", "questions", "reference_answers")}
    payload["mode"] = "replace"
    bad = json.loads(json.dumps(payload)); bad["questions"][0]["images"] = [{"asset_key": "../../private.png", "label": "Bad"}]
    expect(teacher.post(f"/api/sessions/s4/worksheet-import/{preview['import_id']}/commit", json=bad), 404)
    committed = expect(teacher.post(f"/api/sessions/s4/worksheet-import/{preview['import_id']}/commit", json=payload))
    assert committed["question_count"] == 2 and any(m["id"] == mid for m in committed["materials"])
    assert expect(student.get("/api/sessions/s4/workspace"))["version"]["id"] == initial
    expect(student.get(asset["url"]), 404)
    edition = publish(teacher)["id"]
    work = expect(student.get("/api/sessions/s4/workspace?version_id="+edition))
    assert len(work["session"]["questions"]) == 2
    image_url = next(image["url"] for q in work["session"]["questions"] for image in q["images"])
    assert "version_id="+edition in image_url
    assert student.get(image_url).content.startswith(b"\x89PNG")
    expect(student.get(f"/api/materials/{mid}/download?version_id="+edition), 403)
    expect(student.get(f"/api/materials/{mid}/preview?version_id="+edition), 403)
    expect(teacher.delete("/api/sessions/s4/materials/"+mid), 409)
    edited = expect(teacher.get("/api/sessions/s4"))["questions"]
    for q in edited:
        q["images"] = []
    expect(teacher.patch("/api/sessions/s4", json={"questions": edited}))
    expect(teacher.delete("/api/sessions/s4/materials/"+mid))
    new = publish(teacher)["id"]
    expect(student.get(f"/api/materials/{mid}/assets/{asset['key']}?version_id="+new), 404)
    assert student.get(image_url).content.startswith(b"\x89PNG")
    with app.state.database.connect() as conn:
        path = Path(conn.execute("SELECT file_path FROM materials WHERE id=?", (mid,)).fetchone()[0])
        assert path.is_file()
        assert not conn.execute("PRAGMA foreign_key_check").fetchall()


def test_worksheet_unreferenced_assets_scope_and_demo_ai_rejection(app, teacher, student):
    preview = expect(teacher.post("/api/sessions/s4/worksheet-import/preview", files={"file": ("figure.docx", worksheet_docx())}))
    payload = {k: preview["draft"][k] for k in ("groups", "questions", "reference_answers")}
    for q in payload["questions"]:
        q["images"] = []
    payload["mode"] = "replace"
    expect(teacher.post(f"/api/sessions/s2/worksheet-import/{preview['import_id']}/commit", json=payload), 404)
    expect(teacher.post(f"/api/sessions/s4/worksheet-import/{preview['import_id']}/commit", json=payload))
    version = publish(teacher)["id"]
    asset = preview["assets"][0]
    expect(student.get(asset["url"]+"?version_id="+version), 403)
    expect(teacher.post("/api/sessions/s4/worksheet-import/preview", data={"use_ai": "true"}, files={"file": ("x.txt", b"1. Explain reliability?")}), 409)
    with demo_client(app, "admin") as other_owner:
        expect(other_owner.post(f"/api/sessions/s4/worksheet-import/{preview['import_id']}/commit", json=payload), 403)


def test_optional_ai_protocol_scope_real_failure_visible_and_review_required(app, teacher, admin, monkeypatch):
    monkeypatch.setenv("AICADEMIYA_AI_PROVIDER", "aicademiya-router")
    expect(admin.put("/api/teacher/settings", json={"mode": "openai", "base_url": "http://127.0.0.1:18280/v1", "model": "AIcademiya-64kx3-Qwen3.8-27B"}))
    calls = []
    def failed(request):
        calls.append(request)
        return httpx.Response(503, json={"error": "LABELLED_PRIVATE_FIXTURE_DIAGNOSTIC"})
    app.state.tutor.transport = httpx.MockTransport(failed)
    error = expect(teacher.post("/api/sessions/s4/worksheet-import/preview", data={"use_ai": "true"}, files={"file": ("fixture.txt", b"1. Explain assessment reliability?")}), 502)
    assert "HTTP 503" in error["detail"] and "LABELLED_PRIVATE" not in error["detail"]
    assert len(calls) == 1 and json.loads(calls[0].content)["user"].startswith("aicv2-") and calls[0].extensions["timeout"]["read"] == 600
    with app.state.database.connect() as conn:
        assert not conn.execute("SELECT 1 FROM import_staging WHERE kind='worksheet'").fetchone()
    def structured(request):
        calls.append(request)
        q = editable_question(teacher)
        draft = {"title": "Labelled protocol extracted draft", "groups": [{"id": "core", "title": "Core", "kind": "core"}], "questions": [q], "reference_answers": {}}
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(draft)}}]})
    app.state.tutor.transport = httpx.MockTransport(structured)
    preview = expect(teacher.post("/api/sessions/s4/worksheet-import/preview", data={"use_ai": "true"}, files={"file": ("fixture.txt", b"1. Explain assessment reliability?")}))
    assert preview["method"] == "llm" and preview["needs_review"] and preview["draft"]["title"].startswith("Labelled protocol")
    assert json.loads(calls[0].content)["user"] != json.loads(calls[1].content)["user"]
    assert expect(teacher.get(VERSIONS))["current_version_id"] == expect(teacher.get(VERSIONS))["versions"][0]["id"]
