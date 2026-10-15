"""Document previews, private originals, scan detection, and library retention."""
import io
from pathlib import Path
import pytest
from .conftest import expect, publish


def payload_for(suffix):
    output = io.BytesIO()
    if suffix in ("txt", "md"):
        return b"Synthetic material evidence, no real student information."
    if suffix == "pdf":
        import pymupdf
        document = pymupdf.open()
        document.new_page().insert_text((50, 50), "Synthetic PDF teaching evidence.")
        data = document.tobytes()
        document.close()
        return data
    if suffix == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Synthetic DOCX teaching evidence.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Synthetic cell"
        table.cell(0, 1).text = "Evidence"
        document.save(output)
    if suffix == "pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Synthetic slide evidence"
        slide.placeholders[1].text = "No real class content."
        document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("suffix", ["txt", "md", "pdf", "docx", "pptx"])
def test_upload_extract_and_download(suffix, teacher, student):
    data = payload_for(suffix)
    material = expect(teacher.post("/api/sessions/s4/materials", data={"visibility": "student", "kind": "reading"}, files={"file": (f"synthetic.{suffix}", data)}))
    assert material["status"] == "ready" and "Synthetic" in material["excerpt"]
    assert material["download_url"] == f"/api/materials/{material['id']}/download"
    expect(student.get(material["download_url"]), 404)
    publish(teacher)
    assert student.get(material["download_url"]).content == data
    assert any(m["id"] == material["id"] for m in expect(student.get("/api/sessions/s4"))["materials"])


def test_teacher_upload_private_and_detach_without_deleting_original(app, teacher, student):
    item = expect(teacher.post("/api/sessions/s4/materials", data={"visibility": "teacher", "kind": "answer_key"}, files={"file": ("key.txt", b"UPLOADED_TEACHER_ONLY_SECRET")}))
    assert "UPLOADED_TEACHER_ONLY_SECRET" not in student.get("/api/sessions/s4/workspace").text
    expect(student.get(item["download_url"]), 404)
    publish(teacher)
    expect(student.get(item["download_url"]), 403)
    expect(student.post("/api/sessions/s4/materials", files={"file": ("x.txt", b"text")}), 403)
    expect(teacher.post("/api/sessions/s4/materials", data={"visibility": "student", "kind": "answer_key"}, files={"file": ("key.txt", b"No leak")}), 422)
    with app.state.database.connect() as conn:
        original = Path(conn.execute("SELECT file_path FROM materials WHERE id=?", (item["id"],)).fetchone()[0])
    expect(teacher.delete(f"/api/sessions/s4/materials/{item['id']}"))
    assert teacher.get(item["download_url"]).content == b"UPLOADED_TEACHER_ONLY_SECRET"
    assert all(m["id"] != item["id"] for m in expect(teacher.get("/api/sessions/s4"))["materials"])
    retained = next(m for m in expect(teacher.get("/api/courses/synthetic-course/materials"))["materials"] if m["id"] == item["id"])
    assert retained["session_ids"] == []
    assert original.is_file() and original.read_bytes() == b"UPLOADED_TEACHER_ONLY_SECRET"


def test_scanned_pdf_is_needs_review_not_success(teacher):
    import pymupdf
    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    item = expect(teacher.post("/api/sessions/s4/materials", files={"file": ("image-only.pdf", data)}))
    assert item["status"] == "needs_review" and item["excerpt"] == ""
    assert "OCR" in item["warning"]


def test_invalid_format_and_corrupt_file_are_visible(teacher):
    expect(teacher.post("/api/sessions/s4/materials", files={"file": ("not-allowed.exe", b"MZ")}), 415)
    item = expect(teacher.post("/api/sessions/s4/materials", files={"file": ("corrupt.pdf", b"invalid PDF")}))
    assert item["status"] == "error" and item["warning"]
    assert teacher.get(item["download_url"]).content == b"invalid PDF"
    item = expect(teacher.post("/api/sessions/s4/materials", files={"file": ("empty.txt", b"")}))
    assert item["status"] == "needs_review"


def test_frontend_static_serve_and_spa_route(app, student, tmp_path):
    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "index.html").write_text("<title>AIcademiya synthetic build fixture</title>", encoding="utf-8")
    (dist / "app.js").write_text("console.log('synthetic fixture')", encoding="utf-8")
    assert "AIcademiya" in student.get("/").text
    assert "AIcademiya" in student.get("/courses/route").text
    assert "synthetic fixture" in student.get("/app.js").text
    expect(student.get("/api/still-missing"), 404)
    expect(student.get("/%2E%2E%2Fsample_data%2Fcatalog.json"), 404)
