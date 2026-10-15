"""Stage10 passive content regressions: no models/search/external XML fetch."""
import io
import json
import socket
import zipfile
import pytest
from PIL import Image
from docx import Document
from backend.aicademiya import material_extraction, worksheet_import
from .conftest import expect


def test_upload_raw_html_remains_plain_text_and_svg_upload_is_denied(teacher):
    payload = b'<script>alert("SYNTHETIC_XSS")</script><img src=x onerror=alert(1)>'
    item = expect(teacher.post('/api/sessions/s4/materials', files={'file': ('raw-html.md', payload)}))
    response = teacher.get(item['preview_url'])
    assert response.status_code == 200
    assert response.content == payload
    assert response.headers['content-type'].startswith('text/plain')
    assert response.headers['x-content-type-options'] == 'nosniff'
    assert teacher.post('/api/sessions/s4/materials', files={'file': ('active.svg', b'<svg onload="alert(1)"/>')}).status_code == 415


def test_imported_image_is_fresh_png_not_source_metadata_or_active_svg(tmp_path):
    source = io.BytesIO()
    from PIL.PngImagePlugin import PngInfo
    meta = PngInfo()
    meta.add_text('SyntheticUnsafeMetadata', '<script>not executed</script>')
    Image.new('RGB', (40, 40), 'green').save(source, format='PNG', pnginfo=meta)
    warnings = []
    assets = worksheet_import._Assets(tmp_path / 'figures', warnings)
    figure = assets.save(source.getvalue(), 'Synthetic safe raster')
    assert figure
    normalized = (tmp_path / 'figures' / figure['asset_key']).read_bytes()
    assert b'SyntheticUnsafeMetadata' not in normalized and b'<script>' not in normalized
    with Image.open(io.BytesIO(normalized)) as image:
        assert image.format == 'PNG' and image.size == (40, 40)
    assert assets.save(b'<svg onload="alert(1)"/>', 'Synthetic active vector') is None


def test_docx_external_entity_never_discloses_local_file(tmp_path, monkeypatch):
    secret = tmp_path / 'outside-private.txt'
    secret.write_text('SYNTHETIC_XML_OUTSIDE_SECRET', encoding='utf-8')
    original = tmp_path / 'base.docx'
    document = Document()
    document.add_paragraph('1. Explain safe extraction.')
    document.save(original)
    malicious = tmp_path / 'external-entity.docx'
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(malicious, 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'word/document.xml':
                data = ('<!DOCTYPE w:document [<!ENTITY secret SYSTEM "' + secret.as_uri() + '">]>' +
                        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&secret;</w:t></w:r></w:p></w:body></w:document>').encode()
            target.writestr(info.filename, data)
    network_calls = []
    def refuse(*args, **kwargs):
        network_calls.append(1)
        raise AssertionError('External request attempted')
    monkeypatch.setattr(socket, 'create_connection', refuse)
    result = material_extraction.extract_document(malicious)
    worksheet = worksheet_import.preview_worksheet(malicious, assets_dir=tmp_path / 'figures')
    assert 'SYNTHETIC_XML_OUTSIDE_SECRET' not in json.dumps(result)
    assert 'SYNTHETIC_XML_OUTSIDE_SECRET' not in json.dumps(worksheet)
    assert network_calls == []
    assert secret.read_text(encoding='utf-8') == 'SYNTHETIC_XML_OUTSIDE_SECRET'


def test_zip_traversal_member_is_not_extracted_to_filesystem(tmp_path):
    document = Document()
    document.add_paragraph('1. Explain safe traversal boundaries.')
    path = tmp_path / 'traversal.docx'
    document.save(path)
    with zipfile.ZipFile(path, 'a', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('../synthetic-escape.txt', 'SYNTHETIC_TRAVERSAL_CONTENT')
    material = material_extraction.extract_document(path)
    worksheet = worksheet_import.preview_worksheet(path, assets_dir=tmp_path / 'figures')
    assert material['status'] == 'ready'
    assert worksheet['questions']
    assert not (tmp_path.parent / 'synthetic-escape.txt').exists()
    assert not (tmp_path / 'synthetic-escape.txt').exists()
