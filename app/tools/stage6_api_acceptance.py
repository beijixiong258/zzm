"""Stage6 real HTTP acceptance: synthetic demo, real extraction, no app patches.
Run through run_stage6_integration.py or against isolated staged QA with confirmation.
"""
import argparse
import hashlib
import json
import mimetypes
import os
from pathlib import Path
import secrets
import sqlite3
import sys
import traceback

import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
PROFILE = dict(schema_version=1, background='Synthetic private design only', student_profile='Synthetic learners',
               response_strategies='Invite evidence and preserve raw responses.', tone='Supportive', response_detail='balanced',
               phases=['worksheet', 'wrap_up'], reading_focus='Public evidence only.', reading_question_count=1,
               reference_delivery='on_request', idle_hint_seconds=0, presentation='full')
TOKEN = 'SYNTHETIC INSTRUCTION NATIVE TEXT'


def question(ids=(), qid='s6-q'):
    return dict(id=qid, group_id='core', ordinal=1, type='open_response',
                prompt='Explain targetalpha targetbeta with evidence.', options=[], answer_slots=1, images=[],
                source_label='Synthetic duplicate evidence', source_material_ids=list(ids), source_links=[])


def make_fixtures(folder):
    from docx import Document
    import pymupdf
    folder.mkdir(parents=True, exist_ok=True)
    (folder / 'Synthetic Instructions.txt').write_text(TOKEN + '\nPreserve raw learner responses.', encoding='utf-8-sig')
    doc = Document()
    doc.add_paragraph(TOKEN)
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = 'Progressive hints'
    t.cell(0, 1).text = 'No formal grades'
    doc.save(folder / 'Synthetic Instructions.docx')
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((40, 50), TOKEN + '\nKeep private references private.')
        pdf.save(folder / 'Synthetic Instructions.pdf')
    (folder / 'Synthetic Overlimit.txt').write_text('x' * 30010, encoding='utf-8')
    (folder / 'Synthetic Broken.docx').write_bytes(b'Explicit synthetic invalid ZIP, not a document')
    return {p.name: {'sha256': hashlib.sha256(p.read_bytes()).hexdigest(), 'bytes': p.stat().st_size}
            for p in sorted(folder.iterdir()) if p.is_file()}


def records(db):
    with sqlite3.connect(f'file:{Path(db).as_posix()}?mode=ro', uri=True) as c:
        tables = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        return {t: sorted([list(r) for r in c.execute('SELECT * FROM "' + t + '"')], key=repr) for t in tables}


class Acceptance:
    def __init__(self, base, out, fixtures, db=None):
        self.out, self.fixtures, self.db = Path(out), Path(fixtures), db
        self.checks, self.failures, self.events, self.ids, self.details = [], [], [], {}, {}
        self.clients = [httpx.Client(base_url=base.rstrip('/'), trust_env=False, timeout=120) for _ in range(4)]
        self.teacher, self.student, self.other, self.anon = self.clients

    def call(self, client, method, url, body=None, expected=200, **kwargs):
        r = client.request(method, url, json=body, **kwargs)
        self.events.append(dict(method=method, url=url, status=r.status_code, expected=expected))
        if r.status_code != expected:
            raise AssertionError(f'{method} {url}: expected {expected}, got {r.status_code}; {r.text[:500]}')
        return r.json() if r.headers.get('content-type', '').startswith('application/json') else r

    def check(self, name, condition, details=None):
        if not condition:
            raise AssertionError(name + ': ' + str(details))
        self.checks.append(name)
        if details is not None:
            self.details[name] = details
        print('PASS', name, flush=True)
        self.save('running')

    def save(self, status):
        self.out.mkdir(parents=True, exist_ok=True)
        result = dict(status=status, checks=self.checks, failures=self.failures, requests=self.events, ids=self.ids,
                      request_count=len(self.events), details=self.details, real_backend_http=True,
                      synthetic_only=True, real_model=False, intercepted_application_routes=0,
                      sqlite_inspected=bool(self.db), evidence='Explicit demo; native DOCX/TXT/PDF extraction is real. Not real-model quality.')
        (self.out / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')

    def run(self):
        boot = self.call(self.anon, 'GET', '/api/bootstrap')
        health = self.call(self.anon, 'GET', '/api/health')
        spec = self.call(self.anon, 'GET', '/openapi.json')
        self.check('Explicit demo-enabled target and API 0.5.0', boot['demo_enabled'] and boot['ai']['mode'] == 'demo' and spec['info']['version'] == '0.5.0', {'health': health, 'api_version': spec['info']['version']})
        for client, body in ((self.teacher, {'role': 'teacher'}), (self.student, {'role': 'student', 'student_index': 1}), (self.other, {'role': 'admin'})):
            self.call(client, 'POST', '/api/auth/demo', body)
        course = self.call(self.teacher, 'POST', '/api/courses', {'code': 'S6-API-' + secrets.token_hex(4), 'title': 'Synthetic Stage6 HTTP lab'})
        cid = course['id']
        sid = self.call(self.teacher, 'POST', f'/api/courses/{cid}/sessions', {'title': 'Synthetic teacher workbench', 'teaching_profile': PROFILE, 'instructions': 'SAVED_SYNTHETIC_BASE', 'learning_objectives': ['Explain evidence.']})['id']
        self.ids.update(course=cid, session=sid)
        detail = lambda: self.call(self.teacher, 'GET', f'/api/sessions/{sid}')
        work = lambda: self.call(self.teacher, 'GET', f'/api/sessions/{sid}/workspace')
        info = lambda: self.call(self.teacher, 'GET', f'/api/teacher/sessions/{sid}/preview')
        review = lambda body: self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/source-review', body)
        upload = lambda name, text, **meta: self.call(self.teacher, 'POST', f'/api/sessions/{sid}/materials', files={'file': (name, text.encode(), 'text/plain')}, data={'kind': 'reading', 'visibility': 'student', **meta})
        public = [upload('DUPLICATE LABEL.txt', 'noise ' * 900 + ' targetalpha targetbeta ' + 'tail ' * 1000) for i in range(18)]
        private = [upload('DUPLICATE LABEL.txt', 'PRIVATE_CONCEPT ' * 450, visibility='teacher', kind='reference') for i in range(5)]
        key = upload('key.txt', 'PRIVATE_KEY_NEVER_AGGREGATE', visibility='teacher', kind='answer_key')
        sheet = upload('worksheet.txt', 'PRIVATE_WORKSHEET_NEVER_AGGREGATE', visibility='teacher', kind='worksheet')
        unselected = upload('unselected.txt', 'UNSELECTED_SENTINEL')
        detached = upload('library-only.txt', 'UNATTACHED_SENTINEL')
        self.call(self.teacher, 'DELETE', f'/api/sessions/{sid}/materials/{detached["id"]}')
        q = question([m['id'] for m in public + private + [key, sheet]])
        self.call(self.teacher, 'PATCH', f'/api/sessions/{sid}', {'groups': [{'id': 'core', 'title': 'Synthetic core', 'kind': 'core'}], 'questions': [q], 'reference_answers': {'s6-q': 'PRIVATE_CURRENT_REFERENCE'}, 'reference_policy': 'after_submit', 'status': 'published'})
        work()  # Initialize the empty workspace before strict read-only snapshots.
        before = records(self.db) if self.db else None
        initial = info()
        r = review({'question_id': 's6-q', 'phase': 'worksheet'})
        overlay = review({'question_id': 's6-q', 'instructions': 'UNSAVED_OVERLAY'})
        preview = self.call(self.teacher, 'POST', f'/api/sessions/{sid}/teaching-preview', {'instructions': 'UNSAVED_PURE_PREVIEW'})
        self.check('Read-only preview info does not bind/start and teaching overlay does not save', initial['id'] is None and not initial['pending'] and detail()['instructions'] == 'SAVED_SYNTHETIC_BASE' and 'UNSAVED_OVERLAY' in overlay['compiled']['text'] and 'UNSAVED_PURE_PREVIEW' in preview['text'])
        self.check('Exact actual evidence IDs survive duplicate labels and public/private truncation', len(r['public_evidence']) == 16 and len({x['material_id'] for x in r['public_evidence']}) == 16 and {x['source_label'] for x in r['public_evidence']} == {'DUPLICATE LABEL.txt'} and sum(len(x['excerpt']) for x in r['public_evidence']) == 32000 and sum(len(x['excerpt']) for x in r['private_evidence']) == 8000, {'public_ids': [x['material_id'] for x in r['public_evidence']], 'public_chars': 32000, 'private_chars': 8000})
        rows = {x['id']: x for x in r['materials']}
        self.check('Unselected/unattached/private worksheet/key excluded with reasons', detached['id'] not in rows and not rows[unselected['id']]['included'] and not rows[key['id']]['included'] and not rows[sheet['id']]['included'] and 'budget' in rows[public[-1]['id']]['reason'] and r['private_reference']['text'] == 'PRIVATE_CURRENT_REFERENCE')
        if self.db:
            from backend.aicademiya.version_context import session_access
            from backend.aicademiya.authoring import selected_tutor_materials
            from backend.aicademiya.ai import bounded_evidence
            with sqlite3.connect(f'file:{Path(self.db).as_posix()}?mode=ro', uri=True) as c:
                c.row_factory = sqlite3.Row
                actor = dict(c.execute("SELECT * FROM users WHERE id='demo-teacher'").fetchone())
                saved = session_access(c, sid, actor)
                pub, pri = selected_tutor_materials(c, saved, 's6-q')
                self.check('Source review equals actual live selector over real SQLite rows', [{k: v for k, v in x.items() if k != 'material_id'} for x in r['public_evidence']] == bounded_evidence(pub, q) and [{k: v for k, v in x.items() if k != 'material_id'} for x in r['private_evidence']] == bounded_evidence(pri, q, private=True, budget=8000))
        for phase in ('quiz', 'reading'):
            optional = review({'question_id': 's6-q', 'phase': phase})
            self.check(f'{phase} optional phase excludes all private evidence/reference', not optional['private_evidence'] and optional['private_reference'] is None and all('PRIVATE_' not in x['excerpt'] for x in optional['public_evidence']))
        extracted = []
        for ext in ('docx', 'txt', 'pdf'):
            f = self.fixtures / f'Synthetic Instructions.{ext}'
            raw = f.read_bytes()
            result = self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': (f.name, raw, mimetypes.guess_type(f.name)[0] or 'application/octet-stream')})
            extracted.append(result)
            self.check('Actual multipart native extraction ' + ext, TOKEN in result['text'] and result['status'] == 'ready' and result['extraction_report']['method'] == 'native' and result['max_instruction_chars'] == 30000 and f.read_bytes() == raw, result)
        long = self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': ('Synthetic Overlimit.txt', (self.fixtures / 'Synthetic Overlimit.txt').read_bytes())})
        self.check('Over-field-limit extraction remains editable without silent truncation', len(long['text']) == 30010 and long['status'] == 'needs_review' and '30,000' in long['warning'])
        self.call(self.teacher, 'PATCH', f'/api/sessions/{sid}', {'instructions': long['text']}, expected=422)
        broken = self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': ('Synthetic Broken.docx', (self.fixtures / 'Synthetic Broken.docx').read_bytes())})
        self.check('Actual parser error is visible without invented content', broken['status'] == 'error' and not broken['text'] and bool(broken['warning']))
        self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': ('Synthetic Oversize.txt', b'x' * (25 * 1024 * 1024 + 1))}, expected=413)
        self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': ('unsupported.exe', b'not instructions')}, expected=422)
        for client, status in ((self.anon, 401), (self.student, 403), (self.other, 403)):
            for route, method, body in (('preview', 'GET', None), ('preview/reset', 'POST', {'expected_revision': initial['current_revision']}), ('source-review', 'POST', {})):
                self.call(client, method, f'/api/teacher/sessions/{sid}/{route}', body, expected=status)
            self.call(client, 'POST', f'/api/teacher/sessions/{sid}/instruction-import/preview', files={'file': ('Synthetic.txt', b'Synthetic text')}, expected=status)
        self.check('Owner scope / student / unrelated staff / anonymous rejects', True)
        if before is not None:
            self.check('Pure source/compiler/import/denied requests leave every SQLite table unchanged', records(self.db) == before)
        self.check('No automatic attachments or publication from import', len(detail()['materials']) == len(public + private + [key, sheet, unselected]))
        # Create only an explicitly synthetic published history, to prove reset isolation.
        vid = self.call(self.teacher, 'POST', f'/api/teacher/courses/{cid}/versions/publish', {'label': 'Synthetic preserved edition'})['id']
        self.call(self.student, 'POST', '/api/courses/join', {'join_code': course['join_code']})
        swurl = f'/api/sessions/{sid}/workspace?version_id={vid}'
        self.call(self.student, 'POST', f'/api/sessions/{sid}/answers/s6-q?version_id={vid}', {'answer': 'Protected synthetic student raw response'})
        sw = self.call(self.student, 'GET', swurl)
        history = self.call(self.student, 'GET', '/api/history')
        snapshot = self.call(self.teacher, 'GET', f'/api/teacher/courses/{cid}/versions/{vid}')
        reset = lambda revision: self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/preview/reset', {'expected_revision': revision})
        w = reset(info()['current_revision'])
        pid = w['preview']['id']
        start = lambda token, expected=200: self.call(self.teacher, 'POST', f'/api/sessions/{sid}/tutor/start', {'request_id': 's6-opening', 'preview_id': token}, expected=expected)
        start(pid)
        self.call(self.teacher, 'PATCH', f'/api/sessions/{sid}', {'instructions': 'EXPLICIT_SYNTHETIC_SAVED_CHANGE'})
        stale = work()
        self.check('Saved edit makes old preview readable but stale', stale['read_only_reason'] == 'preview_stale' and stale['preview']['id'] == pid and bool(stale['messages']))
        start(pid, 409)
        self.call(self.teacher, 'POST', f'/api/sessions/{sid}/answers/s6-q?preview_id={pid}', {'answer': 'Queued stale'}, expected=409)
        self.call(self.teacher, 'POST', f'/api/teacher/sessions/{sid}/preview/reset', {'expected_revision': w['preview']['revision']}, expected=409)
        fresh = reset(info()['current_revision'])
        newpid = fresh['preview']['id']
        start(pid, 409)
        self.call(self.teacher, 'POST', f'/api/sessions/{sid}/tutor/turn', {'message': 'Queued old tab', 'preview_id': pid, 'request_id': 's6-old'}, expected=409)
        self.check('Reset creates new generation / clears only own preview / rejects old queued tab', newpid != pid and not work()['messages'] and not work()['attempts'])
        self.check('Student raw work/history and immutable edition survive teacher reset', self.call(self.student, 'GET', swurl) == sw and self.call(self.student, 'GET', '/api/history') == history and self.call(self.teacher, 'GET', f'/api/teacher/courses/{cid}/versions/{vid}') == snapshot)
        # Parent-added full-preview synchronous query-token compatibility.
        for token in (None, pid):
            self.call(self.teacher, 'POST', f'/api/sessions/{sid}/answers/s6-q' + (f'?preview_id={token}' if token else ''), {'answer': 'Forbidden token'}, expected=409)
        self.call(self.teacher, 'POST', f'/api/sessions/{sid}/chat', {'message': 'Legacy managed chat', 'intent': 'ask'}, expected=409)
        self.call(self.teacher, 'POST', f'/api/sessions/{sid}/answers/s6-q?preview_id={newpid}', {'answer': 'Explicit full preview saved evidence'})
        self.call(self.teacher, 'POST', f'/api/sessions/{sid}/references/s6-q?preview_id={newpid}', {}, expected=409)
        submitted = self.call(self.teacher, 'POST', f'/api/sessions/{sid}/submit?preview_id={newpid}', {})
        ref = self.call(self.teacher, 'POST', f'/api/sessions/{sid}/references/s6-q?preview_id={newpid}', {})
        self.check('Full managed preview saves/submits/reveals after_submit using exact query token', submitted['submission']['status'] == 'submitted' and work()['attempts']['s6-q']['answer'] == 'Explicit full preview saved evidence' and ref['reference'] == 'PRIVATE_CURRENT_REFERENCE')
        reading_sid = self.call(self.teacher, 'POST', f'/api/courses/{cid}/sessions', {'title': 'Synthetic full-preview reading branch', 'teaching_profile': PROFILE})['id']
        reading_q = question(qid='reading-q')
        reading_q['group_id'] = 'reading-path'
        self.call(self.teacher, 'PATCH', f'/api/sessions/{reading_sid}', {'groups': [{'id': 'reading-path', 'title': 'Synthetic reading choice', 'kind': 'reading'}], 'questions': [reading_q]})
        reading_info = self.call(self.teacher, 'GET', f'/api/teacher/sessions/{reading_sid}/preview')
        reading_w = self.call(self.teacher, 'POST', f'/api/teacher/sessions/{reading_sid}/preview/reset', {'expected_revision': reading_info['current_revision']})
        reading_pid = reading_w['preview']['id']
        self.call(self.teacher, 'POST', f'/api/sessions/{reading_sid}/reading', {'group_id': 'reading-path'}, expected=409)
        chosen = self.call(self.teacher, 'POST', f'/api/sessions/{reading_sid}/reading?preview_id={reading_pid}', {'group_id': 'reading-path'})
        self.check('Full managed preview reading selection needs current query token', chosen['selected_reading'] == 'reading-path')
        self.ids.update(student='demo-student-1', version=vid, final_preview_id=newpid)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--base-url', default=os.getenv('AICADEMIYA_TEST_URL'))
    p.add_argument('--artifact-dir', default=os.getenv('AICADEMIYA_TEST_ARTIFACT_DIR'))
    p.add_argument('--fixture-dir', required=True)
    p.add_argument('--confirm-test-data', action='store_true')
    args = p.parse_args()
    if not args.base_url or not args.artifact_dir or not args.confirm_test_data:
        p.error('Require isolated synthetic demo QA URL, artifact directory and --confirm-test-data. Never production.')
    a = Acceptance(args.base_url, args.artifact_dir, args.fixture_dir, os.getenv('AICADEMIYA_QA_DB'))
    try:
        a.run()
    except Exception as e:
        a.failures.append({'name': 'API scenario interrupted', 'error': str(e), 'traceback': traceback.format_exc()})
        print(traceback.format_exc(), flush=True)
    finally:
        for c in a.clients:
            c.close()
        a.save('failed' if a.failures else 'passed')
    print(json.dumps({'checks': len(a.checks), 'failures': len(a.failures)}))
    return 1 if a.failures else 0


if __name__ == '__main__':
    raise SystemExit(main())
