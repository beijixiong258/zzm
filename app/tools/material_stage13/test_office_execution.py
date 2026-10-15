"""Actual LibreOffice acceptance with synthetic hidden-slide/failure fixtures."""
import argparse
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE.parent/'document_stage12'))
from parser_helpers import sha256,dump,validate_bundle


def run(source,dest,*extra):
    cmd=[sys.executable,str(HERE/'parse_material.py'),'--input',str(source),'--output',str(dest),
         '--engine','native','--dpi','96','--deadline-seconds','60']+list(extra)
    p=subprocess.run(cmd,capture_output=True,text=True,timeout=70)
    return p


def main():
    cli=argparse.ArgumentParser();cli.add_argument('--output',required=True);a=cli.parse_args()
    root=Path(a.output);root.mkdir(parents=True,exist_ok=False)
    from pptx import Presentation
    deck=Presentation()
    for hidden,text in [(False,'Visible slide fixture'),(True,'Hidden slide fixture')]:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0,0,deck.slide_width,deck.slide_height).text=text
        if hidden:slide._element.set('show','0')
        slide.notes_slide.notes_text_frame.text='Fixture speaker note not integrated into student reading.'
    source=root/'hidden.pptx';deck.save(source);before=sha256(source)
    p=run(source,root/'hidden-bundle');assert p.returncode==0,(p.stdout,p.stderr)
    m=json.loads((root/'hidden-bundle/manifest.json').read_text())
    assert m['source_sha256']==before and m['page_count']==2 and m['coverage']['hidden_slides']==1
    assert m['coverage']['notes_count']==2 and m['coverage']['notes_integrated_in_reading'] is False
    assert any('not integrated into student reading' in w for w in m['warnings'])
    rows=[{'case':'real_hidden_slide_conversion','source_unchanged':sha256(source)==before,
           'page_count':2,'hidden_slides':1,'notes_retained':2,'notes_integrated':False,
           'validation':validate_bundle(root/'hidden-bundle',m)}]
    before_manifest=sha256(root/'hidden-bundle/manifest.json')
    p=run(source,root/'hidden-bundle');assert p.returncode!=0 and sha256(root/'hidden-bundle/manifest.json')==before_manifest
    rows.append({'case':'fresh_output_no_overwrite','passed':True})
    large=root/'oversize.docx'
    with large.open('wb') as f:f.truncate(64*1024*1024+1)
    p=run(large,root/'oversize-bundle');assert p.returncode!=0 and not (root/'oversize-bundle').exists()
    rows.append({'case':'source_over64MiB_before_allocation','passed':True})
    p=run(source,root/'deadline-bundle','--deadline-seconds','0.01')
    receipt=json.loads((root/'deadline-bundle/receipt.json').read_text());m=json.loads((root/'deadline-bundle/manifest.json').read_text())
    assert p.returncode!=0 and receipt['status']=='failed' and not m['complete'] and m['source_sha256']==before
    rows.append({'case':'deadline_preserves_failure_and_original_identity','passed':True,'error_type':receipt['error_type']})
    dump(root/'result.json',{'status':'passed','cases':rows,'model_calls':0,'fixture_label':'synthetic; actual Linux LibreOffice conversion'})
    print(json.dumps(rows,indent=2))


if __name__=='__main__':main()
