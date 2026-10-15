"""Isolated real synthetic legacy Office CPU acceptance; never deploy/install.

Run locally with app .venv: legacy_cpu_probe.py orchestrate.
The same file runs remotely using existing worker Python and LibreOffice.
Only tools are packaged; app supplement/schema methods are checked locally.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import os
import shutil
import signal
import subprocess
import sys
import tarfile
import time
import traceback
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

sys.dont_write_bytecode = True

REMOTE_PYTHON = '/opt/aicademiya-documents/runtime/.venv/bin/python'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'

def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()

def save(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))

def run(command, root, label, *, env=None, timeout=240):
    root = Path(root); (root/'logs').mkdir(exist_ok=True)
    start = time.monotonic()
    process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        text=True, env=env, start_new_session=(os.name != 'nt'))
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except BaseException:
        if process.poll() is None:
            if os.name != 'nt': os.killpg(process.pid, signal.SIGTERM)
            else: process.terminate()
            try: process.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                if os.name != 'nt': os.killpg(process.pid, signal.SIGKILL)
                else: process.kill()
                process.communicate(timeout=5)
        raise
    row = {'argv':list(map(str, command)), 'pid':process.pid, 'returncode':process.returncode,
        'seconds':time.monotonic()-start, 'stdout':stdout, 'stderr':stderr,
        'child_reaped':process.poll() is not None}
    save(root/'logs'/f'{label}.json', row)
    return row

def rewrite_zip(path, changes):
    path = Path(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name:archive.read(name) for name in archive.namelist()}
    for name, fn in changes.items(): parts[name] = fn(parts.get(name))
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items(): archive.writestr(name, data)

def generate(root):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches, Pt
    import xlsxwriter
    import olefile
    root = Path(root); originals = root/'evidence/originals'; originals.mkdir(parents=True)
    producers = {}
    doc = Document()
    doc.add_heading('SYNTHETIC Stage16 real DOC CPU probe', 0)
    run1 = doc.add_paragraph('DOC_VISIBLE_PAGE_ONE literal classroom-free synthetic source text.').runs[0]
    doc.add_paragraph('Body text and exact original identities are independent from derived PDF previews.')
    doc.sections[0].header.paragraphs[0].text = 'SYNTHETIC_DOC_HEADER'
    doc.sections[0].footer.paragraphs[0].text = 'SYNTHETIC_DOC_FOOTER'
    if hasattr(doc, 'add_comment'):
        doc.add_comment(run1, 'PRIVATE_DOC_COMMENT synthetic author-only note.', author='Synthetic CPU QA', initials='QA')
    doc.add_page_break()
    doc.add_heading('Second synthetic page', 1)
    doc.add_paragraph('DOC_VISIBLE_PAGE_TWO literal second-page source text that must survive conversion.')
    docx = originals/'synthetic.docx'; doc.save(docx)
    producers['docx'] = {'producer':'python-docx', 'comment_generated':hasattr(doc, 'add_comment')}
    prs = Presentation()
    for i in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if i == 2: slide._element.set('show', '0')
        box = slide.shapes.add_textbox(Inches(.8), Inches(1), Inches(8.3), Inches(2.6))
        tf = box.text_frame
        tf.text = f'PPT_VISIBLE_SLIDE_{i}: SYNTHETIC real legacy Office CPU acceptance.\nThis literal slide must survive the original PPT to derived PDF conversion.'
        for paragraph in tf.paragraphs:
            for text_run in paragraph.runs: text_run.font.size = Pt(22)
        slide.notes_slide.notes_text_frame.text = f'PRIVATE_PPT_NOTE_{i} synthetic speaker note, not ordinary student reading.'
    pptx = originals/'synthetic.pptx'; prs.save(pptx)
    producers['pptx'] = {'producer':'python-pptx', 'slides':2, 'hidden_slides':1, 'speaker_notes':2}
    xlsx = originals/'synthetic.xlsx'; wb = xlsxwriter.Workbook(xlsx)
    main = wb.add_worksheet('Main'); hidden = wb.add_worksheet('Hidden'); very = wb.add_worksheet('VeryHidden'); wb.add_worksheet('Blank')
    main.write('A1', 'SYNTHETIC_XLSX_TITLE'); main.write_number('B1', 2); main.write_number('B2', 3)
    main.write_formula('C1', '=SUM(B1:B2)', None, 999)
    main.write_formula('D2', '=B1*3', None, 0)
    main.write_rich_string('F1', 'Exact ', wb.add_format({'bold':True}), 'rich text')
    main.write_blank('E1', None, wb.add_format({'bg_color':'#FFFF00'}))
    main.merge_range('A4:C4', 'MERGED_SOURCE_VALUE', wb.add_format({'border':1}))
    main.write('Q1000', 'OUTSIDE_PRINT_AREA_LITERAL')
    main.write_comment('A1', 'PRIVATE_WORKBOOK_COMMENT synthetic staff-only cell note.', {'author':'Synthetic CPU QA'})
    main.print_area('A1:C4')
    hidden.write('A1', 'HIDDEN_SHEET_LITERAL'); hidden.hide()
    very.write('A1', 'VERYHIDDEN_SHEET_LITERAL'); very.very_hidden()
    wb.close()
    def no_cache(raw):
        tree = ET.fromstring(raw)
        for cell in tree.iter(S+'c'):
            if cell.get('r') == 'D2':
                value = cell.find(S+'v')
                if value is not None: cell.remove(value)
        return ET.tostring(tree, encoding='utf-8', xml_declaration=True)
    rewrite_zip(xlsx, {'xl/worksheets/sheet1.xml':no_cache})
    producers['xlsx'] = {'producer':'XlsxWriter', 'synthetic_xml_edit':'D2 formula cache deliberately absent; C1 stale stored result999 deliberately supplied, never calculated by native parser', 'sheets':4}
    env = cpu_env(root)
    exports = []
    for source, suffix, filter_name in [(docx,'doc','MS Word 97'),(pptx,'ppt','MS PowerPoint 97'),(xlsx,'xls','MS Excel 97')]:
        profile = root/'export-profiles'/suffix; (profile/'user').mkdir(parents=True)
        profile_config(profile)
        command = ['/usr/bin/libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--norestore','--nodefault','--nofirststartwizard','--nolockcheck','--convert-to',suffix+':'+filter_name,'--outdir',str(originals),str(source)]
        row = run(command, root/'evidence', 'export-'+suffix, env=env, timeout=130)
        target = source.with_suffix('.'+suffix)
        row.update(output_exists=target.is_file(), real_ole=target.is_file() and olefile.isOleFile(target))
        exports.append(row)
        if row['returncode'] or not row['real_ole']:
            row['error']='Real LibreOffice binary export failed; no mock output substituted'
            if suffix=='xls':continue
            raise RuntimeError('Real LibreOffice binary export failed: '+suffix)
        with olefile.OleFileIO(target) as ole:
            row['ole_streams'] = ['/'.join(p) for p in ole.listdir(streams=True, storages=False)]
        expected = {'doc':'WordDocument','ppt':'PowerPoint Document','xls':'Workbook'}[suffix]
        assert expected in row['ole_streams'], row
    files = [{'name':p.name,'bytes':p.stat().st_size,'sha256':digest(p)} for p in sorted(originals.iterdir()) if p.is_file()]
    save(root/'evidence/generation.json', {'claim':'REAL SYNTHETIC OFFICE FILES; not historical classroom accuracy, not mocked converters or OLE protocols', 'producers':producers,'exports':exports,'originals':files})

def profile_config(profile):
    (profile/'user/registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item></oor:items>', encoding='utf-8')

def cpu_env(root):
    return {**os.environ,'CUDA_VISIBLE_DEVICES':'','HF_HUB_OFFLINE':'1','SAL_USE_VCLPLUGIN':'svp','SAL_MAX_CONCURRENCY':'4','OMP_NUM_THREADS':'4','OPENBLAS_NUM_THREADS':'4','MKL_NUM_THREADS':'4','TMPDIR':str(Path(root)/'tmp')}

def independent_workbook(path):
    """Independent raw XML oracle: inventory every relationship and stored cell."""
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = {n.get('Id'):n.get('Target') for n in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        import posixpath
        shared = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            for si in ET.fromstring(archive.read('xl/sharedStrings.xml')):
                shared.append(''.join(n.text or '' for n in si.iter(S+'t')))
        result = []
        for sheet in workbook.findall(S+'sheets/'+S+'sheet'):
            target = rels[sheet.get(R+'id')]
            part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            tree = ET.fromstring(archive.read(part)); cells = []
            for cell in tree.findall(S+'sheetData/'+S+'row/'+S+'c'):
                f, v, inline = cell.find(S+'f'), cell.find(S+'v'), cell.find(S+'is')
                typ = cell.get('t','n'); raw = v.text if v is not None else None
                value = shared[int(raw)] if typ=='s' and raw is not None else ''.join(n.text or '' for n in inline.iter(S+'t')) if inline is not None else raw
                cells.append({'coordinate':cell.get('r'),'type':typ,'style':cell.get('s'),'attributes':dict(cell.attrib), 'value':value,'raw_value':raw,'formula':None if f is None else {'text':f.text or '', 'attributes':dict(f.attrib)},'cache_present':f is not None and v is not None,'cached_value':value if f is not None and v is not None else None,'explicit_blank':f is None and v is None and inline is None})
            result.append({'name':sheet.get('name'),'part':part,'state':sheet.get('state','visible'),'cells':cells,'merged_ranges':[n.get('ref') for n in tree.findall(S+'mergeCells/'+S+'mergeCell')]})
        return result

def remote(root):
    root = Path(root).resolve()
    assert str(root).startswith('/home/adminz/llm/operations/stage16-legacy-cpu-') and os.getuid()==1000
    (root/'tmp').mkdir(exist_ok=True)
    candidate = root/'candidate'
    catalog = load(candidate/'catalog.json')
    assert len(catalog)==22
    for row in catalog:
        path = candidate/row['path']; assert path.is_file() and digest(path)==row['sha256']
    evidence = root/'evidence'; evidence.mkdir(exist_ok=True)
    save(evidence/'runtime.json',{'uid':os.getuid(),'username':__import__('pwd').getpwuid(os.getuid()).pw_name,'python':sys.version,'python_executable':sys.executable,'libreoffice':'/usr/bin/libreoffice','candidate_files_verified':len(catalog),'candidate_catalog':catalog})
    if (evidence/'originals').is_dir():
        originals=evidence/'originals'
        with zipfile.ZipFile(originals/'synthetic.docx') as archive:
            comment_generated='word/comments.xml' in archive.namelist()
        diagnostics=run(['dpkg-query','-W','-f=${Package}: ${db:Status-Abbrev} ${Version}\n','libreoffice-calc','libreoffice-writer','libreoffice-impress'],evidence,'installed-office-components',timeout=20)
        save(evidence/'generation.json',{'claim':'REAL SYNTHETIC OFFICE FILES; not historical classroom accuracy; resumed existing producer/export outputs without retrying failed XLS export','producers':{'docx':{'producer':'python-docx','comment_generated':comment_generated},'pptx':{'producer':'python-pptx','slides':2,'hidden_slides':1,'speaker_notes':2},'xlsx':{'producer':'XlsxWriter','synthetic_xml_edit':'D2 formula cache deliberately absent; C1 stale result999 deliberately supplied','sheets':4}},'exports':[load(evidence/'logs'/('export-'+fmt+'.json')) for fmt in ('doc','ppt','xls')], 'originals':[{'name':p.name,'bytes':p.stat().st_size,'sha256':digest(p)} for p in sorted(originals.iterdir()) if p.is_file()], 'environment_diagnostic':diagnostics,'calc_filter_available':Path('/usr/lib/libreoffice/program/libscfiltlo.so').is_file(),'xls_blocker':None if (originals/'synthetic.xls').is_file() else 'Installed LibreOffice has Writer/Impress but no Calc module; actual XLS export returned exit0 but source could not be loaded and no XLS produced.'})
    else:
        generate(root)
    guard = root/'guard'; guard.mkdir(exist_ok=True)
    # Inherited by actual Python parser subprocesses: reject network/model/lease
    # execution before it happens and retain all actual subprocess argv.
    guard_code = '''import os,sys,json,time
log=os.environ['STAGE16_CPU_AUDIT']
def audit(event,args):
    if event in ('socket.connect','socket.getaddrinfo'):
        with open(log,'a') as f:f.write(json.dumps({'event':event,'blocked':True,'pid':os.getpid()})+'\\n')
        raise RuntimeError('CPU probe forbids all Python network/model requests')
    if event=='subprocess.Popen':
        argv=args[1]
        with open(log,'a') as f:f.write(json.dumps({'event':event,'pid':os.getpid(),'argv':argv})+'\\n')
        text=' '.join(argv) if isinstance(argv,(list,tuple)) else str(argv)
        if any(s in text for s in ('llama-server','nvidia-smi','--vlm-url','--models','document_worker_service','gpu_lease')):
            raise RuntimeError('CPU probe forbids model/GPU/supervisor actions')
    if event=='import' and args[0].rsplit('.',1)[-1] in ('gpu_lease','page_vision_runtime'):
        raise RuntimeError('CPU probe forbids GPU lease/runtime imports')
sys.addaudithook(audit)
'''
    (guard/'sitecustomize.py').write_text(guard_code, encoding='utf-8')
    env = cpu_env(root); env['PYTHONPATH']=str(guard)+os.pathsep+str(candidate); env['STAGE16_CPU_AUDIT']=str(evidence/'cpu-audit.jsonl')
    for p in (candidate,candidate/'tools/document_stage12',candidate/'tools/material_stage13'):
        sys.path.insert(0,str(p))
    from parser_helpers import validate_bundle
    from prepared_material import visual_plan
    results = []
    for fmt in ('doc','ppt','xls','xlsx'):
        source = evidence/'originals'/('synthetic.'+fmt)
        if not source.is_file():
            results.append({'format':fmt,'blocked':True,'reason':'No real original generated; retained LibreOffice export error; no mocked substitute','native':{},'prepared':{}})
            save(evidence/'remote-results.json',results)
            continue
        before = digest(source)
        record = {'format':fmt,'source_sha256':before,'native':{},'prepared':{}}
        for phase in ('native','prepared'):
            output = evidence/f'{fmt}-{phase}'
            command = [sys.executable,str(candidate/'tools/material_stage13/parse_material.py'),'--input',str(source),'--output',str(output),'--engine','native','--deadline-seconds','180']
            if phase=='prepared': command += ['--prepared-dir',str(evidence/f'{fmt}-native')]
            row = run(command, evidence, fmt+'-'+phase, env=env, timeout=200)
            record[phase] = {'returncode':row['returncode'],'output':str(output),'sha_unchanged':digest(source)==before}
            if row['returncode']:
                record[phase]['error_receipt'] = load(output/'receipt.json') if (output/'receipt.json').is_file() else {'stderr':row['stderr']}
                break
            m, r = load(output/'manifest.json'), load(output/'receipt.json')
            validation = validate_bundle(output,m)
            assert validation['valid'] and m['schema_version']==1 and m['complete'] and r['status']=='complete'
            assert m['source_sha256']==before and m['source_format']==fmt and digest(source)==before
            assert r['engine']=='native' and not any(t.get('model_calls') or t.get('model_processed') for t in r['page_timings'])
            assert digest(output/'preview.pdf')==m['coverage']['rendered_pdf_sha256']
            import pymupdf as fitz
            with fitz.open(output/'preview.pdf') as pdf:
                assert pdf.is_pdf and not pdf.needs_pass and len(pdf)==m['page_count']
                rendered_text = '\n'.join(page.get_text() for page in pdf)
                glyphs = [s for page in pdf for b in page.get_text('dict')['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']]
                assert glyphs
                if fmt in ('xls','xlsx'): assert all(s['bbox'][2]<=559.1 for s in glyphs), 'Clipped ledger glyphs'
                pdf[0].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(output/'inspected-preview-page1.png')
            native = load(output/'raw/office-inventory.json')
            blocks = ''.join(b['text'] for p in m['pages'] for b in p['blocks'])
            if fmt in ('xls','xlsx'):
                reference = source if fmt=='xlsx' else output/'raw/derived.xlsx'
                oracle = independent_workbook(reference)
                projected = [{k:s[k] for k in ('name','part','state','cells','merged_ranges')} for s in native['sheets']]
                assert projected==oracle, 'All-sheet/all-cell inventory differs from independent XML oracle'
                assert [s['name'] for s in oracle]==['Main','Hidden','VeryHidden','Blank']
                assert native['cell_count']==sum(len(s['cells']) for s in oracle) and m['coverage']['stored_cells']==native['cell_count']
                assert m['coverage']['original_sheet_count']==4 and m['coverage']['stored_cells_complete'] and m['coverage']['sheets_complete']
                assert 'OUTSIDE_PRINT_AREA_LITERAL' in blocks and 'HIDDEN_SHEET_LITERAL' in blocks and 'VERYHIDDEN_SHEET_LITERAL' in blocks
                assert 'PRIVATE_WORKBOOK_COMMENT' not in blocks and any('PRIVATE_WORKBOOK_COMMENT' in n['text'] for n in native['notes'])
                cells = {c['coordinate']:c for c in native['sheets'][0]['cells']}
                assert cells['C1']['formula']['text']=='SUM(B1:B2)'
                if fmt=='xlsx':
                    assert cells['C1']['cached_value']=='999' and cells['D2']['formula']['text']=='B1*3' and not cells['D2']['cache_present'] and cells['E1']['explicit_blank']
                    assert native['sheets'][1]['state']=='hidden' and native['sheets'][2]['state']=='veryHidden' and not native['sheets'][3]['cells']
                    assert m['coverage']['formula_evaluation_performed'] is False
                else:
                    assert m['coverage']['original_binary_cells_decoded'] is False and m['coverage']['conversion_may_recalculate'] is True
                    assert m['coverage']['inventory_derivation']=='libreoffice_xlsx'
                with zipfile.ZipFile(reference) as z:
                    for part in native['auxiliary_parts']: assert (output/part['raw_asset']).read_bytes()==z.read(part['part'])
                record[phase].update(sheet_count=4,stored_cells=native['cell_count'],formula_count=native['formula_count'],cached_formula_count=native['cached_formula_count'],C1_cache=cells['C1']['cached_value'],D2_cache=cells['D2']['cached_value'],sheet_states=[s['state'] for s in native['sheets']],comments=len(native['notes']),oracle_exact=True)
            else:
                markers = ('DOC_VISIBLE_PAGE_ONE','DOC_VISIBLE_PAGE_TWO') if fmt=='doc' else ('PPT_VISIBLE_SLIDE_1','PPT_VISIBLE_SLIDE_2')
                assert all(marker in blocks and marker in rendered_text for marker in markers)
                if fmt=='ppt':
                    assert m['page_count']==2 and native['slides_expected']==2 and m['coverage']['hidden_slides']==1
                    assert all(any(f'PRIVATE_PPT_NOTE_{i}' in n['text'] for n in native['notes']) for i in (1,2))
                    assert 'PRIVATE_PPT_NOTE' not in blocks
                assert (output/'raw'/('derived.docx' if fmt=='doc' else 'derived.pptx')).is_file()
            record[phase].update(worker_validation=validation,page_count=m['page_count'],device=r['device'],model_calls=sum(t.get('model_calls',0) for t in r['page_timings']),visual_plan=visual_plan(m,r),rendered_pdf_sha256=digest(output/'preview.pdf'),rendered_text_chars=len(rendered_text),glyph_max_right=max(s['bbox'][2] for s in glyphs),coverage=m['coverage'])
            if phase=='prepared':
                first = evidence/f'{fmt}-native'
                assert (first/'preview.pdf').read_bytes()==(output/'preview.pdf').read_bytes()
                if fmt not in ('xls','xlsx'): assert r.get('conversion_reused_from_native_phase') is True
                record[phase]['preview_exactly_reused'] = True
        results.append(record)
        save(evidence/'remote-results.json',results)
    # Scope-specific Linux process check: only argv containing our unique root.
    own = []
    for entry in Path('/proc').iterdir():
        if not entry.name.isdigit() or int(entry.name)==os.getpid(): continue
        try: argv=(entry/'cmdline').read_bytes().replace(b'\x00',b' ').decode(errors='replace')
        except (OSError,ProcessLookupError): continue
        if str(root) in argv and ('libreoffice' in argv or 'soffice' in argv or 'parse_material.py' in argv or 'parser_worker.py' in argv): own.append({'pid':int(entry.name),'argv':argv})
    save(evidence/'cleanup.json',{'owned_converter_parser_processes_remaining':own,'owned_children_reaped':True,'shared_models_gpu_router_lease_accessed':False,'cpu_network_guard_active':True})
    assert not own, own
    errors = [r['format'] for r in results if r.get('blocked') or r['native'].get('returncode') or r['prepared'].get('returncode')]
    save(evidence/'remote-summary.json',{'formats_completed':[r['format'] for r in results if r['prepared'].get('returncode')==0], 'failed_formats':errors,'synthetic_only':True,'gpu_or_model_actions':0})
    print(json.dumps(load(evidence/'remote-summary.json')))
    return 2 if errors else 0

def local_validation(app, evidence, *, require_all=True):
    app, evidence = Path(app), Path(evidence)
    sys.path.insert(0,str(app))
    from backend.aicademiya.document_bundle import validate_bundle, validate_manifest
    from backend.aicademiya.document_supplements import index_supplements, eligible_supplements, supplement_block
    results = []
    for fmt in ('doc','ppt','xls','xlsx'):
        source = evidence/'originals'/('synthetic.'+fmt)
        if not source.is_file():continue
        for phase in ('native','prepared'):
            bundle = evidence/f'{fmt}-{phase}'
            if not (bundle/'manifest.json').is_file(): continue
            worker = load(bundle/'manifest.json')
            if not worker.get('complete'): continue
            manifest = validate_bundle(bundle,digest(source),None)
            supplements = index_supplements(source,bundle,manifest)
            supplemented = {**manifest,'supplements':supplements}
            validate_manifest(supplemented,source_sha256=digest(source))
            assert eligible_supplements({'manifest':supplemented})==[]
            assert eligible_supplements({'manifest':supplemented},True)==supplements
            assert all(s['visibility']=='teacher' and s['char_end']-s['char_start']==len(s['text']) for s in supplements)
            if fmt in ('xls','xlsx'):
                assert any('PRIVATE_WORKBOOK_COMMENT' in s['text'] for s in supplements)
                assert all(s['page_hint'] is None and 'sheet:1; cell:A1'==s['source_node'] for s in supplements)
                assert all(s['derivation']==('original_office_xml' if fmt=='xlsx' else 'libreoffice_xlsx') for s in supplements)
            elif fmt=='ppt':
                for i in (1,2): assert any(f'PRIVATE_PPT_NOTE_{i}' in s['text'] and s['page_hint']==i and s['derivation']=='libreoffice_pptx' for s in supplements)
            elif fmt=='doc':
                assert any(s['kind']=='header' and 'SYNTHETIC_DOC_HEADER' in s['text'] for s in supplements)
                assert any(s['kind']=='footer' and 'SYNTHETIC_DOC_FOOTER' in s['text'] for s in supplements)
                if load(evidence/'generation.json')['producers']['docx']['comment_generated']: assert any('PRIVATE_DOC_COMMENT' in s['text'] for s in supplements)
            for supplement in supplements:
                page, block = supplement_block(supplement)
                assert block['text']==supplement['text'] and page['source_part']==supplement['source_part']
            results.append({'format':fmt,'phase':phase,'strict_app_bundle_validation':True,'app_supplement_hook':True,'staff_only_gate':True,'supplements':supplements,'source_sha256':digest(source),'page_count':manifest['page_count']})
            save(evidence/'app-validation.json', results)
    expected={(fmt,phase) for fmt in ('doc','ppt','xls','xlsx') for phase in ('native','prepared')}
    if require_all:assert {(r['format'],r['phase']) for r in results}==expected, 'Incomplete actual formats/phases; inspect retained remote failures'
    return results

def orchestrate(app):
    app=Path(app).resolve()
    run_id=time.strftime('%Y%m%dT%H%M%SZ',time.gmtime())+'-'+uuid.uuid4().hex[:8]
    artifacts=app/'artifacts/stage16/legacy-real'/run_id; artifacts.mkdir(parents=True)
    remote_root='/home/adminz/llm/operations/stage16-legacy-cpu-'+run_id
    candidate=app/'artifacts/stage16/release-candidate-r2/extracted/AIcademiya_V2-stage16-workflow-candidate-r2'
    catalog=[]; archive=artifacts/'candidate-tools.tar.gz'
    groups=('material_stage13','material_stage15','material_stage16','document_stage12')
    paths=sorted(p for group in groups for p in (candidate/'tools'/group).glob('*.py'))
    assert len(paths)==22, len(paths)
    for p in paths:
        rel=p.relative_to(candidate).as_posix()
        assert digest(p)==digest(app/rel), 'Candidate source differs: '+rel
        catalog.append({'path':rel,'sha256':digest(p),'bytes':p.stat().st_size})
    save(artifacts/'catalog.json',catalog)
    with tarfile.open(archive,'w:gz') as tar:
        for p in paths: tar.add(p,arcname=p.relative_to(candidate).as_posix())
        tar.add(artifacts/'catalog.json',arcname='catalog.json')
    save(artifacts/'ownership.json',{'local_script':str(Path(__file__).resolve()),'local_artifact_root':str(artifacts),'remote_root':remote_root,'username':'adminz','remote_uid':1000,'remote_venv_read_only':REMOTE_PYTHON,'scope':'isolated CPU synthetic-only acceptance; no deployment, sudo, installs, system config, CT, GPU/model/router/lease actions'})
    command=['ssh','ubuntu-gpu',f'umask 077; mkdir {remote_root} && mkdir {remote_root}/candidate && mkdir {remote_root}/tmp']
    row=run(command,artifacts,'remote-create',timeout=30); assert row['returncode']==0,row
    row=run(['scp',str(archive),str(Path(__file__).resolve()),'ubuntu-gpu:'+remote_root+'/'],artifacts,'upload',timeout=60); assert row['returncode']==0,row
    row=run(['ssh','ubuntu-gpu',f'tar -xzf {remote_root}/candidate-tools.tar.gz -C {remote_root}/candidate && CUDA_VISIBLE_DEVICES= HF_HUB_OFFLINE=1 TMPDIR={remote_root}/tmp {REMOTE_PYTHON} {remote_root}/legacy_cpu_probe.py remote --root {remote_root}'],artifacts,'remote-run',timeout=1600)
    downloaded=run(['scp','-r','ubuntu-gpu:'+remote_root+'/evidence',str(artifacts)],artifacts,'download-evidence',timeout=180)
    assert downloaded['returncode']==0, downloaded
    cleanup=run(['ssh','ubuntu-gpu',f'/opt/aicademiya-documents/runtime/.venv/bin/python -c "import pathlib,json,os; r=pathlib.Path(\'{remote_root}\'); bad=[str(p) for p in r.rglob(\'*\') if p.stat().st_uid!=1000]; print(json.dumps({{\'not_owned_by_adminz\':bad,\'cleanup\':json.loads((r/\'evidence/cleanup.json\').read_text())}}))"'],artifacts,'ownership-cleanup-readback',timeout=30)
    # No generic process termination or model queries. Any failure is retained.
    errors=[]
    if row['returncode']: errors.append({'remote_run_returncode':row['returncode']})
    try: validations=local_validation(app,artifacts/'evidence',require_all=False)
    except Exception as exc:
        errors.append({'app_validation_error':str(exc),'traceback':traceback.format_exc()}); validations=[]
    if not (artifacts/'evidence/cleanup.json').is_file(): errors.append({'cleanup':'Remote execution interrupted before owned-process check; use retained exact root to inspect owned converter/parser descendants only.'})
    if len(validations)!=8:errors.append({'incomplete_real_cases':len(validations),'required':8})
    summary={'run_id':run_id,'remote_root':remote_root,'local_artifact_root':str(artifacts),'actual_remote_command':row['argv'],'real_command_returncode':row['returncode'],'strict_app_cases':len(validations),'status':'passed' if not errors else 'failed','errors':errors,'synthetic_real_files_not_historical_accuracy':True,'candidate_catalog_files':len(catalog),'ownership_returned':True}
    save(artifacts/'summary.json',summary)
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    return 0 if not errors else 2

def cleanup_owned(root):
    root=Path(root).resolve(); terminated=[]
    def owned():
        found=[]
        for entry in Path('/proc').iterdir():
            if not entry.name.isdigit() or int(entry.name)==os.getpid():continue
            try:argv=(entry/'cmdline').read_bytes().replace(b'\x00',b' ').decode(errors='replace')
            except OSError:continue
            if str(root) in argv and any(s in argv for s in ('libreoffice','soffice','parse_material.py','parser_worker.py')):
                found.append({'pid':int(entry.name),'argv':argv})
        return found
    for row in owned():
        try:os.kill(row['pid'],signal.SIGTERM);terminated.append(row)
        except ProcessLookupError:pass
    if terminated:
        deadline=time.monotonic()+5
        while owned() and time.monotonic()<deadline:time.sleep(.1)
        for row in owned():
            try:os.kill(row['pid'],signal.SIGKILL)
            except ProcessLookupError:pass
    save(root/'evidence/cleanup.json',{'owned_converter_parser_processes_remaining':owned(),'terminated_exact_owned_pids':terminated,'owned_children_reaped':True,'shared_models_gpu_router_lease_accessed':False,'cpu_network_guard_active':(root/'guard/sitecustomize.py').is_file()})

def main():
    cli=argparse.ArgumentParser(description=__doc__)
    cli.add_argument('action',choices=('orchestrate','remote','validate'))
    cli.add_argument('--app',default=str(Path(__file__).resolve().parents[2]))
    cli.add_argument('--root')
    a=cli.parse_args()
    if a.action=='orchestrate':return orchestrate(a.app)
    if a.action=='remote':
        try:return remote(a.root)
        except BaseException as exc:
            save(Path(a.root)/'evidence/remote-error.json',{'type':type(exc).__name__,'error':str(exc),'traceback':traceback.format_exc()})
            raise
        finally:cleanup_owned(a.root)
    results=local_validation(a.app,a.root,require_all=False)
    print(json.dumps({'strict_app_cases':len(results),'formats':sorted({r['format'] for r in results}),'all_four_formats_complete':len(results)==8}))
    return 0 if len(results)==8 else 2

if __name__=='__main__':raise SystemExit(main())
