"""Stage isolated app-only snapshot + synthetic inputs on CT255, never production."""
from pathlib import Path
import hashlib,io,json,shlex,socket,subprocess,sys,tarfile,uuid
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT),str(ROOT/'tools')]
from tools.build_v2_release import collect
from run_stage12_real_reading import make_catalog
SSH=['ssh','-4','-i','C:/Users/17193/.ssh/id_ed25519_aicademiya_pve','-o','BatchMode=yes','-o','IdentitiesOnly=yes','-o','StrictHostKeyChecking=yes','root@192.168.10.16']

def guest(args,timeout=45):
    r=subprocess.run(SSH+[shlex.join(['pct','exec','255','--']+args)],capture_output=True,text=True,timeout=timeout)
    if r.returncode:raise RuntimeError(r.stderr or r.stdout)
    return r.stdout

def main():
    from pptx import Presentation
    tag=uuid.uuid4().hex[:8];out=ROOT/'artifacts/stage14'/('app-acceptance-'+tag);out.mkdir();fixture=out/'fixture';fixture.mkdir();make_catalog(fixture)
    lines=['SOURCE_LEDGER: synthetic exact-reading acceptance.\n']
    for number in range(1,201):
        lines.append(f'Record {number:03}: This is a synthetic source passage for sequential reading. The learner must cite the stored words rather than invent missing text. Evidence remains tied to its exact document revision and published edition.\n')
    lines.append('Closing record: The amber heron keeps the final copper notebook.\n');literal=''.join(lines);(fixture/'source-ledger.txt').write_bytes(literal.encode())
    note='SYNTHETIC speaker note: The hidden reference marker is APRICOT-LANTERN. Preserve this exact sentence.'
    ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='SOURCE_NOTES: isolated permission test'
    slide.placeholders[1].text='This is synthetic public slide text. Speaker notes are separate source content. The source file is preserved exactly, while student AI context follows explicit sharing and publication.'
    slide.notes_slide.notes_text_frame.text=note;ppt.save(fixture/'source-notes.pptx')
    files=collect(ROOT);selected={n:v['data']for n,v in files.items()if n.startswith(('backend/','frontend/dist/'))or n in ('tools/document_stage12/document_ssh_bridge.py','tools/material_stage13/source_text.py','tools/backup_data.py')}
    # Snapshot only app modules: independent mask/parser edits cannot race this package.
    qa='/var/lib/aicademiya-qa/stage14-app-'+tag;source=qa+'/source';data=qa+'/data';base='http://192.168.32.55:18774';unit='aicademiya-qa-stage14-app-'+tag
    live=guest(['readlink','-f','/opt/aicademiya/current']).strip();python=live+'/.venv/bin/python'
    guest(['python3','-c',"import socket;s=socket.socket();s.bind(('192.168.32.55',18774));s.close()"])
    worker=[python,source+'/tools/document_stage12/document_ssh_bridge.py','--ssh-host','aicademiya-docs@192.168.32.5','--ssh-key','/var/lib/aicademiya/worker-ssh-stage14/id_ed25519','--known-hosts','/var/lib/aicademiya/worker-ssh-stage14/known_hosts','--timeout','1200']
    server='''from pathlib import Path
import os,json,time,uvicorn
from backend.aicademiya.main import create_app
f=Path(__file__).resolve().parent
app=create_app(data_dir=Path(os.environ['AICADEMIYA_DATA_DIR']),catalog_path=f/'catalog.json',frontend_dir=Path(SOURCE)/'frontend/dist',demo_enabled=True,document_worker_command=WORKER,document_worker_timeout=1250)
original=app.state.tutor.completion_message
def record(settings,messages,**kwargs):
 start=time.monotonic();row={'model':settings['model'],'messages':messages,'response_schema':kwargs.get('response_format'),'real_provider_call':True}
 try:response=original(settings,messages,**kwargs);row['response']=response;return response
 except Exception as e:row['error_type']=type(e).__name__;raise
 finally:
  row['seconds']=time.monotonic()-start
  with (f/'model-exchanges.jsonl').open('a')as log:log.write(json.dumps(row,ensure_ascii=False)+'\\n')
app.state.tutor.completion_message=record
uvicorn.run(app,host='192.168.32.55',port=18774,log_level='warning')
'''.replace('SOURCE',repr(source)).replace('WORKER',repr(worker))
    (fixture/'server.py').write_text(server,encoding='utf-8')
    blobs={'source/'+n:b for n,b in selected.items()}
    blobs.update({p.relative_to(fixture).as_posix():p.read_bytes()for p in fixture.rglob('*')if p.is_file()})
    hashes={n:hashlib.sha256(b).hexdigest()for n,b in blobs.items()};buffer=io.BytesIO()
    blobs['snapshot-manifest.json']=json.dumps(hashes,indent=2).encode()
    with tarfile.open(fileobj=buffer,mode='w:gz')as tar:
        for name,blob in blobs.items():
            info=tarfile.TarInfo(name);info.size=len(blob);info.mode=0o644;tar.addfile(info,io.BytesIO(blob))
    code="from pathlib import Path;import sys,tarfile,io,os,json,hashlib,pwd;user=pwd.getpwnam('aicademiya');assert user.pw_uid==999;p=Path("+repr(qa)+");b=sys.stdin.buffer.read();assert hashlib.sha256(b).hexdigest()=="+repr(hashlib.sha256(buffer.getvalue()).hexdigest())+";p.mkdir(mode=0o700);t=tarfile.open(fileobj=io.BytesIO(b),mode='r:gz');t.extractall(p,filter='data');t.close();expected=json.loads((p/'snapshot-manifest.json').read_text());assert all(hashlib.sha256((p/n).read_bytes()).hexdigest()==h for n,h in expected.items());\nfor q in [p,*p.rglob('*')]:os.chown(q,user.pw_uid,user.pw_gid)\nprint(json.dumps({'verified_files':len(expected),'qa':str(p)}))"
    r=subprocess.run(SSH+[shlex.join(['pct','exec','255','--','python3','-c',code])],input=buffer.getvalue(),capture_output=True,timeout=50)
    if r.returncode:raise RuntimeError(r.stderr.decode())
    env={'PYTHONPATH':source,'PYTHONDONTWRITEBYTECODE':'1','LANG':'C.UTF-8','AICADEMIYA_DATA_DIR':data,'AICADEMIYA_DEMO_ENABLED':'true','AICADEMIYA_AI_MODE':'openai','AICADEMIYA_AI_PROVIDER':'aicademiya-router','AICADEMIYA_AI_BASE_URL':'http://127.0.0.1:18280/v1','AICADEMIYA_AI_MODEL':'AIcademiya-64kx3-Qwen3.8-27B','AICADEMIYA_AI_API_KEY':'','AICADEMIYA_DOCUMENT_FORMATS':'pdf,pptx,docx,doc,txt,md,png,jpg,jpeg,webp','AICADEMIYA_WEB_SEARCH_ENABLED':'false','TAVILY_API_KEY':'','AICADEMIYA_ALLOWED_ORIGINS':base,'AICADEMIYA_COOKIE_SECURE':'false','NO_PROXY':'127.0.0.1,localhost,192.168.32.55','no_proxy':'127.0.0.1,localhost,192.168.32.55'}
    start=['systemd-run','--unit='+unit,'--collect','--uid=aicademiya','--gid=aicademiya','--working-directory='+source,'--property=UMask=0077','--property=RuntimeMaxSec=3600','--property=NoNewPrivileges=true','--property=ProtectSystem=strict','--property=ProtectHome=true','--property=PrivateTmp=true','--property=ReadWritePaths='+qa]+['--setenv='+k+'='+v for k,v in env.items()]+[python,qa+'/server.py']
    meta={'scope':'Isolated app snapshot, NOT final release/promotion','tag':tag,'artifact_dir':str(out),'qa':qa,'data':data,'source':source,'base_url':base,'unit':unit,'interpreter':python,'production_release':live,'start_command':start,'environment':env,'source_snapshot_hashes':hashes,'txt_sha256':hashes['source-ledger.txt'],'txt_characters':len(literal),'closing_line':lines[-1].strip(),'pptx_sha256':hashes['source-notes.pptx'],'note':note,'synthetic_accounts_only':True,'worker':'installed r5 restricted key; no root worker replacement','server_started':False}
    (out/'meta.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n');print(json.dumps({'staged':True,'meta':str(out/'meta.json'),'source_snapshot_files':len(selected),'remote_readback':json.loads(r.stdout)}))
if __name__=='__main__':main()
